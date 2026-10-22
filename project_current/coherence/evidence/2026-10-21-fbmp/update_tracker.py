"""D-074: update tracker rows COH-013 and COH-004 in place for the D-073 pier export (rerunnable, idempotent).

    .venv\\Scripts\\python.exe project_current\\coherence\\evidence\\2026-10-21-fbmp\\update_tracker.py

Edits only the named cells of the two rows and the review-date line; never touches another row, a response
(column G), an ID or the formulas. Saves with openpyxl, then recalculates in Excel (geometry_tables/excel_recalc.py)
so the COUNTIFS summary cells keep cached values. Writes tracker_update.json beside this script.
"""
from __future__ import annotations

import hashlib
import json
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[4]
sys.path.insert(0, str(ROOT))
from geometry_tables.excel_recalc import excel_available, recalculate  # noqa: E402

WORKBOOK = ROOT / 'project_current' / 'coherence' / 'Coherence_Tracker.xlsx'
HERE = Path(__file__).resolve().parent
REVIEWED = datetime(2026, 10, 21)
PIER_XML = '8be45bb889f5c955d0547ce6a538118bbce87600d40a65b094255e6b9b76e3b6'

COH_013 = {
    'B': 'Closed',
    'E': ('Resolved: every selected cap report and case names a registered export. End Bent 1 and End Bent 2 (10-19 reports and cases, D-051) '
          'name EB-XML and EB2-XML (the 19 October fixed exports, D-049); the pier (10-21 report and case, D-073) names PIER-XML, the 21 October '
          f'rerun with the 45 ksi pile material (SHA-256 {PIER_XML[:12]}...). Historical finding (12 to 16 October): the 12 and 14 October reports '
          'named exports that were not in project_current (pier c8420fad..., End Bent 1 23bb3afc..., End Bent 2 56f78267...).'),
    'F': 'None for export identity. Native run records and the pile-head reruns remain under COH-004 and COH-028.',
    'J': ('10 Oct: the supplied EB1 export 8414fb... matched the 10 October report. The 12 and 14 October reports then named new exports. 16 Oct: the supplied '
          'pier export c8420fad... was compared (pile heads at EL 41.7717 ft, COH-028) and left unsubmitted by user direction. See '
          'coherence/evidence/2026-10-16-fbmp-pile-heads/pile_heads.json. 19 Oct (D-049, D-051): the three fixed exports were registered and the cap '
          'notebook rerun on them, so the 10-19 reports and cases name them. 21 Oct (D-073): the rerun pier export is registered and the pier cap and pile '
          'cases were rerun from it; machine checks PIER:analysis-identity and EB:analysis-identity read MATCH. Closed by user direction, 21 Oct (D-074). '
          'See coherence/evidence/2026-10-21-fbmp/source_receipt.json, 2026-10-21-cap-rerun/source_receipt.json and 2026-10-21-pile-reports/source_receipt.json.'),
    'K': ('EB-REPORT; EB-XML; coherence/evidence/end_bent_provenance.json; coherence/evidence/2026-10-10/source_update.json; PIER-REPORT; PIER-XML; '
          'EB2-REPORT; EB2-XML; coherence/evidence/2026-10-12/source_update.json; coherence/evidence/2026-10-16-fbmp-pile-heads/pile_heads.json; D-049; D-051; '
          'D-073; D-074; coherence/evidence/2026-10-21-fbmp/source_receipt.json; coherence/evidence/2026-10-21-cap-rerun/source_receipt.json; '
          'coherence/machine_checks.json'),
    'L': REVIEWED,
}
COH_004 = {
    'E': ('Pier report/XML fingerprints match (the 10-21 report and case name the registered 21 October export, D-073) and bearing offsets reconcile. '
          'The 21 October export carries the same nodal loads as the 19 October one (only the pile material changed to 45 ksi), so the D-062 load review stands. '
          'Full current wind, below-pad weights, transfer moments, combinations and convergence remain unverified against T006.'),
    'J': ('16 Oct: the 12 Oct pier report names export c8420fad..., not the registered PIER-XML (COH-013), and the final runs need the current pile heads (COH-028). '
          '21 Oct (D-073): the pier export rerun with the 45 ksi pile material is registered and the 10-21 report and case name it (COH-013 closed, D-074); its nodal '
          'loads equal the 19 October export (coherence/evidence/2026-10-21-fbmp/source_receipt.json, comparison_with_previous). The full load comparison, native '
          'run records and the COH-028 pile-head reruns remain open.'),
    'K': ('PIER-XML; JOURNAL; coherence/evidence/pier_review_2026-10-09.xlsx, item 05; tasks/T006/response-disposition.json, item 05; '
          'coherence/evidence/2026-10-10/source_update.json; COH-013; COH-028; coherence/evidence/2026-10-16-fbmp-pile-heads/pile_heads.json; D-062; D-073; '
          'coherence/evidence/2026-10-19-fbmp-load-review/fbmp_load_review.md; coherence/evidence/2026-10-21-fbmp/source_receipt.json'),
    'L': REVIEWED,
}
REVIEW_LINE = 'Review updated: 21 October 2026. Source selections, recorded responses and closure evidence are preserved here.'


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> int:
    before = sha(WORKBOOK)
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb['Coherence items']
    assert [c.value for c in ws[8]][:3] == ['Item ID', 'Status', 'Area']
    rows = {ws.cell(r, 1).value: r for r in range(9, ws.max_row + 1) if ws.cell(r, 1).value}
    changed = {}
    for item, fields in (('COH-013', COH_013), ('COH-004', COH_004)):
        r = rows[item]
        assert ws.cell(r, 7).value is None or isinstance(ws.cell(r, 7).value, str)      # column G (response) is never written
        for col, value in fields.items():
            cell = ws[f'{col}{r}']
            if cell.value != value:
                changed[f'{item}:{col}{r}'] = {'before': str(cell.value), 'after': str(value)}
                cell.value = value
    if ws['A3'].value != REVIEW_LINE:
        changed['A3'] = {'before': ws['A3'].value, 'after': REVIEW_LINE}
        ws['A3'].value = REVIEW_LINE
    if not changed:
        print('tracker already updated')
        return 0
    wb.save(WORKBOOK)
    if not excel_available():
        raise RuntimeError('Excel is not available to recalculate the summary formulas; the saved workbook has no cached counts.')
    recalculate(WORKBOOK)
    cached = openpyxl.load_workbook(WORKBOOK, data_only=True)['Coherence items']
    summary = {'Outstanding': cached['B4'].value, 'Awaiting source': cached['D4'].value, 'Deferred': cached['F4'].value, 'Closed / N.A.': cached['H4'].value}
    assert all(isinstance(v, int) for v in summary.values()), summary
    record = {'decision': 'D-074', 'date': '2026-10-21', 'workbook': 'project_current/coherence/Coherence_Tracker.xlsx', 'sha256_before': before,
              'sha256_after': sha(WORKBOOK), 'changed_cells': changed, 'summary_after_recalc': summary,
              'rules': 'Edited in place; responses (column G), IDs and formulas untouched; no row added or renumbered; recalculated in Excel.'}
    (HERE / 'tracker_update.json').write_text(json.dumps(record, indent=2, ensure_ascii=False) + '\n', encoding='utf-8', newline='\n')
    print(json.dumps({'changed': sorted(changed), 'summary': summary, 'sha256_after': record['sha256_after'][:12]}, indent=1))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
