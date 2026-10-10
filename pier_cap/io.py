"""Portable case bundles, traceable workbook imports and review-copy exports."""
from copy import deepcopy
from pathlib import Path
from datetime import datetime,timezone
import csv
import hashlib
import io
import json
import re
from .model import INPUTS,DEFINITIONS,GEOMETRY,validate_case,evaluate,default_case,formula_trace,upgrade_case
from .force_audit import strength_audit

def load_case(source):
    try:
        if isinstance(source,(bytes,bytearray,memoryview)):data=json.loads(bytes(source).decode('utf-8-sig'))
        else:data=json.loads(Path(source).read_text(encoding='utf-8-sig'))
    except json.JSONDecodeError as exc:
        raise ValueError(f'Invalid JSON at line {exc.lineno}, column {exc.colno}: {exc.msg}. Select the original exported selected_case.json.') from exc
    except UnicodeDecodeError as exc:
        raise ValueError('The file is not UTF-8 JSON. Select the exported selected_case.json.') from exc
    if isinstance(data,dict) and 'review' in data and 'controls' in data:
        raise ValueError('This is a pile review JSON. Use Load pile review in the pile panel; use selected_case.json here for the cap.')
    if not isinstance(data,dict) or not all(isinstance(data.get(key),dict) for key in ('inputs','units','analysis','screening')):
        raise ValueError('This is not a saved cap case. Select selected_case.json from Export case + checks, not a report, equation trace or pile review JSON.')
    data=upgrade_case(data)
    validate_case(data)
    evaluate(data)
    return data

def write_case(case,path):
    case=upgrade_case(case);validate_case(case);path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    with path.open('x',encoding='utf-8') as f:json.dump(case,f,indent=2,ensure_ascii=False,allow_nan=False)
    return path

def input_formula(name,value):
    if isinstance(value,bool):rhs=str(value).lower()
    else:rhs=f'{value:.12g}'+(' '+INPUTS[name]['unit'] if INPUTS[name]['unit'] else '')
    return name+' = '+rhs

def export_bundle(case,root='exports',search_result=None,search_filter=None):
    e=evaluate(case);case=e.case
    if search_result:
        if search_result.force_mode not in ('matched','fixed'):raise ValueError('Unknown search force mode.')
        from .optimizer import filter_candidates,candidate_dc,candidate_governing,same_design_basis
        if not same_design_basis(search_result.base_case,case):raise ValueError('Search inputs no longer match this case. Run steel search again before exporting alternatives.')
        options=search_filter or {'max_dc':1.0,'scope':'all','objective':search_result.config['objective']}
        matching=filter_candidates(search_result,**options)
    path=Path(root)/datetime.now(timezone.utc).strftime('case-%Y%m%d-%H%M%S-%f')
    path.mkdir(parents=True,exist_ok=False)
    write_case(case,path/'selected_case.json')
    from .end_grid import enabled as end_grid_enabled,geometry as end_grid_geometry
    from .end_grid_views import end_figure,summary_html as end_grid_summary
    if end_grid_enabled(case):
        end_bars=end_grid_geometry(e)
        (path/'end_grid_geometry.json').write_text(json.dumps(end_bars,indent=2,ensure_ascii=False),encoding='utf-8')
        with (path/'end_grid_schedule.csv').open('w',newline='',encoding='utf-8-sig') as f:
            fields=['id','end','direction','bar','diameter','coordinate_in','nominal_coordinate_in','shift_in','plane_in','pitch_in','return_in','inside_diameter_in','hook_mode','placement_mode','placement_ok','placement_note','length_in','fit']
            writer=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore');writer.writeheader();writer.writerows(end_bars)
        for end in ('left','right'):
            end_figure(e,end).write_html(path/f'end_grid_{end}.html',include_plotlyjs=True,full_html=True)
    from .transverse import enabled as actual_transverse,scheduled_bars,bar_shape,shape_issues
    if actual_transverse(case):
        with (path/'transverse_bar_schedule.csv').open('w',newline='',encoding='utf-8-sig') as f:
            writer=csv.DictWriter(f,fieldnames=['id','run','kind','bar','zone','station_in']);writer.writeheader();writer.writerows(scheduled_bars(case))
        shapes=[dict(run=r,geometry=bar_shape(e,r),issues=shape_issues(e,r)) for r in case['transverse_detail']['runs']]
        (path/'transverse_shapes.json').write_text(json.dumps(shapes,indent=2,ensure_ascii=False),encoding='utf-8')
    with (path/'checks.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.writer(f);w.writerow(['Check','Status','D/C','Basis'])
        for c in e.checks:w.writerow([c.label,c.status,c.ratio,c.basis])
    from .model import bar_positions
    from .detailing import spacing_records
    with (path/'longitudinal_bar_positions.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.writer(f);w.writerow(['Region','Layer','Bar size','Across cap (in)','Above underside (in)','Diameter (in)','Additional'])
        for region in 'PB':
            for bar in bar_positions(e,region):
                w.writerow([region,bar['kind'],bar['bar'],bar['x'],bar['y'],bar['diameter'],bar['additional']])
    if e.longitudinal_layout is not None:
        geometry={name:e.value(name) for name in ('dc_N','dc_P','dc_B','d_N','d_P','d_B','dv','SP_N','SP_P','SP_B','SP_skin')}
        geometry.update(units='in',fitted=e.longitudinal_layout['fitted'],issues=e.longitudinal_layout['issues'],
            basis='Actual longitudinal coordinates inside the common envelope of all entered transverse runs; includes conservative envelope for span additions. Existing sectional equations use the resulting centroids and spacing. Hook-end congestion and development remain separate checks.')
        (path/'longitudinal_geometry.json').write_text(json.dumps(geometry,indent=2),encoding='utf-8')
    with (path/'rebar_clear_spacing.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.writer(f);w.writerow(['Region','Bars','Actual clear (in)','Required clear (in)','Status'])
        for region in 'PB':
            for r in spacing_records(e,region,bar_positions(e,region)):
                w.writerow([region,r['label'],r['actual'],r['required'],r['status']])
    from .visuals import section_figure,reinforcement_plan_figure,elevation_figure,hoop_figure,hoop_explanation_html,reinforcement_summary_html,clear_spacing_html,configuration_html
    from .cage_3d import layout_3d
    cage=layout_3d(e)
    cage.write_html(path/'cage_3d.html',include_plotlyjs=True,full_html=True)
    parts=['<!doctype html><html><head><meta charset="utf-8"><title>Cap reinforcement detail review</title><style>.cap-table{border-collapse:collapse;width:100%}.cap-table td,.cap-table th{padding:8px;border-bottom:1px solid #dce5ec;text-align:left}.cap-table th{background:#e7eef4}</style></head><body style="font:14px Arial;max-width:1200px;margin:auto">',configuration_html(e),reinforcement_summary_html(e)]
    for i,fig in enumerate([section_figure(e,'P'),section_figure(e,'B'),reinforcement_plan_figure(e),elevation_figure(e),hoop_figure(e)]):
        if i==4:parts.append(hoop_explanation_html(e))
        parts.append(fig.to_html(full_html=False,include_plotlyjs=(i==0)))
    if not actual_transverse(case):
        parts.append(cage.to_html(full_html=False,include_plotlyjs=False))
    if end_grid_enabled(case):
        parts.append(end_grid_summary(e))
        for end in ('left','right'):
            parts.append(end_figure(e,end).to_html(full_html=False,include_plotlyjs=False))
    parts.extend([clear_spacing_html(e),'</body></html>'])
    (path/'reinforcement_detail.html').write_text('\n'.join(parts),encoding='utf-8')
    from .geometry_dimensions import dimensions_figure,dimensions_html
    dimensions=dimensions_figure(e).to_html(full_html=True,include_plotlyjs=True)
    dimensions=dimensions.replace('<body>','<body>'+dimensions_html(e),1)
    (path/'geometry_dimensions.html').write_text(dimensions,encoding='utf-8')
    prefix=('UNIFORM-CAGE REFERENCE ONLY. Actual hoop/U runs are in selected_case.json and transverse_bar_schedule.csv; these scalar inputs do not represent their topology.\n\n' if actual_transverse(case) else '')
    if case.get('added_bar_layout'):
        prefix+='ADDED-BAR LAYOUT: independent row spacing/offsets are saved in selected_case.json and longitudinal_bar_positions.csv, not represented by these legacy scalar spacing inputs. Recheck spacing in the notebook.\n\n'
    (path/'blockpad_inputs.txt').write_text(prefix+'\n'.join(input_formula(k,v) for k,v in case['inputs'].items())+'\n',encoding='utf-8')
    (path/'formula_trace.json').write_text(json.dumps(formula_trace(e),indent=2,ensure_ascii=False),encoding='utf-8')
    if e.lrfd:
        (path/'lrfd_calculations.json').write_text(json.dumps(e.lrfd,indent=2,ensure_ascii=False,allow_nan=False),encoding='utf-8')
        for key in ('regions','longitudinal','intervals','inventory','transverse_development','faces'):
            rows=e.lrfd[key]
            with (path/('lrfd_'+key+'.csv')).open('w',newline='',encoding='utf-8-sig') as f:
                if rows:
                    writer=csv.DictWriter(f,fieldnames=list(dict.fromkeys(k for r in rows for k in r)))
                    writer.writeheader();writer.writerows(rows)
    if case['analysis'].get('xml_audit',{}).get('end_records'):
        from .force_diagrams import cap_force_figure,diagram_notice
        import html
        fig=cap_force_figure(case,show_resistance=True,evaluation=e)
        page=fig.to_html(full_html=True,include_plotlyjs=True)
        page=page.replace('<body>','<body><p style="font:14px Arial;margin:20px">'+html.escape(diagram_notice(case))+'</p>'
            '<p style="font:14px Arial;margin:20px">'+html.escape(fig.layout.meta['resistance_notice'])+'</p>',1)
        (path/'force_diagrams.html').write_text(page,encoding='utf-8')
    audit=strength_audit(case)
    for filename,headers,rows in [('strength_loads.csv',audit['headers'],audit['rows']),
                                  ('strength_governing.csv',audit['provenance_headers'],audit['provenance'])]:
        with (path/filename).open('w',newline='',encoding='utf-8-sig') as f:
            w=csv.writer(f);w.writerow(headers);w.writerows(rows)
    manifest={'status':e.status,'cage_issues':e.issues,'stale_geometry':e.stale,'max_dc':e.max_dc,'estimated_gross_steel_lb':e.weight_lb,
        'limitations':'Sectional checks only. Service III/fatigue readiness, D-regions, hook development/cutoffs, end anchorage, pile-head hoops and full code/detail review remain explicit. Steel includes drawn span-hook bends/tails; excludes laps, end anchorage, hoop bends and waste. Hoop quantity uses tighter spacing over the full cap; stationing is unresolved.',
        'case_sha256':hashlib.sha256((path/'selected_case.json').read_bytes()).hexdigest()}
    if e.lrfd:
        manifest['limitations']='Shared sectional LRFD working is in lrfd_calculations.json and lrfd_*.csv. Unknown load paths get no direct-loading exception. Numerical capacity remains pending where anchorage or source data is unresolved; consult the check register. Weight excludes hoop closure extensions, laps and waste. U-bars receive no closed-path torsion credit. End zones, cutoffs, D-regions and full 3D congestion retain explicit review statuses.'
        if not actual_transverse(case):manifest['limitations']+=' Uniform placement is conditional: cover plus half the bar diameter, at the larger G/L pitch; enter an actual schedule before acceptance.'
    if search_result:
        write_case(search_result.base_case,path/'search_base_case.json')
        manifest['search']={k:getattr(search_result,k) for k in ('config','total','evaluated','passed','elapsed','exhaustive','rejection_counts','force_mode')}
        manifest['search']['base_case_file']='search_base_case.json'
        manifest['search']['basis']='Completed search inputs; candidate checks do not describe manual edits to selected_case.json.'
        if search_result.force_mode=='fixed':
            manifest['search']['basis']+=' Width/depth trial with current forces held unchanged. Self-weight and stiffness changes were not reanalyzed; original analysis geometry is retained.'
            g=case['analysis']['geometry'];p=case['inputs']
            (path/'trial_geometry.txt').write_text(
                'CAP WIDTH / DEPTH TRIAL — FORCES UNCHANGED\n'
                f'Force source: {case["analysis"]["id"]}\n'
                f'Analyzed section: {g["b"]:g} x {g["h"]:g} in\n'
                f'Trial section: {p["b"]:g} x {p["h"]:g} in\n'
                'The steel search uses the current load envelopes for the trial cap size.\n'
                'Self-weight and stiffness effects have not been reanalyzed.\n'
                'Passing sectional checks do not make this a new FBMP analysis.\n',encoding='utf-8')
        manifest['search']['filter']={**options,'matching_count':len(matching)}
        mode_label='Fixed forces — width/depth trial' if search_result.force_mode=='fixed' else 'Matching analyzed geometry'
        with (path/'alternatives.csv').open('w',newline='',encoding='utf-8-sig') as f:
            w=csv.writer(f);w.writerow(['Candidate ID','Layout','Estimated gross steel (lb)','All-check D/C','Strength D/C','All-check governing check','Strength governing check','Complexity score','Force basis'])
            for i,c in enumerate(search_result.candidates,1):w.writerow([i,c.label,c.weight_lb,c.max_dc,c.strength_dc,c.governing_check,c.strength_governing_check,c.complexity,mode_label])
        with (path/'filtered_alternatives.csv').open('w',newline='',encoding='utf-8-sig') as f:
            w=csv.writer(f);w.writerow(['Filtered rank','Candidate ID','Layout','Estimated gross steel (lb)','Filter D/C','All-check D/C','Strength D/C','Controls filter D/C','D/C scope','D/C target','Force basis'])
            for rank,i in enumerate(matching,1):
                c=search_result.candidates[i]
                w.writerow([rank,i+1,c.label,c.weight_lb,candidate_dc(c,options['scope']),c.max_dc,c.strength_dc,candidate_governing(c,options['scope']),options['scope'],options['max_dc'],mode_label])
    (path/'review.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False),encoding='utf-8')
    from .calculation_report import write_calculation_report
    write_calculation_report(case,path/'calculation_report.html',evaluation=e,
                             search_result=search_result,search_filter=search_filter)
    return path

def _blockpad_template(source):
    """Read a local path or uploaded bytes without resolving external XML data."""
    from lxml import etree as E
    parser=E.XMLParser(resolve_entities=False,no_network=True)
    if isinstance(source,(bytes,bytearray,memoryview)):
        tree=E.parse(io.BytesIO(bytes(source)),parser)
    else:
        source=Path(source)
        if not source.is_file():
            raise FileNotFoundError('This notebook cannot find the source journal. In Colab, use Upload Blockpad journal; a Windows path is not accessible from the Colab runtime.')
        tree=E.parse(str(source),parser)
    reports=tree.findall("report[@name='PierCapDesign']")
    if len(reports)!=1:raise ValueError('This exporter requires exactly one C005 Live Design report named PierCapDesign.')
    report=reports[0]
    found={}
    for el in report.iter('dynexp'):
        f=el.get('formula','')
        if ' = ' in f:found[f.split(' = ')[0]]=el
    added={'Bar_P','Bar_B','Ready_pile','Pile_embed','C_pile'}
    required=set(INPUTS)-added
    required|={'Bar_pos'} if 'Bar_P' not in found else {'Bar_P','Bar_B'}
    missing=(required|{'Status_layout','Status_section'})-set(found)
    if missing:raise ValueError('C005 input definitions missing: '+', '.join(sorted(missing)))
    return tree,report,found

def validate_blockpad_source(source):
    """Confirm that a journal contains the required C005 module before export."""
    _blockpad_template(source)

def export_blockpad(source,case,destination):
    """Patch only C005 in a new copy, accepting local paths or uploaded bytes."""
    from lxml import etree as E
    case=upgrade_case(case);validate_case(case);e=evaluate(case);destination=Path(destination)
    from .end_grid import enabled as end_grid_enabled
    if end_grid_enabled(case):
        raise ValueError('The Blockpad scalar template cannot represent end-face U grids. Export the case JSON and full review bundle to retain the complete detail.')
    if case.get('transverse_detail',{}).get('enabled'):
        raise ValueError('The Blockpad scalar template cannot represent actual hoop/U stations and open-bottom topology. Export the case JSON and full review bundle to retain the complete detail.')
    same_source=not isinstance(source,(bytes,bytearray,memoryview)) and Path(source).resolve()==destination.resolve()
    if destination.exists() or same_source:raise ValueError('Choose a new output filename; originals are never overwritten.')
    tree,report,found=_blockpad_template(source);root=tree.getroot()
    other=[E.tostring(n) for n in root if n is not report]
    def replace(name,formula):
        el=found[name];el.set('formula',formula)
        for child in list(el):el.remove(child)
        E.SubElement(el,'expbody').text=formula
    # Upgrade only C005 to the current equations. Older journals have one
    # Bar_pos input; exported positive-region sizes must never be collapsed.
    extra=E.Element('table',name='NotebookRegionalDefinitions',capture='False',colwidths='1:310.00, 2:600.00')
    for definition in DEFINITIONS:
        name=definition['name']
        formula=input_formula(name,case['inputs'][name]) if definition['input'] else definition['formula']
        if name in found:
            replace(name,formula)
        else:
            row=E.SubElement(extra,'row')
            label=E.SubElement(row,'textcell',capture='False')
            E.SubElement(label,'paragraph',fontfamily='Times New Roman',fontsize='10').text=definition['caption'] or name
            cell=E.SubElement(row,'textcell',capture='False')
            paragraph=E.SubElement(cell,'paragraph',fontfamily='Times New Roman',fontsize='10')
            exp=E.SubElement(paragraph,'dynexp',formula=formula)
            E.SubElement(exp,'expbody').text=formula
            found[name]=exp
    if len(extra):report.insert(0,extra)
    from .blockpad_regions import update_regional_views
    update_regional_views(report,e)
    g=case['analysis']['geometry']
    def source_test(names):
        terms=[]
        for n in names:
            unit=INPUTS[n]['unit'];u=(' '+unit) if unit else ''
            terms.append(f'Abs({n}-{g[n]:.12g}{u}) < 0.000001{u}')
        return 'And('+','.join(terms)+')'
    replace('Status_layout','Status_layout = If('+source_test(('N_pile','S_pile','D_pile','E_clear','E_detail'))+',"SOURCE PILE LAYOUT","REIMPORT ANALYSIS FORCES")')
    replace('Status_section','Status_section = If('+source_test(('b','h'))+',"SOURCE SECTION","RECHECK MODEL SELF-WEIGHT / FORCES")')
    # Replace our prior audit when re-exporting a review copy, without changing
    # any other project report or the C005 calculation equations.
    for old in list(report):
        if old.get('name') in ('NotebookStrengthAudit','NotebookStrengthControllers','NotebookStrengthBasis'):
            report.remove(old)
    audit=strength_audit(case)
    def audit_table(name,headers,rows,widths):
        table=E.Element('table',name=name,capture='False',colwidths=', '.join(f'{i}:{w:.2f}' for i,w in enumerate(widths,1)))
        last=chr(ord('A')+len(headers)-1)
        E.SubElement(table,'stylerule',range=f'A1:{last}{len(rows)+1}',fontfamily='Times New Roman',fontsize='9',
                     border='Border("AllSides", "Solid", 0.5, "#7189A6"); Border("BetweenRows", "Solid", 0.4, "#7189A6");')
        E.SubElement(table,'stylerule',range=f'A1:{last}1',background='#DCE7F3',fontweight='bold')
        for ri,values in enumerate([headers]+rows):
            row=E.SubElement(table,'row')
            for ci,value in enumerate(values):
                cell=E.SubElement(row,'textcell',name=f'{name}R{ri}C{ci}',capture='False')
                E.SubElement(cell,'paragraph',fontfamily='Times New Roman',fontsize='9',spacingafter='3',spacingbefore='3').text=str(value)
        return table
    basis=E.Element('paragraph',name='NotebookStrengthBasis',fontfamily='Times New Roman',fontsize='10')
    basis.text='STRENGTH LOAD AUDIT — snapshot at export. '+ ' '.join(audit['notes'])+' If loads are edited in Blockpad, this source audit remains the import snapshot.'
    report.insert(0,basis)
    report.insert(1,audit_table('NotebookStrengthAudit',audit['headers'],audit['rows'],[180,100,100,100,100,100]))
    report.insert(2,audit_table('NotebookStrengthControllers',audit['provenance_headers'],audit['provenance'],[150,70,460]))
    # Clarify the original input-table heading where present.
    for paragraph in report.iter('paragraph'):
        if paragraph.text and paragraph.text.strip() == 'Strength':
            paragraph.text='Combined strength envelope'
    note=E.Element('paragraph',fontfamily='Times New Roman',fontsize='11',background='#FFF5DE')
    note.text=f'NOTEBOOK REVIEW COPY — {case["name"]}. Analysis: {case["analysis"]["id"]}. {e.status}. C005 inputs and equations match this notebook revision; other project sections are unchanged. Recalculate in Blockpad. Cage drawings are labeled geometry snapshots; re-export after any input changes. Pile collision, minimum spacing, hook fit and cage-fit screening are notebook checks: rerun/export after changing reinforcement or pile-head dimensions. Current cage issues: {"; ".join(e.issues) or "none in the notebook screen"}. Earlier narrative examples/source notes may describe the original model; reconcile them before finalizing.'
    report.insert(0,note)
    assert other==[E.tostring(n) for n in root if n is not report]
    destination.parent.mkdir(parents=True,exist_ok=True)
    with destination.open('xb') as f:tree.write(f,encoding='utf-8',xml_declaration=True)
    return destination

def import_workbooks(moment,shear,torsion,base=None,analysis_id=None,low_interval_labels=None):
    """Strict reader for the supplied Max_PierCap_*_Design workbook family.

    Independent envelopes are retained, not represented as concurrent actions.
    Rejects unknown headers, missing caches, unequal pile spacing and mismatched
    station lists. Geometry b/h/nominal pile width remains user-declared.
    """
    import openpyxl
    case=upgrade_case(base or default_case());audit={};datasets={};coordinate_lists=[]
    for kind,source in [('Moment',moment),('Shear',shear),('Torsion',torsion)]:
        path=Path(source);wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
        try:
            name='Max_PierCap_'+kind+'_Design'
            if name not in wb.sheetnames or 'Coords' not in wb.sheetnames:raise ValueError(f'{path.name}: unsupported workbook layout.')
            rows=list(wb[name].values);header=None
            for i,row in enumerate(rows):
                if len(row)>6 and str(row[1]).strip()=='LOCATION' and str(row[6]).strip()=='AASHTO-LRFD':header=i;break
            if header is None:raise ValueError(f'{path.name}: missing expected headers.')
            expected={10:'M33'} if kind=='Moment' else {9:'F22',13:'TORSION'}
            if any(str(rows[header][i]).strip()!=v for i,v in expected.items()):raise ValueError(f'{path.name}: force-axis headers differ.')
            records=[];locations=[]
            for index,row in enumerate(rows[header+2:],header+3):
                label=str(row[1] or '').strip();state=str(row[6] or '').strip().upper()
                if not re.match(r'^(Pil|Brg)\s+\d',label) or not (state.startswith('STRENGTH') or state.startswith('SERVICE')):continue
                axis=10 if kind=='Moment' else 9 if kind=='Shear' else 13
                value=row[axis]
                if not isinstance(value,(int,float)):raise ValueError(f'{path.name}: row {index} has no cached numeric force; recalculate and save in Excel.')
                records.append({'location':label,'state':state,'value':float(value),'row':index,'combination':row[5]})
                if label not in locations:locations.append(label)
            datasets[kind]=records
            audit[kind]={'filename':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheet':name,'records':len(records)}
            # Moment Coords has one coordinate per station in the first block.
            coords=[];started=False
            for row in wb['Coords'].values:
                v=row[0]
                if isinstance(v,(int,float)):coords.append(float(v));started=True
                elif started:break
            coordinate_lists.append(coords)
            if kind=='Moment':moment_locations=locations
        finally:wb.close()
    # Shear/torsion contain separate L/R rows at a station. Compare unique coordinates.
    reference=sorted(set(coordinate_lists[0]))
    for coords in coordinate_lists[1:]:
        if sorted(set(coords))!=reference:raise ValueError('Workbook coordinate sets differ; use all three exports from the same analysis run.')
    if len(moment_locations)!=len(coordinate_lists[0]):raise ValueError('Cannot map moment stations to Coords unambiguously.')
    pairs=list(zip(moment_locations,coordinate_lists[0]));centers=[]
    for label,x in pairs:
        if re.fullmatch(r'Pil\s+\d+\s+CL',label):centers.append(x)
    if len(centers)<2:raise ValueError('At least two labeled pile centerlines are needed.')
    gaps=[b-a for a,b in zip(centers,centers[1:])]
    if min(gaps)<=0 or max(gaps)-min(gaps)>.01:raise ValueError('Unequal pile spacing requires an expanded geometry model.')
    p=case['inputs'];p['N_pile']=len(centers);p['S_pile']=sum(gaps)/len(gaps)
    def envelope(kind,predicate,key,absolute=True):
        group=[r for r in datasets[kind] if predicate(r)]
        if not group:raise ValueError(f'Missing {key} envelope in {kind} workbook.')
        winner=max(group,key=lambda r:abs(r['value']) if absolute else r['value'])
        p[key]=abs(winner['value']) if absolute else winner['value'];audit[key]=winner
    for prefix,state in [('Mu','STRENGTH'),('MI','SERVICE-I')]:
        def state_ok(r):return r['state'].startswith(state) if state=='STRENGTH' else r['state']=='SERVICE-I'
        envelope('Moment',lambda r:state_ok(r) and r['value']<0,prefix+'_N')
        envelope('Moment',lambda r:state_ok(r) and r['location'].startswith('Pil') and r['value']>=0,prefix+'_P')
        envelope('Moment',lambda r:state_ok(r) and r['location'].startswith('Brg') and r['value']>=0,prefix+'_B')
    envelope('Shear',lambda r:r['state'].startswith('STRENGTH'),'Vu_G')
    envelope('Torsion',lambda r:r['state'].startswith('STRENGTH'),'Tu')
    if low_interval_labels:
        labels=set(low_interval_labels)
        if not labels.issubset({r['location'] for r in datasets['Shear']}):raise ValueError('Low interval labels were not found in this export.')
        envelope('Shear',lambda r:r['state'].startswith('STRENGTH') and r['location'] in labels,'Vu_L')
    else:p['Vu_L']=p['Vu_G']
    for k in ('Ready_III','Ready_fatigue'):p[k]=False
    for k in p:
        if k.startswith(('MIII_','MDL_','DMLL_')):p[k]=0
    case['analysis']={'id':analysis_id or Path(moment).stem,'geometry':{k:p[k] for k in GEOMETRY},'workbook_audit':audit,
        'notes':'b/h, nominal pile width and end allowances are user-declared. Workbook coordinate matching cannot prove files share a run; confirm analysis ID. Independent design envelopes, not concurrent actions. Low interval defaults to global shear unless explicit labels are supplied.'}
    case['name']=case['analysis']['id'];validate_case(case);evaluate(case)
    case.pop('section_study',None)  # These are newly imported analysis forces.
    return case
