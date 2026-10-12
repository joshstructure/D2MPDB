"""Geotechnical Excel report with summary, soil properties, and minimum-tip trials."""
from pathlib import Path
import math
import re
import textwrap

from .output_labels import table_reference
from .soil_properties import soil_report_rows


def handoff_filename(trial_source=''):
    """Use the study label as a portable filename, never as a directory path."""
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '-', trial_source.strip()).strip(' .')
    suffix = Path(name).suffix
    if suffix.lower() in ('.csv', '.tsv', '.txt', '.xls', '.xlsx', '.xlsm'):
        name = name[:-len(suffix)]
    name = name[:120].rstrip(' .') or 'geotechnical_report'
    if re.fullmatch(r'CON|PRN|AUX|NUL|COM[1-9¹²³]|LPT[1-9¹²³]', name.split('.')[0].rstrip(), re.I):
        name = '_'+name
    return name+'.xlsx'


def handoff_tables(data, result, ground):
    """Share the exact final rows between the notebook preview and Excel export."""
    tip = result.get('tip_elevation_ft') if result and data['selected'] else None
    # Required embedment in the calculation is before optional tip rounding.
    # The handoff length must reach the adopted (possibly rounded) elevation.
    embedment = ground - tip if ground is not None and tip is not None else None
    elevations = data.get('elevations', {})
    section = list(data['section'])
    section.extend([
        dict(item='Soil-top elevation (ft)', value=elevations.get('soil_top_ft')),
        dict(item='Pile-head elevation (ft)', value=elevations.get('pile_head_ft')),
        dict(item='Total pile length at minimum tip (ft)', value=result.get('total_length_ft') if tip is not None else None),
    ])
    # Show adopted datums too if they differ from the physical XML elevations.
    for label, value, original in [
        ('Design ground / scour elevation (ft)', ground, elevations.get('soil_top_ft')),
        ('Adopted pile cutoff elevation (ft)', elevations.get('cutoff_ft'), elevations.get('pile_head_ft')),
    ]:
        if value is not None and (original is None or not math.isclose(value, original, rel_tol=0, abs_tol=1e-6)):
            section.append(dict(item=label, value=value))
    return [
        dict(reference='handoff_section', columns=[('item', 'Pile information'), ('value', 'Value')],
             rows=section),
        dict(reference='handoff_tip', columns=[('criterion', 'Controlling criterion'),
             ('embedment_ft', 'Final embedment below design ground / scour (ft)'),
             ('tip_elevation_ft', 'Final tip elevation (ft)')],
             rows=[dict(criterion=result['controlling_criterion'] if tip is not None else 'Not calculated',
                        embedment_ft=embedment, tip_elevation_ft=tip)]),
        dict(reference='handoff_loads', columns=[('load', 'Pile-head design load'), ('fdl_tons', 'Final FDL (short tons)')],
             rows=[dict(load=r['load'], fdl_tons=r['short_tons']) for r in data['loads']]),
    ]


def handoff_qualification(data, result):
    """Keep an incomplete calculation from being presented as a final design."""
    if not data['selected'] or not result or result.get('tip_elevation_ft') is None:
        return 'Minimum tip not calculated.'
    if not result['comparison_complete']:
        return 'Minimum-tip comparison incomplete; reported values use available criteria only.'
    return ''


def write_handoff_xlsx(path, data, result, ground, notes='', *, trial_source=''):
    """Write a standalone, values-only workbook using the notebook's Excel dependency."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.page import PageMargins

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Pile summary'
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_properties.outlinePr.summaryRight = False
    sheet.sheet_properties.tabColor = '213649'
    for col, width in [('A', 43), ('B', 40), ('C', 30)]:
        sheet.column_dimensions[col].width = width

    navy, ink, pale, line = '213649', '243746', 'EAF1F6', 'DDE3E8'

    def cell(row, col, value, *, bold=False, fill=None, color=ink, size=11, number_format=None):
        target = sheet.cell(row, col, 'Unavailable' if value is None else value)
        # User-entered section descriptions and notes are text, never Excel formulas.
        if isinstance(target.value, str):
            target.data_type = 's'
        target.font = Font(name='Calibri', size=size, bold=bold, color=color)
        target.alignment = Alignment(horizontal='right' if isinstance(value, (int, float)) else 'left',
                                     vertical='center', wrap_text=True,
                                     indent=1 if isinstance(value, (int, float)) else 0)
        if fill:
            target.fill = PatternFill('solid', fgColor=fill)
        if number_format:
            target.number_format = number_format
        return target

    def wide_row(row, value, end=3, **style):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end)
        cell(row, 1, value, **style)

    title = trial_source.strip() or 'Geotechnical report'
    workbook.properties.title = title
    wide_row(1, title, bold=True, fill=navy, color='FFFFFF', size=20)
    title_lines = sum(max(1, len(textwrap.wrap(line, width=60))) for line in title.splitlines())
    sheet.row_dimensions[1].height = min(409, max(40, 26*title_lines+12))
    row = 3
    qualification = handoff_qualification(data, result)
    for table in handoff_tables(data, result, ground):
        wide_row(row, table_reference(table['reference']), bold=True, size=13)
        sheet.row_dimensions[row].height = 28
        row += 1
        for col, (_, label) in enumerate(table['columns'], 1):
            cell(row, col, label, bold=True, fill=pale)
        if len(table['columns']) == 2:
            sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        sheet.row_dimensions[row].height = 42 if len(table['columns']) == 3 else 27
        row += 1
        for index, values in enumerate(table['rows']):
            for col, (key, _) in enumerate(table['columns'], 1):
                value = values[key]
                fmt = ('0' if key == 'value' and isinstance(value, int) else '#,##0.000;-#,##0.000;0.000')
                target = cell(row, col, value, fill='F5F8FA' if index % 2 else 'FFFFFF',
                              number_format=fmt if isinstance(value, (int, float)) else None)
                target.border = Border(bottom=Side(style='hair', color=line))
            if len(table['columns']) == 2:
                sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
            # Criterion ties and section descriptions need room at normal zoom.
            longest = max(len(str(v or '')) for v in values.values())
            sheet.row_dimensions[row].height = max(29, 15 * ((longest + 39) // 40) + 10)
            row += 1
        if table['reference'] == 'handoff_tip' and qualification:
            wide_row(row, qualification, color='805600', fill='FFF4D6')
            sheet.row_dimensions[row].height = 31
            row += 1
        if table['reference'] == 'handoff_section' and result and result.get('total_length_ft') is not None:
            cutoff = data.get('elevations', {}).get('cutoff_ft')
            tip = result.get('tip_elevation_ft')
            if cutoff is not None and tip is not None and not math.isclose(result['total_length_ft'], cutoff-tip, abs_tol=1e-8):
                wide_row(row, 'Total pile length is rounded up to whole feet.', size=10)
                sheet.row_dimensions[row].height = 24
                row += 1
        if table['reference'] == 'handoff_loads':
            wide_row(row, 'FDL = factored design load. Compression and uplift are positive magnitudes. 1 short ton = 2 kip.', size=10)
            sheet.row_dimensions[row].height = 28
            row += 1
        row += 1

    if notes.strip():
        wide_row(row, 'Project notes', bold=True, size=13)
        sheet.row_dimensions[row].height = 28
        row += 1
        # Split long notes into ordinary rows so printing can continue on the next page.
        for paragraph in notes.strip().splitlines():
            for part in textwrap.wrap(paragraph, width=110) or ['']:
                wide_row(row, part)
                sheet.row_dimensions[row].height = 30
                row += 1

    from openpyxl.utils import get_column_letter

    def finish_sheet(last_row, columns):
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.sheet_properties.tabColor = navy
        sheet.print_options.horizontalCentered = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins = PageMargins(left=.3, right=.3, top=.35, bottom=.35, header=.15, footer=.15)
        sheet.print_area = f'A1:{get_column_letter(columns)}{last_row}'
        sheet.oddFooter.center.text = 'Page &P of &N'

    finish_sheet(row-1, 3)

    def detail_sheet(name, columns, rows, empty_message):
        """One filterable table, with fixed identifying columns and units in its headers."""
        nonlocal sheet
        sheet = workbook.create_sheet(name)
        count = len(columns)
        wide_row(1, title, end=count, bold=True, fill=navy, color='FFFFFF', size=20)
        sheet.row_dimensions[1].height = min(409, max(40, 26*title_lines+12))
        wide_row(2, name, end=count, bold=True, size=13)
        sheet.row_dimensions[2].height = 28
        for col, (_, label, width) in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(col)].width = width
            cell(4, col, label, bold=True, fill=pale)
        sheet.row_dimensions[4].height = 42
        for index, values in enumerate(rows):
            row = index+5
            line_count = 1
            for col, (key, _, width) in enumerate(columns, 1):
                value = values.get(key)
                value = 'Yes' if value is True else 'No' if value is False else value
                value = '—' if value is None or value == '' else value
                target = cell(row, col, value, fill='F5F8FA' if index % 2 else 'FFFFFF',
                              number_format=('0' if value == int(value) else '0.000E+00' if abs(value)<1e-5 else '0.#####;-0.#####;0')
                              if isinstance(value, (int, float)) else None)
                if value in ('—', 'Yes', 'No'):
                    target.alignment = Alignment(horizontal='center', vertical='center')
                target.border = Border(bottom=Side(style='hair', color=line))
                line_count = max(line_count, sum(max(1, len(textwrap.wrap(t, width=max(8, int(width)-3))))
                                                  for t in str(value).splitlines()))
            sheet.row_dimensions[row].height = max(25, 15*line_count+8)
        last = 4+len(rows)
        if not rows:
            wide_row(5, empty_message, end=count, color='805600', fill='FFF4D6')
            sheet.row_dimensions[5].height = 30
            last = 5
        else:
            sheet.auto_filter.ref = f'A4:{get_column_letter(count)}{last}'
        sheet.freeze_panes = 'D5'
        sheet.print_title_rows = '1:4'
        finish_sheet(last, count)

    detail_sheet('Soil properties', [
        ('category', 'Property group', 25), ('soil_set', 'Soil set', 9), ('layer', 'Layer', 9),
        ('property', 'Property', 43), ('units', 'Units', 13), ('top', 'Layer top', 15),
        ('bottom', 'Layer bottom', 15), ('value', 'Layer / tip value', 18),
    ], soil_report_rows(data.get('soil_properties', [])), 'Soil properties are unavailable for this study.')

    trials = data.get('trials', [])
    columns = [('trial', 'Trial', 10), ('combination', 'Load combination', 16), ('pile', 'Pile', 10),
               ('embedment_ft', 'Embedment (ft)', 20), ('displacement_in', 'Displacement (in)', 22)]
    if any(r.get('series', 'Envelope') != 'Envelope' for r in trials):
        columns.insert(0, ('series', 'Series', 22))
    for key, label, width in [('dc', 'Model D/C', 14), ('dx_in', 'DX (in)', 17), ('dy_in', 'DY (in)', 17),
                              ('converged', 'Converged', 15), ('iterations', 'Iterations', 13),
                              ('tolerance_kip', 'Solver tolerance (kip)', 22)]:
        if any(r.get(key) is not None for r in trials):
            columns.append((key, label, width))
    columns.extend([('delta_in', 'Change to next shallower trial (in)', 26),
                    ('critical_trial', 'Trial Lcrit', 14)])
    detail_sheet('Minimum tip trials', columns, trials, 'No minimum-tip trial data are available for this study.')

    path = Path(path)
    workbook.save(path)
    return path
