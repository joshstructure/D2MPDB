"""Verify a recalculated workbook: no formula errors, QC passes, values match the live notebook."""
from __future__ import annotations

from pathlib import Path

import numpy as np
import openpyxl

from . import sources, workbook

ERRORS = ('#REF!', '#VALUE!', '#DIV/0!', '#NAME?', '#N/A', '#NUM!', '#NULL!')


def named_value(wb, name):
    dn = wb.defined_names[name]
    (sheet, ref), = list(dn.destinations)
    return wb[sheet][ref.replace('$', '')].value


def formula_errors(wb) -> list[str]:
    found = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value in ERRORS:
                    found.append(f'{ws.title}!{cell.coordinate} {cell.value}')
    return found


def qc_rows(wb):
    ws = wb['QC']
    rows = []
    for r in range(1, ws.max_row + 1):
        status = ws.cell(row=r, column=7).value
        if status in ('PASS', 'FAIL', 'CHECK'):
            rows.append((ws.cell(row=r, column=2).value, ws.cell(row=r, column=3).value, ws.cell(row=r, column=4).value, status))
    return rows


def verify(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    lines, ok = [], True
    errors = formula_errors(wb)
    if errors:
        ok = False
        lines.append(f'FORMULA ERRORS ({len(errors)}): ' + '; '.join(errors[:20]))
    if named_value(wb, 'Build_Fingerprint') != workbook.fingerprint(workbook.source_hashes()):
        ok = False
        lines.append('Fingerprint does not match the current sources/generator.')
    rows = qc_rows(wb)
    fails = [r for r in rows if r[3] == 'FAIL']
    checks = [r for r in rows if r[3] == 'CHECK']
    lines.append(f'QC: {len(rows) - len(fails) - len(checks)} pass, {len(fails)} fail, {len(checks)} check; overall {named_value(wb, "QC_Overall")}')
    for label, got, expected, status in fails[:30]:
        lines.append(f'  FAIL {label}: workbook {got} vs notebook {expected}')
    for label, *_ in checks:
        lines.append(f'  CHECK {label}')
    if fails or named_value(wb, 'QC_Overall') != 'PASS':
        ok = False
    # Independent live comparison (not the build-time snapshot).
    cfg, model, _ = sources.load_notebook_model()
    live = {
        'Sup_Drop': model.support_drop, 'Pier_Top': model.pier['top'], 'EB1_Top': model.end_bents[0]['top'],
        'EB2_Top': model.end_bents[1]['top'], 'Haunch_Min': model.haunch_min * 12,
    }
    if np.isfinite(model.clearance_min):
        live['Clr_Min'] = model.clearance_min
    for name, value in live.items():
        got = named_value(wb, name)
        if got is None or abs(float(got) - float(value)) > 1e-8:
            ok = False
            lines.append(f'LIVE MISMATCH {name}: workbook {got} vs notebook {value}')
    lines.append('Verification ' + ('passed.' if ok else 'FAILED.'))
    return dict(ok=ok, lines=lines, errors=errors, fails=fails)
