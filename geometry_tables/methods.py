"""Method tabs: math notation, explanation, schematics and live charts for every workbook tab.

Pictures here are equations and label-only schematics (figures.py). Every figure
that shows project geometry or numbers is a live Excel chart (charts.py), and
every sentence quoting a current value is an Excel formula, so the Method tabs
recalculate with the workbook. A banner on each tab reports whether any input
differs from the notebook build.
"""
from __future__ import annotations

import io
import math

from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill

from . import charts, figures
from .charts import NUM, STA, TXT, live

FONT = 'Arial'
COLS = 12               # B..M
COL_WIDTH = 11.5
TEXT_CHARS = 138        # characters per merged line across B..M
ROW_PX = 20
F_P = Font(name=FONT, size=10)
F_H = Font(name=FONT, size=12, bold=True, color='1F2A36')
F_T = Font(name=FONT, size=16, bold=True, color='1F2A36')
F_CAP = Font(name=FONT, size=9, italic=True, color='5A6570')
F_NOTE = Font(name=FONT, size=10, color='9C5700')
F_VH = Font(name=FONT, size=9, bold=True)
F_BANNER = Font(name=FONT, size=9, bold=True, color='1F2A36')
FILL_H = PatternFill('solid', fgColor='EEF1F4')
FILL_NOTE = PatternFill('solid', fgColor='FFF4E5')
FILL_VH = PatternFill('solid', fgColor='D9E1F2')
FILL_LIVE = PatternFill(fill_type='solid', start_color='C6EFCE', end_color='C6EFCE')     # conditional fills need bgColor
FILL_EDITED = PatternFill(fill_type='solid', start_color='FFEB9C', end_color='FFEB9C')
WRAP = Alignment(wrap_text=True, vertical='top')

BANNER = ('=IF(Inputs_Edited=0,'
          + TXT('LIVE  Charts, tables and the numbers quoted in this text are workbook formulas. '
                'Inputs match the notebook build (model ') + '&Build_Model&' + TXT(', built ') + '&Build_Time&' + TXT(').')
          + ',Inputs_Edited&' + TXT(' INPUT(S) EDITED SINCE THE NOTEBOOK BUILD  The charts and tables follow the edited inputs, '
                                    'which no longer match the geometry notebook (QC ') + '&QC_Overall&'
          + TXT('). Change the notebook CONFIG and rebuild.') + ')')


class Renderer:
    def __init__(self, book, ws):
        self.book, self.ws, self.row = book, ws, 1
        self.area = None
        ws.sheet_view.showGridLines = False
        ws.column_dimensions['A'].width = 2
        for k in range(2, 2 + COLS):
            ws.column_dimensions[chr(64 + k)].width = COL_WIDTH

    def data(self):
        """Figure-data area for this tab's charts (created on first use)."""
        if self.area is None:
            self.area = charts.DataArea(self.ws)
        return self.area

    def _merge_text(self, text, font=F_P, fill=None, chars=TEXT_CHARS, min_height=15):
        ws, r = self.ws, self.row
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=1 + COLS)
        cell = ws.cell(row=r, column=2, value=text)
        cell.font = font
        cell.alignment = WRAP
        if fill:
            for c in range(2, 2 + COLS):
                ws.cell(row=r, column=c).fill = fill
        shown = str(text)
        if shown.startswith('='):        # a live sentence: estimate its displayed length
            shown = ' ' * int(len(shown) * 0.7)
        lines = sum(max(1, math.ceil(len(part) / chars)) for part in shown.split('\n'))
        ws.row_dimensions[r].height = max(min_height, 13.2 * lines + 3)
        self.row += 1

    def title(self, text, sub):
        ws = self.ws
        ws.cell(row=1, column=2, value=text).font = F_T
        ws.cell(row=2, column=2, value=sub).font = F_CAP
        ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=1 + COLS)
        banner = ws.cell(row=3, column=2, value=BANNER)
        banner.font = F_BANNER
        banner.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[3].height = 28
        ref = f'B3:{chr(65 + COLS)}3'
        ws.conditional_formatting.add(ref, FormulaRule(formula=['Inputs_Edited=0'], fill=FILL_LIVE))
        ws.conditional_formatting.add(ref, FormulaRule(formula=['Inputs_Edited>0'], fill=FILL_EDITED))
        self.row = 4

    def h(self, text):
        self.row += 1
        ws, r = self.ws, self.row
        ws.cell(row=r, column=2, value=text).font = F_H
        for c in range(2, 2 + COLS):
            ws.cell(row=r, column=c).fill = FILL_H
        ws.row_dimensions[r].height = 18
        self.row += 1

    def p(self, text):
        self._merge_text(text)

    def bullets(self, items):
        for item in items:
            self._merge_text(item if str(item).startswith('=') else '•  ' + item)

    def note(self, text):
        self._merge_text(text, font=F_NOTE, fill=FILL_NOTE, min_height=18)

    def _reserve(self, height_px):
        rows = math.ceil((height_px + 8) / ROW_PX)
        for r in range(self.row, self.row + rows):
            self.ws.row_dimensions[r].height = 15
        self.row += rows

    def image(self, png: bytes, scale: float):
        img = XLImage(io.BytesIO(png))
        img.width, img.height = int(img.width * scale), int(img.height * scale)
        img.anchor = f'B{self.row}'
        self.ws.add_image(img)
        self._reserve(img.height)

    def eq(self, latex, text=None, size=13):
        self.image(figures.equation(latex, size), 0.5)
        if text:
            self._merge_text(text, font=F_CAP)

    def fig(self, png, caption=None):
        """A label-only schematic picture (no project values)."""
        self.image(png, 0.56)
        if caption:
            self._merge_text(caption, font=F_CAP)

    def chart(self, chart, caption=None):
        """A live Excel chart drawn from this tab's figure-data block."""
        self.ws.add_chart(chart, f'B{self.row}')
        self._reserve(chart.height / 2.54 * 96)
        if caption:
            self._merge_text(caption, font=F_CAP)

    def vars(self, rows):
        ws = self.ws
        r = self.row
        heads = [(2, 2, 'Symbol'), (3, 7, 'Meaning'), (8, 10, 'Workbook name / location'), (11, 13, 'Units / value')]
        for a, b, text in heads:
            ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
            cell = ws.cell(row=r, column=a, value=text)
            cell.font = F_VH
            for c in range(a, b + 1):
                ws.cell(row=r, column=c).fill = FILL_VH
        r += 1
        for sym, meaning, where, unit in rows:
            for (a, b, _), text in zip(heads, (sym, meaning, where, unit)):
                ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
                cell = ws.cell(row=r, column=a, value=text)
                cell.font = F_P
                cell.alignment = WRAP
            ws.row_dimensions[r].height = max(15, 13.2 * math.ceil(len(meaning) / 58) + 2)
            r += 1
        self.row = r + 1


def _range(sheet, column, rows):
    return f"'{sheet}'!${column}${min(rows)}:${column}${max(rows)}"


# --------------------------------------------------------------- content --
def inputs_method(R, ctx):
    m = ctx['model']
    R.title('Inputs - Method', 'Where every input comes from and how the workbook is kept consistent with the notebook')
    R.h('1. Authority and data flow')
    R.p('The Colab geometry notebook (Bridge_Geometry_2 Beam_V2.ipynb) is the source of truth for all bridge geometry. '
        'Its CONFIG dictionary and BridgeModel equations define the profile, deck, beam layout, bearings, caps, '
        'pedestals and piles. The FreeCAD model and these CAD tables are built from the notebook, never the reverse. '
        'The generator executes the notebook cells, writes CONFIG into this tab, and writes every calculation as an Excel formula.')
    R.fig(figures.data_flow(), 'Figure 1. Source-to-table flow. Each arrow is a formula chain you can trace with Excel "Trace Precedents".')
    R.h('2. Sign conventions and units')
    R.bullets([
        'Stations s increase from End Bent 1 to End Bent 2 (ft). Offsets y are measured from the PGL, positive to the right looking up-station (ft).',
        'Elevations z are in feet on the project datum. Inputs ending in "in" are inches; slopes and grades are ft/ft.',
        'Deflections and camber are in inches, positive upward (Mathcad convention). The camber C is the peak upward value at midspan.',
        'Workbook names (Formulas > Name Manager) mirror the notebook keys, for example Prof_Sta0 = CONFIG["profile"]["start_sta_ft"].',
    ])
    R.h('3. Mathcad beam-design values')
    R.p('Beam-design inputs are the saved results stored in the two registered Mathcad runs. The generator parses the worksheet XML '
        'and reads each result by variable name; Mathcad is not recalculated. The notebook camber input equals the Mathcad NetCamber:')
    R.eq(r'C = \max(\Delta_{120\,day}) + \min(\Delta_{DL,nc}) + \min(\Delta_{DL,c})',
         live('Span 1: C = ', NUM('Sp1_C', '0.000000'), ' in;  Span 2: C = ', NUM('Sp2_C', '0.000000'), ' in.  '
              'QC confirms the deflection sum, the Mathcad NetCamber and the notebook CONFIG camber agree.'))
    R.vars([
        ('C', 'Net camber at 120 days after all dead load', 'Sp1_C, Sp2_C (Inputs)', 'in'),
        ('Δ120', '120-day prestress + beam weight deflection (max)', 'MC1_D120, MC2_D120', 'in'),
        ('ΔDL,nc', 'Non-composite dead load (deck pour) deflection (min)', 'MC1_Dnc, MC2_Dnc', 'in'),
        ('ΔDL,c', 'Composite dead load (barrier) deflection (min)', 'MC1_Dc, MC2_Dc', 'in'),
        ('ψ, εsh', 'Creep coefficient and shrinkage strain at 120 days', 'MC1_Psi, MC1_Esh ...', '-'),
    ])
    R.h('4. Designer assignments and template constants')
    R.p('A few table entries are not geometry and are not in Mathcad: concrete class, strand pattern number, plan view case and pad type. '
        'They are explicit inputs at the bottom of the Inputs tab. Standard bar-bending constants for the Table of Beam Variables are '
        'taken from the BCC example workbook formulas so the tables are produced the same way as before.')
    R.h('5. Keeping it current')
    R.bullets([
        'Change geometry in the notebook CONFIG, then run: .venv\\Scripts\\python.exe scripts\\build_geometry_cad_tables.py',
        'The script rebuilds every tab, has Excel recalculate, and verifies the results against the notebook before saving.',
        'Every figure that shows project values is a live Excel chart of the workbook cells; its figure data (formulas) sits to the '
        'right of the text on each Method tab. Pictures are limited to equations and label-only schematics.',
        'Editing a blue input here recalculates the tables and charts. The banner at the top of every Method tab, the Cover and the QC tab '
        'then report that inputs differ from the notebook build, and QC flags any result that no longer matches the notebook.',
        f'Build fingerprint {ctx["fingerprint"]}; notebook model ID {m.model_id}. The regression test fails if any source or the generator changes without a rebuild.',
    ])


def profile_method(R, ctx):
    R.title('Profile & Deck - Method', 'Profile grade line, deck cross slope, beam offsets and haunch check points')
    R.h('1. Profile grade line (PGL)')
    R.p(live('The notebook profile mode is "', ('Prof_Mode',), '". In brokeback mode the PGL is two straight tangents of nominal grade g '
             'that meet at the tangent intersection station. The PI elevation averages the two tangents projected to the PI, '
             'which preserves both end elevations:'))
    R.eq(r'z_{PI} = \frac{\left[z_0 + g\,(s_{PI}-s_0)\right] + \left[z_1 + g\,(s_1-s_{PI})\right]}{2}')
    R.eq(r'g_1 = \frac{z_{PI}-z_0}{s_{PI}-s_0}\qquad g_2 = \frac{z_1-z_{PI}}{s_1-s_{PI}}')
    R.eq(r'z_{PGL}(s) = z_0 + g_1\,(s-s_0)\quad (s \leq s_{PI}),\qquad z_{PGL}(s) = z_{PI} + g_2\,(s-s_{PI})\quad (s > s_{PI})')
    R.p(live('Current values: g1 = ', NUM('Prof_G1*100', '+0.000000;-0.000000'), '%, g2 = ', NUM('Prof_G2*100', '+0.000000;-0.000000'),
             '%, z_PI = ', ('IF(Prof_Mode="brokeback",TEXT(Prof_PIZ,"0.000000")&" ft","n/a (parabolic)")',), ' at ', (STA('Prof_PISta'),), '.'))
    R.p('Parabolic mode (also in the notebook) keeps both end elevations and the entered grade change Δg, with x = s - s0 and L = s1 - s0:')
    R.eq(r'g_1 = \frac{z_1-z_0}{L} - \frac{\Delta g}{2},\quad g_2 = g_1+\Delta g,\quad z_{PGL}(s) = z_0 + g_1 x + k_q x^2,\quad k_q = \frac{g_2-g_1}{2L}')
    R.chart(charts.profile(R.data(), ctx), 'Figure 1. Profile along the bridge, drawn from the PGL formulas (vertical scale exaggerated by the chart shape).')
    R.h('2. Deck surface')
    R.p('The deck top at offset y adds a cross-slope term to the PGL. Crowned mode slopes both sides away from the PGL; uniform mode '
        'applies one slope across the deck. (The notebook transition mode is not implemented in this workbook; QC fails if it is selected.)')
    R.eq(r'z_{deck}(s,y) = z_{PGL}(s) + \delta(y),\qquad \delta(y) = s_c\,|y|\ \ (\mathrm{crowned}),\qquad \delta(y) = s_u\,y\ \ (\mathrm{uniform})')
    R.chart(charts.typical_section(R.data(), ctx),
            live('Figure 2. Section at Span 1 midspan (', (STA('Sp1_Mid'),), '), true scale: deck, FIB-36 beams, haunch (shaded) and the '
                 'four haunch check points on each flange (red).'))
    R.h('3. Beam offsets and haunch check points')
    R.eq(r'y_1 = -\frac{1}{2}\sum_k \mathrm{gap}_k,\qquad y_{k+1} = y_k + \mathrm{gap}_k')
    R.p('The haunch is checked at four transverse points on each top flange: both flange edges, the beam centerline, and the crown '
        'point clipped onto the flange (it captures a crown break that falls on the flange). The lowest deck point over the flange governs:')
    R.eq(r'y \in \left\{y_k-\frac{b_f}{2},\ y_k,\ y_k+\frac{b_f}{2},\ \max\left(y_k-\frac{b_f}{2},\ \min\left(0,\ y_k+\frac{b_f}{2}\right)\right)\right\},\qquad \delta_{min,k} = \min_y \delta(y)')
    R.vars([
        ('s0, s1, sPI', 'Profile start, end and tangent-intersection stations', 'Prof_Sta0, Prof_Sta1, Prof_PISta', 'ft'),
        ('z0, z1', 'Profile end elevations', 'Prof_Z0, Prof_Z1', 'ft'),
        ('g1, g2', 'Tangent grades', 'Prof_G1, Prof_G2', 'ft/ft'),
        ('kq', 'Parabolic coefficient (0 in brokeback mode)', 'Prof_Kq', '1/ft'),
        ('sc, su', 'Crown slope, uniform slope', 'Deck_Crown, Deck_Uni', 'ft/ft'),
        ('yk', 'Beam k centerline offset', 'Bm1_Off, Bm2_Off', 'ft'),
        ('δmin,k', 'Lowest cross-slope drop over beam k flange', 'Bm1_dMin, Bm2_dMin', 'ft'),
    ])


def beam_geometry_method(R, ctx):
    R.title('Beam Geometry - Method', 'Physical beam ends, bearings and the Table of Beam Variables dimensions')
    R.h('1. Physical beam ends and bearings')
    R.p('Spans run between datums: the front face of backwall (FFBW) at end bents and the pier centerline. Beams stop short of each datum '
        'by the beam-end clearance e at an end bent, or by half the pier gap g at the pier. Bearings sit a setback j inside each beam end. '
        'The deck stops a clearance cd inside each FFBW and is continuous over the pier:')
    R.eq(r's_{beam,0} = s_a + \left\{e\ \mathrm{or}\ \frac{g}{2}\right\},\quad s_{beam,1} = s_b - \left\{e\ \mathrm{or}\ \frac{g}{2}\right\},\quad s_A = s_{beam,0}+j,\quad s_B = s_{beam,1}-j')
    R.p(live('Current: e = ', NUM('Phys_EndClr', 'General'), ' in, g = ', NUM('Phys_PierGap', 'General'), ' in, j = ',
             NUM('Phys_Setback', 'General'), ' in, cd = ', NUM('Phys_DeckClr', 'General'), ' in. Span 1 bearings ', (STA('Sp1_BrgA'),),
             ' to ', (STA('Sp1_BrgB'),), '; Span 2 bearings ', (STA('Sp2_BrgA'),), ' to ', (STA('Sp2_BrgB'),), '.'))
    R.fig(figures.support_details(), 'Figure 1. Beam-end, bearing and pier-joint dimensions (elevation, not to scale). Current values are in the table below.')
    R.vars([
        ('e', 'Beam end clearance to the FFBW', 'Phys_EndClr', live(NUM('Phys_EndClr', 'General'), ' in')),
        ('J', 'Bearing setback, beam end to CL bearing (Dim J)', 'Phys_Setback', live(NUM('Phys_Setback', 'General'), ' in')),
        ('K1 = e + J', 'FFBW to CL bearing at an end bent (Dim K1 / K2)', 'Beam Geometry cols L, M', live(NUM('Phys_EndClr+Phys_Setback', 'General'), ' in')),
        ('g', 'Pier beam gap, beam end to beam end', 'Phys_PierGap', live(NUM('Phys_PierGap', 'General'), ' in')),
        ('K = g/2 + J', 'CL pier to CL bearing', 'Beam Geometry cols L, M', live(NUM('Phys_PierGap/2+Phys_Setback', 'General'), ' in')),
    ])
    R.h('2. Bearing chord and sloped length')
    R.eq(r'L_b = s_B - s_A,\qquad g_c = \frac{z_{PGL}(s_B)-z_{PGL}(s_A)}{L_b},\qquad L_s = (s_{beam,1}-s_{beam,0})\sqrt{1+g_c^2}')
    R.p('The beam soffit chord between bearings is parallel to the PGL chord (the notebook drops all bearings by the same D; see Camber & Haunch). '
        'QC compares the horizontal beam length with the Mathcad L.beam.')
    R.h('3. Table of Beam Variables dimensions (BCC template method)')
    R.eq(r'\mathrm{Dim\ J} = j,\qquad \mathrm{Dim\ K1} = 12\,(s_A - s_a),\qquad \mathrm{Dim\ K2} = 12\,(s_b - s_B)\quad[\mathrm{in}]')
    R.eq(r'\mathrm{Dim\ P} = 0\ \ \mathrm{if}\ \ |g_c\,d_b| < 1\ \mathrm{in},\ \ \mathrm{else}\ \ -g_c\,d_b;\qquad \mathrm{End\ elev.\ cond.} = 1\ (P=0),\ 2\ (P<0),\ 3\ (P>0)')
    R.p('Elastic plus time-dependent shortening R uses the template expression with the Mathcad prestress, creep and shrinkage values:')
    R.eq(r'R = \frac{(1+\psi_{120})\,P\,L_s}{E_{ci}\,A_g}\cdot 12 + |\varepsilon_{sh,120}|\,L_s\cdot 12\ \ [\mathrm{in}],\qquad P = n_s P_j + n_t P_t')
    R.eq(r'\mathrm{Dim\ L} = L_s + \frac{R}{12}\ \ (\mathrm{BCC\ template\ rule})')
    R.note('Dim L adds R to the sloped length, exactly as the BCC template did. Confirm against the current FDOT Table of Beam Variables '
           'instructions before issue; the sloped length without R is shown in the Ls column.')
    R.fig(figures.beam_end_dims(), 'Figure 2. End-of-beam dimensions used in the Table of Beam Variables.')
    R.vars([
        ('sa, sb', 'Span datum stations (FFBW or CL pier)', 'Sp1_Sta0, Sp1_Sta1 ...', 'ft'),
        ('sA, sB', 'Bearing stations', 'Sp1_BrgA, Sp1_BrgB ...', 'ft'),
        ('Lb, gc', 'Bearing span and chord grade', 'Sp1_Lb, Sp1_Gc ...', 'ft, ft/ft'),
        ('P', 'Total prestress force at jacking', 'Beam Geometry col P', 'kip'),
        ('Ag, Eci', 'FIB-36 gross area; release modulus (Mathcad)', 'Bm_Area, MC1_Eci', 'in2, ksi'),
    ])


def camber_method(R, ctx):
    R.title('Camber & Haunch - Method', 'Common support drop D, build-up dimensions and the minimum haunch')
    R.h('1. Beam soffit with camber')
    R.p('Within a span, t runs from 0 at bearing A to 1 at bearing B. The soffit follows the PGL chord between bearings, dropped by one '
        'common distance D for every support, plus the Mathcad camber as an upward parabola and any entered beam raise r:')
    R.eq(r't = \frac{s-s_A}{s_B-s_A},\qquad q(s) = z_{PGL}(s_A) + \left[z_{PGL}(s_B)-z_{PGL}(s_A)\right]t,\qquad c(s) = 4\,\frac{C}{12}\,t(1-t)')
    R.eq(r'z_{bot}(s) = q(s) - D + c(s) + r,\qquad z_{top}(s) = z_{bot}(s) + d_b')
    R.h('2. Haunch and the exact minimum')
    R.eq(r'h(s,y) = z_{PGL}(s) + \delta(y) - t_d - z_{top}(s)')
    R.p('Because the PGL is straight (brokeback) or a single parabola over each beam, z_PGL - q is itself -k_q L_b^2 t(1-t). The haunch '
        'over a beam is therefore an exact quadratic in t, and its minimum on the beam (t from t_lo to t_hi) has a closed form:')
    R.eq(r'h(t,y) = \delta(y) - t_d - r - d_b + D - m\,t(1-t),\qquad m = k_q L_b^2 + \frac{C}{3}\ \ [\mathrm{ft}]')
    R.eq(r'W = \max_{t_{lo}\leq t\leq t_{hi}} m\,t(1-t) = \max\left(\frac{m}{4},\ m\,t_{lo}(1-t_{lo}),\ m\,t_{hi}(1-t_{hi})\right)')
    R.p('t = 1/2 always lies on the beam, so m/4 covers upward camber; the end terms cover a net sag. The notebook finds the same minimum '
        'by locating every quadratic extremum between breakpoints. One common drop then enforces the required minimum haunch everywhere:')
    R.eq(r'D = h_{req} - \min_{spans,\,beams}\left[\delta_{min,k} - t_d - r_k - d_b - W\right]')
    a, b = ctx['haunch_min_range']
    beam = f"INDEX('Camber & Haunch'!$B${a}:$B${b},MATCH(Haunch_Min,'Camber & Haunch'!$F${a}:$F${b},0))"
    R.p(live('Current: D = ', NUM('Sup_Drop', '0.0000000000'), ' ft. The minimum flange haunch ', NUM('Haunch_Min', '0.000000'),
             ' in (required ', NUM('Bm_Hmin', 'General'), ' in) is on beam ', (beam,), '.'))
    for k, i in enumerate((2, 1), start=1):
        R.chart(charts.camber_haunch(R.data(), ctx, i),
                live(f'Figure {k}. Span {i} over CL beam 1: deck soffit, cambered beam top (C = ', NUM(f'Sp{i}_C', '0.000'),
                     ' in) and the haunch between them (shaded), vertical scale x40 about the zero-camber beam top. '
                     'Dims B, C and D are read from section D.'))
    R.h('3. Build-up dimensions and case')
    R.p('Dims B, C and D are the haunch over the beam centerline at the begin bearing, midspan and end bearing. The build-up case follows '
        'the BCC template rule using the vertical-curve mid-ordinate of the FGE over the beam centerline:')
    R.eq(r'R_m = 12\left[z(s_{mid}) - \frac{z(s_A)+z(s_B)}{2}\right];\quad \mathrm{Case\ 1}: R_m=0,\ \mathrm{Case\ 2}: R_m<0,\ \mathrm{Case\ 3}: 0<R_m<C,\ \mathrm{Case\ 4}: R_m>C')
    rows = ctx['buildup_rows'].values()
    cases = _range('Camber & Haunch', 'H', rows)
    R.p(live((f'IF(COUNTIF({cases},"Case 1")={len(rows)},' + TXT('A straight PGL in each span gives R_m = 0, so every beam is Case 1.')
              + ',' + TXT('Not every beam is Case 1; see section D of the Camber & Haunch tab.') + ')',)))
    R.vars([
        ('C', 'Net camber (Mathcad NetCamber)', 'Sp1_C, Sp2_C', 'in'),
        ('D', 'Common support drop', 'Sup_Drop', 'ft'),
        ('td, db, hreq', 'Deck thickness, beam depth, minimum haunch', 'Deck_T, Bm_D, Bm_Hmin', 'ft, ft, in'),
        ('m, W', 'Haunch curvature term and its maximum on the beam', 'Sp1_m, Sp1_W ...', 'ft'),
        ('B, C, D dims', 'Build-up over CL beam at begin, mid, end', 'Camber & Haunch section D', 'in'),
    ])


def fge_method(R, ctx):
    R.title('FGE - Method', 'Finished grade elevations at every line and station')
    R.h('1. Stations')
    R.p('Each span lists its datum, the two CL bearings and equal increments between bearings (tenth points with the default 10 increments):')
    R.eq(r's_k = s_A + k\,\frac{L_b}{N},\qquad k = 1,\ \ldots,\ N-1,\qquad N = 10')
    R.h('2. Elevations')
    R.eq(r'\mathrm{FGE}(s,y) = z_{PGL}(s) + \delta(y)')
    R.p(live('Lines are the two deck edges (copings) at y = yc +/- W/2, each beam centerline and the PGL. The FFBW columns are the PGL '
             'extended: the deck itself ends ', NUM('Phys_DeckClr', 'General'), ' in inside each FFBW (row "Deck present?"). Station labels '
             'use the notebook format (hundredths, e.g. 16+46.96).'))
    R.chart(charts.fge_grid(R.data(), ctx, 1),
            'Figure 1. FGE grid for Span 1 in plan, read from the FGE table: every red point is one finished-grade elevation.')
    R.note('A barrier-face (gutter) line is not tabulated: the notebook does not define the barrier. Add the barrier to the notebook '
           'CONFIG first if that line is wanted.')
    R.vars([
        ('N', 'Increments between bearings', 'Des_FGEInc', '10'),
        ('W, yc', 'Deck width and center offset', 'Deck_W, Deck_C', 'ft'),
        ('δ(y)', 'Cross-slope drop at the line offset', 'FGE column D', 'ft'),
    ])


def seats_method(R, ctx):
    R.title('Bearing Seats - Method', 'Seat planes, bearing slopes, pedestal tops and bearing plate data')
    R.h('1. Seat plane over the pad')
    R.p('The pad top is the chord of the cambered soffit between the pad ends, so the pad bears on a plane rather than a point. The '
        'chord gap at the bearing centerline is reported (it is a few millionths of an inch):')
    R.eq(r's_{lo,hi} = s_{brg} \mp \frac{L_p}{2},\qquad z_{pad} = \frac{z_{bot}(s_{lo}) + z_{bot}(s_{hi})}{2},\qquad \beta = \frac{z_{bot}(s_{hi}) - z_{bot}(s_{lo})}{L_p}')
    R.eq(r'z_{ped} = z_{pad} - t_p,\qquad \mathrm{gap} = 12\,\left[z_{bot}(s_{brg}) - z_{pad}\right]')
    R.fig(figures.pad_chord(), 'Figure 1. Pad-top plane chords the soffit (curvature exaggerated).')
    R.h('2. Pedestal tops')
    R.p('Each pedestal top is the seat plane extended longitudinally and level transversely (with bevels at the transverse edges). '
        'End-bent pedestals run from the FFBW (corners A, C) to the span-side face (corners B, D). The pier carries one pedestal per beam line: '
        'Span 1 seat plane up to the Span 1 pad end (E), a straight transition to the Span 2 pad start (F), then the Span 2 seat plane.')
    R.eq(r'z(s) = z_{ped} + \beta\,(s - s_{brg});\qquad \mathrm{EL\ A=C} = z(s_{back}),\quad \mathrm{EL\ B=D} = z(s_{front})')
    R.fig(figures.pedestal_plan(), 'Figure 2. Corner letters used in the CAD pedestal tables.')
    R.chart(charts.pier_pedestal(R.data(), ctx),
            live('Figure 3. Pier pedestal top along beam line 1, from Bearing Seats section C (the cap top, ', NUM('Pier_Top', '0.0000'),
                 ' ft, is below this window).'))
    R.h('3. Bearing plate data (BCC template rule)')
    R.eq(r'\mathrm{Beveled\ plate} = \mathrm{YES}\ \ \mathrm{if}\ \ |\beta| > 0.02;\qquad X = 0.5 + C\beta\ (\beta>0),\quad Y = 0.5 - C\beta\ (\beta<0),\quad C = 12\ \mathrm{in\ (pad\ types\ D\ to\ J)}')
    slopes = _range('Bearing Seats', 'T', ctx['seat_rows'].values())
    largest = f'MAX(MAX({slopes}),-MIN({slopes}))'
    R.note(live('Difference from the BCC template: the template levels seats when |slope| < 0.5%. The notebook slopes every seat to the actual '
                'soffit chord, so the pedestal tops here follow the notebook. ',
                (f'IF({largest}<=Des_BevelLimit,' + TXT('All current slopes are within the ') + '&TEXT(Des_BevelLimit*100,"0.0")&'
                 + TXT('% beveled-plate limit (largest ') + f'&TEXT({largest}*100,"0.000")&' + TXT('%).') + ','
                 + TXT('Some seat slopes exceed the beveled-plate limit; see the CAD Bearing Plates tab.') + ')',)))
    R.vars([
        ('Lp, tp', 'Pad length and thickness', 'EB1_PadL, EB1_Pad, Pier_PadL, Pier_Pad', 'in'),
        ('β', 'Bearing (seat) slope, + up-station', 'Bearing Seats col T', 'ft/ft'),
        ('zped', 'Pedestal top at CL bearing', 'Bearing Seats col V', 'ft'),
    ])


def substructure_method(R, ctx):
    R.title('Substructure - Method', 'Cap placement, pedestal heights, cap corners and pile cutoffs')
    R.h('1. Cap location and length')
    R.eq(r's_{cap} = s_{FFBW} \pm \left(\frac{w}{2} - t_{bw} - i\right)\ \ (\mathrm{end\ bents}),\qquad s_{cap} = s_{pier}\ \ (\mathrm{pier})')
    R.eq(r'L_{cap} = (n-1)\,p + d_p + 2c\qquad (p = 3d\ \mathrm{at\ the\ pier},\ p = 54\ \mathrm{in\ at\ end\ bents})')
    R.h('2. Cap top from the lowest pedestal')
    R.p('Unless a fixed cap top is entered, the cap top is set so the lowest pedestal (measured at its CL bearing) has the minimum height:')
    R.eq(r'z_{cap} = \min_j\left[z_{ped,j} - \sigma\,(y_j - y_c)\right] - \frac{h_{ped,min}}{12},\qquad h_j = 12\left[z_{ped,j} - \left(z_{cap} + \sigma\,(y_j-y_c)\right)\right]')
    R.h('3. Piles')
    R.eq(r'y_i = y_c + \left(i - \frac{n+1}{2}\right)p,\qquad z_{cutoff,i} = z_{cap} + \sigma\,(y_i - y_c) - d_{cap} + e_{emb}')
    R.chart(charts.cap_section(R.data(), ctx, 'EB1'),
            'Figure 1. End Bent 1 transverse section, drawn from the Substructure and Bearing Seats tables: cap, piles, pedestals, pads, '
            'cheek webs with bottom lugs, and the backwall behind the cap (dashed).')
    R.chart(charts.cap_section(R.data(), ctx, 'Pier'), 'Figure 2. Pier transverse section, drawn from the Substructure and Bearing Seats tables.')
    R.h('4. End-bent backwalls, cheek webs and lugs (section F)')
    R.p(live('The ', NUM('EB1_Len*12', '0.0'), '-in cap body follows from the pile layout. Two ', NUM('EB1_CkT', 'General'),
             '-in cheek webs sit outside the cap ends over the full cap width, so the cage envelope is ', NUM('EB1_OutIn', '0.0'),
             ' in (D-027). The backwall sits on the cap rear between the inset and the FFBW, across the full cap length:'))
    R.eq(r'L_{out} = L_{cap} + 2t_c,\qquad z_{bw}(y) = z_{PGL}(s_{FFBW}) + \delta(y) - d_{seat},\qquad h_{bw}(y) = z_{bw}(y) - z_{cap}(y)')
    R.eq(r'V_{bw} = t_{bw}\int_{y_L}^{y_R} h_{bw}(y)\,dy\quad(\mathrm{exact\ trapezoids:\ }h_{bw}\ \mathrm{is\ linear\ between\ the\ cap\ ends\ and\ the\ crown})')
    R.p(live('Each cheek runs from the local cap underside at its cap end to its entered top (the D-033 barrier-return base, an MSE interface '
             'input). A bottom lug hangs below the cap and a rear lug stands behind it, both ', NUM('EB1_LugP', 'General'), ' in deep and ',
             NUM('EB1_LugW', 'General'), ' in wide on the outer face:'))
    R.eq(r'z_{cheek,bot} = z_{cap}(y_{end}) - d_{cap},\qquad z_{lug,bot} = z_{cheek,bot} - p,\qquad s_{lug} = s_{rear} \pm p')
    R.eq(r'V_{cheek} = w\,t_c\,(z_{top}-z_{bot}) + w\,w_l\,p + p\,w_l\,(z_{top}-z_{bot}+p)')
    R.note('D-053: both backwall tops have two straight 2% slopes down from the PGL, level through their 12-in thickness, and lie 4-1/4 in below the FFBW deck-top lines. Ramp slabs thicken locally (about 4.05-4.95 in at checked CAD points). Recheck fit and steel cover after profile edits; downstream slab/backwall design and drawings require reconciliation. Cheek geometry is unchanged.')
    R.vars([
        ('σ, yc', 'Cap cross slope and center offset', 'Pier_Slope, Pier_Ctr ...', 'ft/ft, ft'),
        ('hped,min', 'Minimum pedestal height at CL bearing', 'Pier_Reveal, EB1_Reveal', 'in'),
        ('n, p, dp, c', 'Pile count, spacing, size and end clearance', 'Substructure section A', '-, ft, in, in'),
        ('eemb', 'Pile embedment above the cap underside', 'Pier_Emb, EB1_Emb', 'in'),
    ])


def beam_data_method(R, ctx):
    R.title('Beam Data - Method', 'Bars 5K spacing zones and FIB standard bar dimensions')
    R.h('1. Mapping the Mathcad stirrup layout to FDOT notation')
    R.p('The Mathcad stirrup design gives zones A1-A3 and S1-S4 for half the beam. They map to the FIB Bars 5K table columns as follows; '
        'zones A1 and A2 share one spacing (checked in Beam Data column O) and together form S*, and A3 forms Sj (FIB-36 standard: S* = 4, Sj = 12):')
    zone = lambda z: live(NUM(f'MC1_{z}_n', '0'), ' @ ', NUM(f'MC1_{z}_s', 'General'))
    R.vars([
        ('S*, V*', 'A1 + A2 spaces at the A1 spacing', 'MC1_A1_n + MC1_A2_n, MC1_A1_s',
         live(NUM('MC1_A1_n+MC1_A2_n', '0'), ' @ ', NUM('MC1_A1_s', 'General'), ' in')),
        ('Sj, Vj', 'A3 spaces and spacing', 'MC1_A3_n, MC1_A3_s', zone('A3') + '&" in"'),
        ('S1-S3, V1-V3', 'Mathcad zones S1-S3', 'MC1_S1_n ... MC1_S3_s',
         zone('S1') + '&", "&' + zone('S2')[1:] + '&", "&' + zone('S3')[1:] + '&" in"'),
        ('S4, V4', 'Remaining spaces to midspan at V4', 'computed here', live('@ ', NUM('MC1_S4_s', 'General'), ' in')),
        ('es', 'Beam end to first stirrup (Mathcad EndSpace)', 'MC1_EndSp', live(NUM('MC1_EndSp', 'General'), ' in')),
    ])
    R.eq(r'S_4 = \left\lfloor \frac{6L - e_s - \left(S^*V^* + S_jV_j + S_1V_1 + S_2V_2 + S_3V_3\right)}{V_4} \right\rfloor,\qquad \mathrm{center\ space} = 2\left[6L - e_s - \sum S_iV_i - S_4V_4\right]')
    R.eq(r'N_{5K} = 2\left(1 + S^* + S_j + S_1 + S_2 + S_3 + S_4\right)\qquad (L\ \mathrm{in\ ft},\ \mathrm{dimensions\ in\ in})')
    R.chart(charts.stirrup_zones(R.data(), ctx),
            'Figure 1. Bars 5K zones along one half of beam 1-1, read from Beam Data sections A and B (S4 follows the current beam length).')
    R.h('2. Standard bar dimensions (BCC template constants, square ends)')
    R.eq(r'C_{3C} = \frac{17.5}{\sin\phi},\quad A_{3D} = \frac{8.5}{\sin\phi},\quad B_{3D} = \frac{15}{\sin\phi},\quad D_{4M} = \frac{44}{\sin\phi}\quad[\mathrm{in}],\qquad \phi = 90^\circ - \mathrm{skew}')
    R.eq(r'N_{3D3} = 2\left(1 + \left\lceil \frac{12L - 2\,(e_s + S^*V^* + S_jV_j + 18)}{18} \right\rceil\right),\qquad N_{4M3} = 1 + \left\lceil \frac{12L - 2\,(3.5 + 2\cdot 7 + 6k_M + 12)}{12} \right\rceil')
    R.note('These constants and counts reproduce the BCC example workbook formulas (FIB standard bar bending). The template used a 2.5 in end '
           'distance; this workbook uses the Mathcad EndSpace. Verify against the current FDOT FIB standard plans before issue.')


def clearance_method(R, ctx):
    R.title('Clearance - Method', 'Minimum vertical clearance between the beam bottoms and the roadway profile')
    R.p('Clearance is the beam bottom minus the roadway elevation directly below, evaluated only where a beam exists (beam ends and the open '
        'pier joint are excluded). The roadway comes from Geometry Report.xml, mapped to bridge stations by the notebook assumption '
        '(bridge station at XML zero = Rd_Shift).')
    R.eq(r'c(s) = z_{bot}(s) - z_{road}(s),\qquad z_{road}(s) = z_0 + g\,(s-s_0) + q\,(s-s_0)^2')
    R.p('On every roadway segment piece under a beam both terms are quadratic, so c is quadratic. Fitting it through the piece ends and midpoint '
        'gives the exact minimum:')
    R.eq(r'A = 2\,(c_b - 2c_m + c_a),\quad B = c_b - c_a - A,\quad t^* = -\frac{B}{2A}\ (A>0),\qquad c_{min} = \min\left(c_a,\ c_b,\ c_a + Bt^* + At^{*2}\right)')
    R.chart(charts.clearance(R.data(), ctx),
            'Figure 1. Clearance c(s) under the beams against the required clearance, drawn from the roadway segments and beam formulas; '
            'the minimum is the section A result.')
    R.p(live('Current: minimum ', NUM('Clr_Min', '0.0000'), ' ft at ', (STA('Clr_Sta'),), ' (required ', NUM('Rd_Req', 'General'),
             ' ft). The XML covers ', NUM('Clr_Cover', '0.000'), ' of ', NUM('Prof_L', '0.000'), ' ft, so the status is "', ('Clr_Status',),
             '" (no extrapolation).'))


def cad_fge_method(R, ctx):
    R.title('CAD FGE - Method', 'Finish Grade Elevations table for CADD')
    R.p('Each span table lists the increment distance, the station labels and the finished grade elevation of every line, linked directly to the '
        'FGE tab. Values keep full precision and display to 0.001 ft. Line names follow the template (copings, CL beams, PGL).')
    R.eq(r'\mathrm{FGE}(s,y) = z_{PGL}(s) + \delta(y)')
    R.chart(charts.fge_grid(R.data(), ctx, 2), 'Figure 1. Span 2 FGE grid in plan, read from the FGE table.')
    R.bullets(['INCR. DIST = (CL brg B - CL brg A)/10.',
               live('•  FFBW columns are the PGL extended; the deck ends ', NUM('Phys_DeckClr', 'General'), ' in inside each FFBW.'),
               'The barrier face is not in the notebook, so no gutter/barrier line is listed.'])


def cad_buildup_method(R, ctx):
    R.title('CAD Build-Up - Method', 'Build-Up & Deflection Data Table for prestressed I-beams')
    R.vars([
        ('Dim B, C, D', 'Required theoretical build-up over the CL beam at begin bearing, midspan, end bearing', 'Camber & Haunch section D', 'in'),
        ('Release', 'Net beam camber (prestress - beam weight) at release: Mathcad max(Defl.release)', 'MC1_Drel, MC2_Drel', 'in'),
        ('120 days', 'Net beam camber (prestress - beam weight) at 120 days: Mathcad max(Defl.120day)', 'MC1_D120, MC2_D120', 'in'),
        ('Dim A', 'Dead load deflection during deck pour: |Mathcad min(Defl.dl.nc)|', 'MC1_Dnc, MC2_Dnc', 'in'),
        ('Case', 'Build-up case from the mid-ordinate rule', 'Camber & Haunch col H', '-'),
    ])
    R.eq(r'C = \Delta_{120} - \mathrm{Dim\,A} - |\Delta_{DL,c}|')
    R.p('The build-up uses the Mathcad NetCamber C, which also subtracts the composite (barrier) dead-load deflection. That is why the '
        '120-day camber minus Dim A is larger than C by the composite deflection.')
    R.chart(charts.camber_haunch(R.data(), ctx, 2), 'Figure 1. Dims B, C and D over CL beam 1, Span 2 (vertical scale x40).')


def cad_beam_variables_method(R, ctx):
    R.title('CAD Beam Variables - Method', 'Florida-I Beam Table of Beam Variables')
    R.vars([
        ('Beam type, class', 'FIB-36 (Mathcad); concrete class designer input', 'Bm_Type, Des_ConcClass', '-'),
        ("f'c, f'ci", 'Mathcad beam strengths x 1000', 'MC1_fc, MC1_fci', 'psi'),
        ('Stnd. ptrn. type', 'Designer strand-pattern number', 'Des_Pattern', '-'),
        ('End elev. cond.', '1 if Dim P = 0, 2 if P < 0, 3 if P > 0', 'Beam Geometry col O', '-'),
        ('Plan view case', 'Case 1 = square end (zero skew)', 'Des_PlanCase', '-'),
        ('Brg. plate mark', 'span-beam-end (end 1 = begin, end 2 = end)', 'CAD Bearing Plates', '-'),
        ('Angle phi', '90 deg - skew', 'Beam Geometry cols I, J', 'deg'),
        ('Dim P, J, K1, K2', 'See Beam Geometry Method', 'Beam Geometry cols K-N', 'in, ft-in'),
        ('Dim L, R', 'Sloped length + R (template) and shortening R', 'Beam Geometry cols R-U', 'ft-in, in'),
        ('3C1 ... 4M3', 'Standard bar dimensions and counts', 'Beam Data section B', 'ft-in, no.'),
        ('5K, S1-S4, V1', 'Total Bars 5K, spaces per zone, V1 spacing', 'Beam Data sections A-B', 'no., in'),
    ])
    R.fig(figures.beam_end_dims(), 'Figure 1. End-of-beam dimensions.')
    R.p('Feet-inch values are split into a feet text cell ("1\'-") and an inches number, as in the BCC template, so the table pastes into the '
        'existing CADD cell layout.')


def cad_plates_method(R, ctx):
    R.title('CAD Bearing Plates - Method', 'Bearing Plate Data Table - Type 2')
    R.vars([
        ('Brg. plate mark', 'span-beam-end', 'Bearing Seats section D', '-'),
        ('Pad type', 'Designer input checked against the pad dimension table', 'Des_PadType', '-'),
        ('Beam end', '1 = begin of span, 2 = end of span', '-', '-'),
        ('Slope (%)', '100 x seat-plane slope beta from the notebook method, + up-station', 'Bearing Seats col T', '%'),
        ('Beveled plate', 'YES when |beta| > 2% (template); C, X, Y only then', 'Des_BevelLimit', '-'),
    ])
    R.fig(figures.pad_chord(), 'Figure 1. The slope is the chord of the soffit over the pad length.')


def cad_pedestals_method(R, ctx):
    R.title('CAD Pedestal Elevs - Method', 'Top of pedestal, cap and pile cutoff elevations')
    R.p('End-bent pedestal tops are planar: A and C on the FFBW edge, B and D on the span-side edge (A, B on the left side of the beam line). '
        'Tops are level transversely, so A = C and B = D. The pier pedestal on each beam line is two seat planes joined by a straight '
        'transition, so the table adds the break-line elevations E and F and the elevation at the pier centerline.')
    R.eq(r'z(s) = z_{ped} + \beta\,(s - s_{brg})')
    R.fig(figures.pedestal_plan(), 'Figure 1. Corner and break-line letters.')
    R.chart(charts.pier_pedestal(R.data(), ctx), 'Figure 2. Pier pedestal top along beam line 1, from Bearing Seats section C.')
    R.p('Cap tops, undersides and pile cutoffs come from the Substructure tab (cutoff = local cap top - cap depth + embedment).')
    R.chart(charts.cap_section(R.data(), ctx, 'Pier'), 'Figure 3. Pier section, drawn from the Substructure and Bearing Seats tables.')
    R.p(live('The end-bent wall table lists the backwall top at both cap ends and at the PGL (deck-top line at the FFBW less the ',
             NUM('EB1_BwDrop', 'General'), '-in EB1 / ', NUM('EB2_BwDrop', 'General'), '-in EB2 seat), both cheek tops (D-033 return bases), the cheek bottom at the cap underside, the lug '
             'bottom and the rear-lug face station where the MSE copings stop. All values link to Substructure section F.'))
    R.chart(charts.cap_section(R.data(), ctx, 'EB2'), 'Figure 4. End Bent 2 section with cheeks and the backwall behind the cap (dashed).')


def qc_method(R, ctx):
    R.title('QC - Method', 'How the workbook proves it reproduces the notebook')
    R.p('Each check passes when the workbook formula result agrees with the independent notebook value within its tolerance:')
    R.eq(r'\mathrm{PASS}\ \Leftrightarrow\ \left|\,x_{workbook} - x_{notebook}\,\right| \leq \tau,\qquad \mathrm{Overall\ PASS}\ \Leftrightarrow\ \mathrm{no\ FAIL}')
    R.fig(figures.data_flow(), 'Figure 1. The notebook values are captured by the generator; Excel recomputes the workbook values from the inputs.')
    R.bullets([
        'Section A compares each key workbook result with the value the notebook model computed when the workbook was built (blue snapshot column).',
        'Tolerances are 1e-9 ft for stations and elevations, 1e-6 in for haunch and pedestal heights, and 0.001 ft only where Mathcad rounds the beam spacing.',
        'Section B compares every FGE cell with the notebook deck_top at the same station and offset (hidden snapshot columns L:M).',
        'Section C confirms the configuration stays inside what the closed-form Excel formulas represent (profile and deck modes, a straight PGL over each beam, span continuity).',
        'Section C also compares every generated input with the value written at build (hidden columns O:R). An edit shows CHECK here and '
        'turns the banner on every Method tab amber; the charts and tables follow the edit.',
        'Overall PASS requires every check to pass. A FAIL after editing an input means the workbook no longer matches the notebook: update CONFIG and rebuild.',
        'The regression test tests/test_geometry_cad_tables.py also recomputes the notebook model live and compares the saved workbook values.',
    ])


CONTENT = {
    'Inputs': inputs_method, 'Profile & Deck': profile_method, 'Beam Geometry': beam_geometry_method,
    'Camber & Haunch': camber_method, 'FGE': fge_method, 'Bearing Seats': seats_method, 'Substructure': substructure_method,
    'Beam Data': beam_data_method, 'Clearance': clearance_method, 'CAD FGE': cad_fge_method, 'CAD Build-Up': cad_buildup_method,
    'CAD Beam Variables': cad_beam_variables_method, 'CAD Bearing Plates': cad_plates_method,
    'CAD Pedestal Elevs': cad_pedestals_method, 'QC': qc_method,
}


def build_method(book, ctx, name):
    ws = book.sheet(f'{name} Method', '8A96A3')
    R = Renderer(book, ws)
    CONTENT[name](R, ctx)
    # Print the text and charts (A:M) one page wide; the figure data to the right stays on screen only.
    ws.print_area = f'A1:{chr(65 + COLS)}{R.row}'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    return ws
