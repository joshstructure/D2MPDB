"""Live Excel charts for the Method tabs.

Every Method-tab figure that shows project geometry or numbers is a native
Excel XY chart, not a picture. Each plotted point and each numeric label is a
formula in a figure-data block to the right of the text (column P onward),
built from the same defined names as the calculation tabs. When an input
changes, Excel recalculates the block and redraws the chart with the tables.

Labels are one-point series whose name is a text formula; the chart shows the
series name as a data label, so label numbers update with the workbook.
"""
from __future__ import annotations

from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import NumDataSource, NumFmt, NumRef
from openpyxl.chart.error_bar import ErrorBars
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import SeriesLabel, StrRef
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineEndProperties, LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RichTextProperties
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .workbook import BOTTOM, DELTA, PGL

DATA_COL = 16           # column P
EMU_PT = 12700
INK, BLUE, RED, GREEN, GREY = '1F2A36', '2B6CB0', 'C0392B', '2F855A', '8A96A3'
STEEL, CONCRETE, HAUNCH, CHEEK = '6B7785', '5A6570', 'E3B85C', '5B7F3A'
ZONES = ('C0392B', 'E67E22', '2B6CB0', '2F855A', '6B46C1', 'B7791F')
F_HEAD = Font(name='Arial', size=10, bold=True, color='1F2A36')
F_SUB = Font(name='Arial', size=8, bold=True, color='5A6570')
F_DATA = Font(name='Arial', size=8, color='5A6570')

# FIB-36 outline in inches (48-in top flange, 36-in depth), from the BCC section sketch.
FIB_OUTLINE = [(-24, 36), (24, 36), (24, 32.5), (7, 31), (3.5, 27.5), (3.5, 10), (19, 7), (19, 0),
               (-19, 0), (-19, 7), (-3.5, 10), (-3.5, 27.5), (-7, 31), (-24, 32.5), (-24, 36)]


# ----------------------------------------------------------- text helpers --
def STA(x: str) -> str:
    """Station label (hundredths) of an Excel expression, e.g. 16+46.96."""
    return f'TEXT(INT(ROUND({x}*100,0)/10000),"0")&"+"&TEXT(MOD(ROUND({x}*100,0),10000)/100,"00.00")'


def TXT(text: str) -> str:
    """Excel string literal(s): line breaks become CHAR(10); each literal stays under Excel's 255-character limit."""
    pieces = []
    for k, line in enumerate(text.split('\n')):
        if k:
            pieces.append('CHAR(10)')
        parts = [line[i:i + 250] for i in range(0, len(line), 250)] or ([''] if not pieces else [])
        pieces += ['"' + part.replace('"', '""') + '"' for part in parts]
    return '&'.join(pieces) or '""'


def live(*parts) -> str:
    """Formula joining literal text (str) and Excel expressions (wrapped in a 1-tuple)."""
    return '=' + '&'.join(p[0] if isinstance(p, tuple) else TXT(p) for p in parts)


def NUM(expr: str, fmt: str) -> tuple:
    return (f'TEXT({expr},"{fmt}")',)


def _text(size, color, bold=False, rot=None):
    props = CharacterProperties(sz=int(size * 100), b=bold, solidFill=color)
    body = RichTextProperties(rot=rot, vert='horz') if rot is not None else RichTextProperties()
    return RichText(bodyPr=body, p=[Paragraph(pPr=ParagraphProperties(defRPr=props), endParaRPr=props)])


def _formula(expr) -> str:
    return '=' + str(expr)


# ------------------------------------------------------------- data area --
class DataArea:
    """Allocates figure-data blocks down column P of one Method tab."""

    def __init__(self, ws):
        self.ws, self.row = ws, 4
        ws.cell(row=1, column=DATA_COL, value='Figure data').font = Font(name='Arial', size=12, bold=True, color='1F2A36')
        ws.cell(row=2, column=DATA_COL, value='Live formulas that draw the charts on this tab. Do not edit; '
                'change the notebook CONFIG and rebuild.').font = F_DATA
        for k, width in enumerate((46, 14, 14, 12)):
            ws.column_dimensions[get_column_letter(DATA_COL + k)].width = width


class Handle:
    def __init__(self, ws, letters, r0, r1):
        self.r0, self.r1, self.L = r0, r1, letters
        self.x = f'{letters[1]}{r0}:{letters[1]}{r1}'
        self.y = f'{letters[2]}{r0}:{letters[2]}{r1}'

    def pt(self, k):
        """(x, y) cell references of point k, for placing a label on the line."""
        return f'${self.L[1]}${self.r0 + k}', f'${self.L[2]}${self.r0 + k}'


class Figure:
    """One chart: series blocks (name | x | y | error+) written as formulas, then a ScatterChart."""

    def __init__(self, area: DataArea, title: str):
        self.area, self.ws = area, area.ws
        self.L = [get_column_letter(DATA_COL + k) for k in range(4)]
        r = area.row
        self.ws.cell(row=r, column=DATA_COL, value=title).font = F_HEAD
        for k, head in enumerate(('Series / label text', 'x', 'y', 'error +')):
            self.ws.cell(row=r + 1, column=DATA_COL + k, value=head).font = F_SUB
        self.row = r + 2
        self.series = []

    def _put(self, row, k, value, fmt='0.0000'):
        cell = self.ws.cell(row=row, column=DATA_COL + k, value=value)
        cell.font = F_DATA
        cell.number_format = fmt
        return cell

    def helper(self, label, expr, fmt='0.000000') -> str:
        r = self.row
        self.row += 1
        self._put(r, 0, label, 'General')
        self._put(r, 1, _formula(expr), fmt)
        return f'${self.L[1]}${r}'

    def _block(self, name, pts, err=None):
        r0 = self.row
        self._put(r0, 0, name, 'General')
        for k, point in enumerate(pts):
            if point is None:
                continue                      # empty row = gap in the line
            x, y = point
            xref = f'${self.L[1]}${r0 + k}'
            self._put(r0 + k, 1, _formula(x))
            self._put(r0 + k, 2, _formula(y(xref) if callable(y) else y))
            if err is not None:
                self._put(r0 + k, 3, _formula(err(xref)))
        r1 = r0 + len(pts) - 1
        self.row = r1 + 2
        return r0, r1

    def _series(self, r0, r1):
        sheet = self.ws.title
        xs = Reference(self.ws, min_col=DATA_COL + 1, min_row=r0, max_row=r1)
        ys = Reference(self.ws, min_col=DATA_COL + 2, min_row=r0, max_row=r1)
        s = Series(ys, xs)
        s.tx = SeriesLabel(strRef=StrRef(f"'{sheet}'!${self.L[0]}${r0}"))
        s.smooth = False
        if s.spPr is None:
            s.spPr = GraphicalProperties(ln=LineProperties())
        if s.spPr.ln is None:
            s.spPr.ln = LineProperties()
        self.series.append(s)
        return s

    def line(self, name, pts, color=INK, width=1.0, dash=None, marker=None, size=4,
             arrows=False, err=None, err_color=HAUNCH, err_width=3.0) -> Handle:
        r0, r1 = self._block(name, pts, err)
        s = self._series(r0, r1)
        ln = s.spPr.ln
        if width:
            ln.solidFill = color
            ln.w = int(width * EMU_PT)
            if dash:
                ln.prstDash = dash
            if arrows:
                ln.headEnd = LineEndProperties(type='triangle', w='sm', len='sm')
                ln.tailEnd = LineEndProperties(type='triangle', w='sm', len='sm')
        else:
            ln.noFill = True
        if marker:
            s.marker = Marker(symbol=marker, size=size)
            s.marker.spPr = GraphicalProperties(solidFill=color, ln=LineProperties(solidFill=color))
        else:
            s.marker = Marker(symbol='none')
        if err is not None:
            ref = f"'{self.ws.title}'!${self.L[3]}${r0}:${self.L[3]}${r1}"
            s.errBars = ErrorBars(errDir='y', errBarType='plus', errValType='cust', noEndCap=True,
                                  plus=NumDataSource(numRef=NumRef(f=ref)),
                                  spPr=GraphicalProperties(ln=LineProperties(solidFill=err_color, w=int(err_width * EMU_PT))))
        return Handle(self.ws, self.L, r0, r1)

    def label(self, text, x, y, pos='t', color=INK, size=7.5, bold=False, rot=None, marker=None):
        r0, r1 = self._block(text, [(x, y)])
        s = self._series(r0, r1)
        s.spPr.ln.noFill = True
        if marker:
            s.marker = Marker(symbol=marker, size=5)
            s.marker.spPr = GraphicalProperties(solidFill=color, ln=LineProperties(solidFill=color))
        else:
            s.marker = Marker(symbol='none')
        s.dLbls = DataLabelList(showSerName=True, showVal=False, showLegendKey=False, showCatName=False,
                                showPercent=False, showBubbleSize=False, dLblPos=pos, txPr=_text(size, color, bold, rot))

    def bounds(self, x0, x1, y0, y1):
        """Invisible corner points so Excel's automatic axes leave room for labels."""
        self.line('frame (invisible, sets the chart window)', [(x0, y0), (x1, y1)], width=0)

    def chart(self, width=24.0, height=9.0, x_title=None, y_title=None, x_axis=True, y_axis=True,
              x_fmt='General', y_fmt='General', grid=True) -> ScatterChart:
        """Lines are labelled directly (Excel ignores legend-entry deletions in generated charts), so no legend."""
        ch = ScatterChart()
        ch.scatterStyle = 'lineMarker'
        ch.width, ch.height = width, height
        ch.display_blanks = 'gap'
        ch.visible_cells_only = False
        for s in self.series:
            ch.series.append(s)
        for ax, pos, title, show, fmt in ((ch.x_axis, 'b', x_title, x_axis, x_fmt), (ch.y_axis, 'l', y_title, y_axis, y_fmt)):
            ax.axPos = pos                    # openpyxl writes 'l' for both XY axes unless told
            ax.delete = not show
            ax.title = title if show else None
            if ax.title is not None:
                ax.title.overlay = False
            ax.numFmt = NumFmt(formatCode=fmt, sourceLinked=False)
            ax.majorTickMark = 'out'
            ax.tickLblPos = 'nextTo'
            ax.crosses = 'min'
            ax.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill='E6E9ED', w=6350))) if show and grid else None
            ax.txPr = _text(8, '5A6570')
        ch.legend = None
        self.area.row = self.row + 1
        return ch


def _gaps(segments):
    """[(p, q), ...] -> p, q, gap, p, q, gap ... for one series of separate segments."""
    pts = []
    for seg in segments:
        pts += list(seg) + [None]
    return pts[:-1]


# ---------------------------------------------------------------- figures --
def profile(area, ctx):
    f = Figure(area, 'Profile along the bridge')
    n = 12
    left = [(f'Prof_Sta0+{k}/{n}*(Prof_PISta-Prof_Sta0)', PGL) for k in range(n + 1)]
    right = [(f'Prof_PISta+{k}/{n}*(Prof_Sta1-Prof_PISta)', PGL) for k in range(1, n + 1)]
    pgl = f.line('PGL', left + right, color=BLUE, width=1.75)
    lo = f.helper('Lowest PGL less 0.25 ft', f'MIN({pgl.y})-0.25')
    hi = f.helper('Highest PGL plus 0.1 ft', f'MAX({pgl.y})+0.1')
    spans = ctx['spans']
    datums = [('Sp1_Sta0', 'FFBW End Bent 1')] + [(f'Sp{i}_Sta1', 'CL Pier' if i < len(spans) else 'FFBW End Bent 2')
                                                   for i in range(1, len(spans) + 1)]
    f.line('Span datums', _gaps([((s, lo), (s, hi)) for s, _ in datums]), color=GREY, width=0.75, dash='dash')
    for s, name in datums:
        f.label(live(name + '  ', (STA(s),)), s, lo, pos='b', size=7)
    for s, pos in (('Prof_Sta0', 'r'), ('Prof_Sta1', 'l')):
        f.label(live('EL ', NUM(PGL(s), '0.000')), s, PGL(s), pos=pos, marker='circle')
    f.label(live('PI  ', (STA('Prof_PISta'),), '\nEL ', NUM(PGL('Prof_PISta'), '0.000')), 'Prof_PISta', PGL('Prof_PISta'),
            pos='t', marker='circle')
    for g, a, b in (('g1', 'Prof_Sta0', 'Prof_PISta'), ('g2', 'Prof_PISta', 'Prof_Sta1')):
        mid = f'({a}+{b})/2'
        f.label(live(g + ' = ', NUM(f'Prof_G{g[1]}*100', '+0.0000;-0.0000'), '%'), mid, PGL(mid), pos='t', color=BLUE, size=8.5)
    f.bounds('Prof_Sta0-5', 'Prof_Sta1+5', f'{lo}-0.3', f'{hi}+0.2')
    return f.chart(x_title='Station (ft)', y_title='Elevation (ft)', x_fmt='0', y_fmt='0.0')


def typical_section(area, ctx):
    f = Figure(area, 'Section at Span 1 midspan (true scale)')
    n = ctx['model'].count
    zp = f.helper('PGL at Span 1 midspan (ft)', PGL('Sp1_Mid'))
    top = lambda y: f'{zp}+{DELTA(y)}'
    crown = 'MAX(Deck_YL,MIN(0,Deck_YR))'
    ys = ['Deck_YL', crown, 'Deck_YR']
    deck = [(y, top(y)) for y in ys] + [(y, f'{top(y)}-Deck_T') for y in reversed(ys)] + [('Deck_YL', top('Deck_YL'))]
    f.line('Deck', deck, color=CONCRETE, width=1.25)
    zb = {k: f.helper(f'Beam {k} bottom at midspan (ft)', BOTTOM(1, '0.5', f'Bm{k}_Raise')) for k in range(1, n + 1)}
    f.line('FIB-36 beams', _gaps([[(f'Bm{k}_Off+({px})/12*Bm_Bf/48', f'{zb[k]}+({py})/36*Bm_D') for px, py in FIB_OUTLINE]
                                  for k in range(1, n + 1)]), color=INK, width=1.0)
    for k in range(1, n + 1):
        samples = [(f'Bm{k}_FL+{j}/24*(Bm{k}_FR-Bm{k}_FL)', f'{zb[k]}+Bm_D') for j in range(25)]
        f.line(f'Haunch over beam {k}', samples, width=0, err=lambda X: f'{top(X)}-Deck_T-({zb[k]}+Bm_D)', err_width=4.0)
    checks = []
    for k in range(1, n + 1):
        for y in (f'Bm{k}_FL', f'Bm{k}_Off', f'Bm{k}_FR', f'MAX(Bm{k}_FL,MIN(0,Bm{k}_FR))'):
            checks.append((y, f'{zb[k]}+Bm_D'))
    f.line('Haunch check points', checks, width=0, marker='triangle', color=RED, size=5)
    low = f.helper('Lowest beam bottom (ft)', 'MIN(' + ','.join(zb.values()) + ')')
    f.line('Beam centerlines', _gaps([((f'Bm{k}_Off', f'{low}-0.25'), (f'Bm{k}_Off', f'{top(f"Bm{k}_Off")}+0.25'))
                                      for k in range(1, n + 1)]), color=GREY, width=0.6, dash='dashDot')
    f.line('PGL', [(0, f'{low}-0.25'), (0, f'{top(0)}+0.35')], color=BLUE, width=0.75, dash='dash')
    for k in range(1, n + 1):
        haunch = f'({top(f"Bm{k}_Off")}-Deck_T-{zb[k]}-Bm_D)*12'
        f.label(live(f'CL Beam {k}  ', NUM(f'Bm{k}_Off', '+0.0000;-0.0000'), ' ft\nhaunch @ CL ', NUM(haunch, '0.00'), ' in'),
                f'Bm{k}_Off', zb[k], pos='b', size=7)
    for y, pos in (('Deck_YL', 't'), ('Deck_YR', 't')):
        f.label(live('deck edge\n', NUM(y, '+0.0000;-0.0000'), ' ft'), y, top(y), pos=pos, size=7)
    f.label(live('PGL  EL ', NUM(zp, '0.000'), '\nSta ', (STA('Sp1_Mid'),)), 0, f'{top(0)}+0.35', pos='t', color=BLUE, size=7)
    slope = 'ABS(IF(Deck_Mode="uniform",Deck_Uni,Deck_Crown))*100'
    for y in ('Deck_YL/2', 'Deck_YR/2'):
        f.label(live(NUM(slope, '0.0'), '%'), y, top(y), pos='t', color=RED, size=8)
    ys_ = f.helper('Spacing dimension line (ft)', f'{low}-0.95')
    yd = f.helper('Deck-width dimension line (ft)', f'{low}-1.3')
    for k in range(1, n):
        f.line(f'Spacing {k}', [(f'Bm{k}_Off', ys_), (f'Bm{k+1}_Off', ys_)], width=0.75, arrows=True)
        f.label(live('S = ', NUM(f'Bm{k+1}_Off-Bm{k}_Off', '0.0000'), ' ft'), f'(Bm{k}_Off+Bm{k+1}_Off)/2', ys_, pos='t', size=7)
    f.line('Deck width', [('Deck_YL', yd), ('Deck_YR', yd)], width=0.75, arrows=True)
    f.label(live('deck width ', NUM('Deck_W*12', '0.0'), ' in (', NUM('Deck_W', '0.0000'), ' ft)'), 'Deck_C', yd, pos='b', size=7)
    f.bounds('Deck_YL-0.6', 'Deck_YR+0.6', f'{yd}-0.35', f'{top(0)}+0.9')
    return f.chart(height=10.5, x_title='Offset from PGL (ft), looking up-station', y_title='Elevation (ft)',
                   x_fmt='0', y_fmt='0.0')


def camber_haunch(area, ctx, i):
    f = Figure(area, f'Span {i}: camber and haunch over CL beam 1')
    E = f.helper('Vertical exaggeration', '40', 'General')
    N = 40
    t = lambda X: f'(({X})-Sp{i}_BrgA)/Sp{i}_Lb'
    chord = lambda X: f'(Sp{i}_ZA+(Sp{i}_ZB-Sp{i}_ZA)*{t(X)})'
    ref = lambda X: f'({chord(X)}-Sup_Drop+Bm_D+Bm1_Raise/12)'          # beam top with zero camber
    soffit = lambda X: f'({PGL(X)}+Bm1_dCL-Deck_T)'
    top = lambda X: f'({BOTTOM(i, t(X), "Bm1_Raise")}+Bm_D)'
    xs = [f'Sp{i}_Bm0+{k}/{N}*(Sp{i}_Bm1-Sp{i}_Bm0)' for k in range(N + 1)]
    deck = f.line('Deck soffit over CL beam 1', [(x, lambda X: f'({soffit(X)}-{ref(X)})*{E}') for x in xs], color=BLUE, width=1.5)
    beam = f.line('Cambered beam top', [(x, lambda X: f'({top(X)}-{ref(X)})*{E}') for x in xs],
                  color=INK, width=1.5, err=lambda X: f'({soffit(X)}-{top(X)})*{E}', err_width=3.0)
    f.line('Beam top with zero camber', [(f'Sp{i}_Bm0', 0), (f'Sp{i}_Bm1', 0)], color=GREY, width=1.0, dash='dash')
    f.label('deck soffit over CL beam 1', *deck.pt(N // 4), pos='t', color=BLUE, size=7.5)
    f.label(live('cambered beam top, C = ', NUM(f'Sp{i}_C', '0.000'), ' in'), *beam.pt(3 * N // 8), pos='t', color=INK, size=7.5)
    f.label('beam top with zero camber (chord less D)', f'Sp{i}_Bm0+0.75*(Sp{i}_Bm1-Sp{i}_Bm0)', 0, pos='b', color=GREY, size=7.5)
    r = ctx['buildup_rows'][(i, 1)]
    for s, name, c in ((f'Sp{i}_BrgA', 'Dim B', 'C'), (f'Sp{i}_Mid', 'Dim C', 'D'), (f'Sp{i}_BrgB', 'Dim D', 'E')):
        y0 = f'({top(s)}-{ref(s)})*{E}'
        y1 = f'({soffit(s)}-{ref(s)})*{E}'
        f.line(f'{name} arrow', [(s, y0), (s, y1)], color=RED, width=0.9, arrows=True)
        f.label(live(name + '\n', NUM(f"'Camber & Haunch'!{c}{r}", '0.000'), ' in'), s, y1, pos='t', color=RED, size=7.5)
    f.bounds(f'Sp{i}_Bm0-2', f'Sp{i}_Bm1+2', -1.5, f'MAX({deck.y})+2.5')
    return f.chart(height=8.5, x_title='Station (ft)', y_axis=False, x_fmt='0', grid=False)


def fge_grid(area, ctx, i):
    f = Figure(area, f'FGE grid, Span {i} (plan)')
    fg = ctx['fge'][i]
    st = [f"'FGE'!{c}{fg['station_row']}" for c in fg['cols']]
    lines = list(fg['lines'].items())
    offsets = [f"'FGE'!$C${lr}" for _, lr in lines]
    ymin = f.helper('Lowest line offset (ft)', 'MIN(' + ','.join(offsets) + ')')
    ymax = f.helper('Highest line offset (ft)', 'MAX(' + ','.join(offsets) + ')')
    f.line('Station lines', _gaps([((s, ymin), (s, ymax)) for s in st]), color=GREY, width=0.5)
    for (name, lr), off in zip(lines, offsets):
        pgl = name == 'PGL'
        f.line(f"='FGE'!B{lr}", [(s, off) for s in st], color=BLUE if pgl else INK, width=0.9 if pgl else 1.1,
               dash='dash' if pgl else None, marker='circle', size=3)
        f.label(f"='FGE'!B{lr}", f'{st[0]}-1.5', off, pos='l', size=7)
    last = len(st) - 1
    for k, (s, c) in enumerate(zip(st, fg['cols'])):
        text = live(fg['labels'][k] + '  ', (f"'FGE'!{c}{fg['label_row']}",))
        if k in (0, last):
            f.label(text, s, ymin, pos='b', size=6.5, rot=-5400000)
        else:
            f.label(text, s, ymax, pos='t', size=6.5, rot=-5400000)
    f.bounds(f'{st[0]}-16', f'{st[-1]}+2', f'{ymin}-3.2', f'{ymax}+3.2')
    return f.chart(height=10.0, x_title='Station (ft)', y_axis=False, x_fmt='0', grid=False)


def pier_pedestal(area, ctx):
    f = Figure(area, 'Pier pedestal top along beam line 1')
    r = ctx['pier_ped_rows'][1]
    B = lambda c: f"'Bearing Seats'!${c}${r}"
    pts = [(B('C'), B('G')), (B('D'), B('H')), (B('E'), B('I')), (B('F'), B('J'))]
    lo = f.helper('Window bottom (ft)', f'MIN({B("G")},{B("H")},{B("I")},{B("J")})-0.006')
    hi = f.helper('Window top (ft)', f'MAX({B("G")},{B("H")},{B("I")},{B("J")})+0.004')
    f.line('Pedestal outline', [(B('C'), lo)] + pts + [(B('F'), lo), (B('C'), lo)], color=CONCRETE, width=1.0)
    f.line('Pedestal top', pts, color=RED, width=1.4, marker='circle', size=4)
    f.line('CL pier', [('Pier_Sta', lo), ('Pier_Sta', hi)], color=BLUE, width=0.8, dash='dash')
    f.label('CL pier', 'Pier_Sta', hi, pos='b', color=BLUE, size=7)
    for (x, z), name in zip(pts, ('A/C', 'E', 'F', 'B/D')):
        f.label(live(name + '\n', NUM(z, '0.0000')), x, z, pos='b', size=7.5)
    return f.chart(height=8.0, x_title='Station (ft)', y_title='Elevation (ft)', x_fmt='0.0', y_fmt='0.000')


def cap_section(area, ctx, tag):
    name = {'Pier': 'Pier', 'EB1': 'End Bent 1', 'EB2': 'End Bent 2'}[tag]
    f = Figure(area, f'{name}: transverse section (elevations relative to the cap top)')
    last = len(ctx['spans'])
    ref = f.helper('Reference: cap top at the cap center (ft)', f'{tag}_Top')
    rel = lambda z: f'({z})-{ref}'
    cap_top = lambda y: f'({tag}_Top+{tag}_Slope*(({y})-{tag}_Ctr))'
    xl, xr = f'{tag}_Ctr-{tag}_Len/2', f'{tag}_Ctr+{tag}_Len/2'
    corners = [(xl, cap_top(xl)), (xr, cap_top(xr)), (xr, f'{cap_top(xr)}-{tag}_Depth'), (xl, f'{cap_top(xl)}-{tag}_Depth'),
               (xl, cap_top(xl))]
    under = f.helper('Lowest cap underside (relative, ft)', f'MIN({rel(cap_top(xl))},{rel(cap_top(xr))})-{tag}_Depth')
    if tag != 'Pier':
        W = lambda key: f"'Substructure'!${'F' if tag == 'EB1' else 'G'}${ctx['wall_rows'][key]}"
        f.line('Backwall (behind the cap, at the FFBW)', [(W('yl'), rel(W('ctl'))), (W('yr'), rel(W('ctr'))), (W('yr'), rel(W('btr'))),
                                                         (f'MAX({W("yl")},MIN(0,{W("yr")}))', rel(W('bt0'))), (W('yl'), rel(W('btl'))),
                                                         (W('yl'), rel(W('ctl')))], color=GREY, width=1.0, dash='dash')
        t_c, w_l = f'{tag}_CkT/12', f'{tag}_LugW/12'
        left = [(f'{W("yl")}-{t_c}', rel(W('lbl'))), (f'{W("yl")}-{t_c}+{w_l}', rel(W('lbl'))), (f'{W("yl")}-{t_c}+{w_l}', rel(W('zbl'))),
                (W('yl'), rel(W('zbl'))), (W('yl'), rel(W('ckl'))), (f'{W("yl")}-{t_c}', rel(W('ckl'))), (f'{W("yl")}-{t_c}', rel(W('lbl')))]
        right = [(f'{W("yr")}+{t_c}-{w_l}', rel(W('lbr'))), (f'{W("yr")}+{t_c}', rel(W('lbr'))), (f'{W("yr")}+{t_c}', rel(W('ckr'))),
                 (W('yr'), rel(W('ckr'))), (W('yr'), rel(W('zbr'))), (f'{W("yr")}+{t_c}-{w_l}', rel(W('zbr'))),
                 (f'{W("yr")}+{t_c}-{w_l}', rel(W('lbr')))]
        f.line('Cheek webs with bottom lugs', _gaps([left, right]), color=CHEEK, width=1.25)
        f.label(live('backwall top ', NUM(W('bt0'), '0.000'), '\n(behind, at FFBW)'), 0, rel(W('bt0')), pos='t', color=GREY, size=7)
        for side, y in (('ckl', f'{W("yl")}-{t_c}/2'), ('ckr', f'{W("yr")}+{t_c}/2')):
            f.label(live('cheek top\n', NUM(W(side), '0.000')), y, rel(W(side)), pos='t', color=CHEEK, size=7)
    f.line('Cap', [(x, rel(z)) for x, z in corners], color=INK, width=1.5)
    count = int(ctx['cfg']['pier']['pile_count'] if tag == 'Pier'
                else ctx['cfg']['end_bents']['start' if tag == 'EB1' else 'end']['pile_count'])
    width = 'Pier_PileD/12' if tag == 'Pier' else f'{tag}_PileW/12'
    piles = []
    for j in range(1, count + 1):
        pr = ctx['pile_rows'][(tag, j)]
        x, cut, bot = f"'Substructure'!$D${pr}", f"'Substructure'!$J${pr}", f"'Substructure'!$H${pr}"
        piles.append([(f'{x}-{width}/2', rel(cut)), (f'{x}+{width}/2', rel(cut)), (f'{x}+{width}/2', f'{rel(bot)}-2.2'),
                      (f'{x}-{width}/2', f'{rel(bot)}-2.2'), (f'{x}-{width}/2', rel(cut))])
    f.line('Piles (shown 2.2 ft below the cap)', _gaps(piles), color=STEEL, width=1.0)
    pr1 = ctx['pile_rows'][(tag, 1)]
    f.label(live('pile cutoff\n', NUM(f"'Substructure'!$J${pr1}", '0.0000')), f"'Substructure'!$D${pr1}", rel(f"'Substructure'!$J${pr1}"),
            pos='t', color=RED, size=7)
    seat_span = last if tag == 'EB2' else 1
    peds, pads = [], []
    pad_w = 'Pier_PadW' if tag == 'Pier' else f'{tag}_PadW'
    pad_t = 'Pier_Pad' if tag == 'Pier' else f'{tag}_Pad'
    ped_w = 'Pier_PedW' if tag == 'Pier' else f'{tag}_PedW'
    for k in range(1, ctx['model'].count + 1):
        s = ctx['seat_rows'][(tag, seat_span, k)]
        v = f"'Bearing Seats'!$V${s}"
        a, b = f'Bm{k}_Off-{ped_w}/24', f'Bm{k}_Off+{ped_w}/24'
        peds.append([(a, rel(cap_top(a))), (b, rel(cap_top(b))), (b, rel(v)), (a, rel(v)), (a, rel(cap_top(a)))])
        pa, pb = f'Bm{k}_Off-{pad_w}/24', f'Bm{k}_Off+{pad_w}/24'
        pads.append([(pa, rel(v)), (pb, rel(v)), (pb, f'{rel(v)}+{pad_t}/12'), (pa, f'{rel(v)}+{pad_t}/12'), (pa, rel(v))])
        f.label(live('pad top\n', NUM(f"'Bearing Seats'!$S${s}", '0.000')), f'Bm{k}_Off', f'{rel(v)}+{pad_t}/12', pos='t', size=7)
    f.line('Pedestals', _gaps(peds), color=CONCRETE, width=1.0)
    f.line('Bearing pads', _gaps(pads), color=STEEL, width=1.0)
    f.label(live('cap top ', NUM(f'{tag}_Top', '0.0000')), f'{tag}_Ctr', rel(cap_top(f'{tag}_Ctr')), pos='b', color=BLUE, size=7.5)
    f.label(live('cap underside\n', NUM(f'{tag}_Under', '0.0000')), f'{tag}_Ctr', f'{rel(cap_top(f"{tag}_Ctr"))}-{tag}_Depth',
            pos='t', size=7)
    y1 = f.helper('Cap-length dimension line (relative, ft)', f'{under}-2.7')
    f.line('Cap length', [(xl, y1), (xr, y1)], width=0.75, arrows=True)
    f.label(live('cap length ', NUM(f'{tag}_Len*12', '0.0'), ' in'), f'{tag}_Ctr', y1, pos='t', size=7)
    bottom = y1
    tops = [f'{ped[2][1]}+{pad_t}/12' for ped in peds]      # pedestal top (relative) plus the pad
    if tag != 'Pier':
        y2 = f.helper('Out-to-out dimension line (relative, ft)', f'{under}-3.2')
        f.line('Out-to-out', [(f'{xl}-{tag}_CkT/12', y2), (f'{xr}+{tag}_CkT/12', y2)], width=0.75, arrows=True)
        f.label(live('out-to-out with cheeks ', NUM(f'{tag}_OutIn', '0.0'), ' in'), f'{tag}_Ctr', y2, pos='b', size=7)
        bottom = f'{y2}-0.5'
        tops += [rel(W('ckl')), rel(W('ckr'))]
    upper = f.helper('Window top: tallest pad or cheek plus labels (relative, ft)', 'MAX(' + ','.join(tops) + ')+1.0')
    f.bounds(f'{xl}-1.3', f'{xr}+1.3', f'{bottom}-0.3', upper)
    return f.chart(height=10.0, x_title='Offset from PGL (ft), looking up-station', y_axis=False, x_fmt='0', grid=False)


def stirrup_zones(area, ctx):
    f = Figure(area, 'Bars 5K zones, beam 1-1 (half beam)')
    z, b = ctx['zone_rows'][1], ctx['bar_rows'][(1, 1)]
    D = lambda c, r=z: f"'Beam Data'!${c}${r}"
    half = f.helper('Half beam length (in)', D('D', b), '0.00')
    zones = [('S*', D('C'), D('D')), ('Sj', D('E'), D('F')), ('S1', D('G'), D('H')), ('S2', D('I'), D('J')),
             ('S3', D('K'), D('L')), ('S4', D('F', b), D('M'))]
    x = f.helper('First bar from beam end (in)', D('N'), '0.00')
    f.line('Beam (half length)', [(0, 0), (half, 0)], color=INK, width=2.0)
    f.line('Beam end', [(0, -0.3), (0, 0.6)], color=INK, width=1.4)
    f.line('Midspan', [(half, -0.3), (half, 0.6)], color=BLUE, width=1.0, dash='dash')
    f.label('beam end', 0, 0.6, pos='t', size=7)
    f.label('midspan', half, 0.6, pos='t', color=BLUE, size=7)
    for k, ((zone, n, v), color) in enumerate(zip(zones, ZONES)):
        end = f.helper(f'Zone {zone} end (in)', f'{x}+{n}*{v}', '0.00')
        f.line(f'Zone {zone}', [(x, 0), (end, 0)], color=color, width=9.0)
        f.label(live(zone + ': ', NUM(n, '0'), ' @ ', NUM(v, 'General'), ' in'), f'({x}+{end})/2',
                0.22 if k % 2 == 0 else -0.22, pos='t' if k % 2 == 0 else 'b', color=color, size=7.5)
        x = end
    f.bounds(-12, f'{half}+12', -1.0, 1.1)
    return f.chart(height=6.0, x_title='Distance from beam end (in)', y_axis=False, x_fmt='0', grid=False)


def clearance(area, ctx):
    f = Figure(area, 'Vertical clearance along the beams')
    a, b = ctx['clear_segs']
    E, F = f"'Clearance'!$E${a}:$E${b}", f"'Clearance'!$F${a}:$F${b}"
    col = lambda c: f"'Clearance'!${c}${a}:${c}${b}"
    cov0 = f.helper('Roadway coverage start (ft)', f'MAX(Prof_Sta0,MIN({E}))', '0.0000')
    cov1 = f.helper('Roadway coverage end (ft)', f'MIN(Prof_Sta1,MAX({F}))', '0.0000')

    def road(X):
        m = f'MATCH({X},{E},1)'
        d = f'(({X})-INDEX({E},{m}))'
        return f'(INDEX({col("G")},{m})+INDEX({col("H")},{m})*{d}+INDEX({col("I")},{m})*{d}^2)'

    raise_min = 'MIN(' + ','.join(f'Bm{k}_Raise' for k in range(1, ctx['model'].count + 1)) + ')'
    N = 30
    tops = []
    for i in range(1, len(ctx['spans']) + 1):
        lo = f.helper(f'Span {i} plotted from (ft)', f'MAX(Sp{i}_Bm0,{cov0})', '0.0000')
        hi = f.helper(f'Span {i} plotted to (ft)', f'MIN(Sp{i}_Bm1,{cov1})', '0.0000')
        pts = [(f'{lo}+{k}/{N}*({hi}-{lo})', lambda X, i=i: f'{BOTTOM(i, f"(({X})-Sp{i}_BrgA)/Sp{i}_Lb", raise_min)}-{road(X)}')
               for k in range(N + 1)]
        h = f.line(f'Clearance under Span {i} beams', pts, color=INK, width=1.5)
        x, _ = h.pt(N // 4)
        f.label(f'Span {i} beams', x, f'MAX({h.y})', pos='t', size=7.5)
        tops.append(f'MAX({h.y})')
    f.line('Required clearance', [('Prof_Sta0', 'Rd_Req'), ('Prof_Sta1', 'Rd_Req')], color=RED, width=1.0, dash='dash')
    f.label(live('required ', NUM('Rd_Req', '0.00'), ' ft'), 'Prof_Sta0+0.1*Prof_L', 'Rd_Req', pos='b', color=RED, size=7.5)
    f.line('Roadway coverage limits', _gaps([((cov0, 0), (cov0, 'Rd_Req*0.5')), ((cov1, 0), (cov1, 'Rd_Req*0.5'))]),
           color=GREEN, width=0.75, dash='sysDot')
    f.label('XML coverage ends', cov1, 'Rd_Req*0.5', pos='t', color=GREEN, size=7)
    f.line('Minimum leader', [('Clr_Sta', 'Clr_Min'), ('Clr_Sta', 'Rd_Req-2')], color=RED, width=0.6)
    f.line('Minimum point', [('Clr_Sta', 'Clr_Min')], width=0, marker='circle', color=RED, size=6)
    f.label(live('min ', NUM('Clr_Min', '0.000'), ' ft @ ', (STA('Clr_Sta'),)), 'Clr_Sta', 'Rd_Req-2', pos='b', color=RED, size=8)
    f.bounds('Prof_Sta0-3', 'Prof_Sta1+3', 0, 'MAX(' + ','.join(tops) + ')+3')
    return f.chart(height=9.0, x_title='Station (ft)', y_title='Clearance = beam bottom - roadway (ft)', x_fmt='0', y_fmt='0')
