"""Build the B4 geometry CAD-table workbook.

Geometry authority: the Colab geometry notebook (CONFIG + BridgeModel). Every
calculation tab is written as live Excel formulas over named input cells, so
the workbook recalculates when an input changes. The QC tab compares each
result with the notebook value captured at build time.
"""
from __future__ import annotations

import datetime as _dt
import hashlib
import json
import math
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from . import figures, sources

GENERATOR_VERSION = '1.1.0'
OUTPUT = sources.ROOT / 'geometry_tables' / 'B4_Geometry_CAD_Tables.xlsx'
FGE_INCREMENTS = 10
PACKAGE = Path(__file__).resolve().parent

FONT = 'Arial'
F_BASE = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color='0000FF')
F_LINK = Font(name=FONT, size=10, color='008000')
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=9, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True, color='1F2A36')
F_H2 = Font(name=FONT, size=11, bold=True, color='1F2A36')
F_NOTE = Font(name=FONT, size=9, italic=True, color='5A6570')
F_WARN = Font(name=FONT, size=9, bold=True, color='9C5700')
FILL_INPUT = PatternFill('solid', fgColor='FFF2CC')
FILL_HEAD = PatternFill('solid', fgColor='D9E1F2')
FILL_SECTION = PatternFill('solid', fgColor='EEF1F4')
FILL_CAD = PatternFill('solid', fgColor='F2F2F2')
FILL_PASS = PatternFill('solid', fgColor='C6EFCE')
FILL_FAIL = PatternFill('solid', fgColor='FFC7CE')
FILL_WARN = PatternFill('solid', fgColor='FFEB9C')


def cf_fill(rgb: str) -> PatternFill:
    """Conditional-format fill: Excel paints a differential solid fill from bgColor, so set both colors."""
    return PatternFill(fill_type='solid', start_color=rgb, end_color=rgb)


CF_PASS, CF_FAIL, CF_WARN = cf_fill('C6EFCE'), cf_fill('FFC7CE'), cf_fill('FFEB9C')
THIN = Side(style='thin', color='9AA5B1')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
WRAP = Alignment(vertical='top', wrap_text=True)

FMT_STA = '0.0000'
FMT_EL = '0.0000'
FMT_FT = '0.0000'
FMT_IN = '0.000'
FMT_GRADE = '0.000000'


def col(n: int) -> str:
    return get_column_letter(n)


# ---------------------------------------------------------------- formulas --
def PGL(x: str) -> str:
    """Notebook BridgeModel.pgl as an Excel expression of station x."""
    x = f'({x})'
    return (f'IF(Prof_Mode="parabolic",Prof_Z0+Prof_G1*({x}-Prof_Sta0)+Prof_Kq*({x}-Prof_Sta0)^2,'
            f'IF({x}<=Prof_PISta,Prof_Z0+Prof_G1*({x}-Prof_Sta0),Prof_PIZ+Prof_G2*({x}-Prof_PISta)))')


def DELTA(y: str) -> str:
    """Notebook BridgeModel.deck_delta (crowned or uniform) of offset y."""
    return f'IF(Deck_Mode="uniform",Deck_Uni*({y}),Deck_Crown*ABS({y}))'


def BOTTOM(i: int, t: str, raise_in: str) -> str:
    """Beam soffit: chord between bearings - common drop + parabolic camber + raise."""
    return (f'(Sp{i}_ZA+(Sp{i}_ZB-Sp{i}_ZA)*({t})-Sup_Drop+Sp{i}_C/3*({t})*(1-({t}))+({raise_in})/12)')


def ft_in(value: float) -> str:
    feet = math.floor(value + 1e-9)
    return f"{feet}'-{(value - feet) * 12:.2f}\""


class Book:
    """Thin openpyxl wrapper: styles, defined names and cell bookkeeping."""

    def __init__(self):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.names: dict[str, str] = {}
        self.refs: dict = {}
        self.images: list = []
        self.inputs: list = []      # (sheet, ref, value written at build) for the QC input snapshot

    def sheet(self, title: str, tab_color: str | None = None):
        ws = self.wb.create_sheet(title)
        ws.sheet_view.showGridLines = False
        if tab_color:
            ws.sheet_properties.tabColor = tab_color
        return ws

    def define(self, name: str, ws, ref: str):
        letters = ''.join(ch for ch in ref if ch.isalpha())
        digits = ''.join(ch for ch in ref if ch.isdigit())
        target = f"'{ws.title}'!${letters}${digits}"
        if name in self.names:
            raise ValueError(f'Duplicate defined name {name}')
        self.names[name] = target
        self.wb.defined_names[name] = DefinedName(name, attr_text=target)

    def put(self, ws, ref, value, font=F_BASE, fill=None, fmt=None, align=None, border=None, bold=False, track=True):
        cell = ws[ref]
        cell.value = value
        if fill is FILL_INPUT and track and not (isinstance(value, str) and value.startswith('=')):
            self.inputs.append((ws.title, ref, value))
        cell.font = F_BOLD if bold else font
        if fill:
            cell.fill = fill
        if fmt:
            cell.number_format = fmt
        if align:
            cell.alignment = align
        if border:
            cell.border = border
        return cell

    @staticmethod
    def xref(ws, ref) -> str:
        return f"'{ws.title}'!{ref}"


def _kind_font(kind):
    return {'input': F_INPUT, 'link': F_LINK}.get(kind, F_BASE)


def title_block(book: Book, ws, title: str, subtitle: str, width: int = 12):
    book.put(ws, 'A1', title, font=F_TITLE)
    book.put(ws, 'A2', subtitle, font=F_NOTE)
    labels = [('Project', '=""&TB_Project'), ('Job No.', '=""&TB_JobNo'), ('Calc. by / date', '=TB_CalcBy&"  "&TB_CalcDate'),
              ('Check. by / date', '=TB_CheckBy&"  "&TB_CheckDate')]
    c = width
    for k, (label, formula) in enumerate(labels):
        book.put(ws, f'{col(c)}{1 + k}', label, font=F_HEAD)
        book.put(ws, f'{col(c + 1)}{1 + k}', formula, font=F_LINK)
    return 6


def section(book: Book, ws, row: int, text: str, span: int = 10):
    book.put(ws, f'A{row}', text, font=F_H2, fill=FILL_SECTION)
    for c in range(2, span + 1):
        ws.cell(row=row, column=c).fill = FILL_SECTION
    return row + 1


def header_row(book: Book, ws, row: int, start_col: int, headers, height=30):
    for k, h in enumerate(headers):
        book.put(ws, f'{col(start_col + k)}{row}', h, font=F_HEAD, fill=FILL_HEAD, align=CENTER, border=BOX)
    ws.row_dimensions[row].height = height
    return row + 1


def check_formatting(ws, rng: str):
    ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"PASS"'], fill=CF_PASS))
    ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"FAIL"'], fill=CF_FAIL))
    ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"CHECK"'], fill=CF_WARN))


# =============================================================== Inputs ==
def _value(v):
    if v is None:
        return None
    if isinstance(v, (bool,)):
        return str(v)
    try:
        import numpy as np
        if isinstance(v, np.generic):
            return v.item()
    except ImportError:
        pass
    return v


def param_table(book: Book, ws, row: int, items, name_col=True):
    """items: dicts(label, sym, value, unit, name, src, kind, fmt)."""
    row = header_row(book, ws, row, 2, ['Parameter', 'Symbol', 'Value', 'Unit', 'Workbook name', 'Source / note'], 18)
    for it in items:
        kind = it.get('kind', 'input')
        book.put(ws, f'B{row}', it['label'], align=WRAP)
        book.put(ws, f'C{row}', it.get('sym', ''))
        v = _value(it.get('value'))
        cell = book.put(ws, f'D{row}', v, font=_kind_font(kind), fmt=it.get('fmt'),
                        fill=FILL_INPUT if kind == 'input' else None, border=BOX,
                        track=not str(it.get('name') or '').startswith('TB_'))   # title block is filled in by hand
        is_text = isinstance(v, str) and not v.startswith('=')
        cell.alignment = Alignment(horizontal='left' if is_text else 'right', wrap_text=is_text)
        if is_text and len(v) > 18:
            ws.row_dimensions[row].height = 26
        book.put(ws, f'E{row}', it.get('unit', ''))
        if it.get('name'):
            book.define(it['name'], ws, f'D{row}')
            book.put(ws, f'F{row}', it['name'], font=F_NOTE)
        book.put(ws, f'G{row}', it.get('src', ''), font=F_NOTE, align=WRAP)
        row += 1
    return row + 1


def build_inputs(book: Book, ctx):
    cfg, m = ctx['cfg'], ctx['model']
    ws = book.sheet('Inputs', '2B6CB0')
    for c, w in zip('ABCDEFG', (3, 46, 22, 16, 16, 18, 78)):
        ws.column_dimensions[c].width = w
    book.put(ws, 'A1', 'Inputs', font=F_TITLE)
    book.put(ws, 'A2', 'Blue values on yellow are inputs. Every geometry input is generated from the Colab geometry '
             'notebook CONFIG (the geometry authority); beam-design inputs come from the saved Mathcad runs. '
             'Edit the notebook and rebuild rather than editing these cells: the tables and Method-tab charts follow any '
             'edit, and the banner on every Method tab and QC report that inputs no longer match the notebook build.', font=F_NOTE)
    ws.merge_cells('A2:G2'); ws['A2'].alignment = WRAP; ws.row_dimensions[2].height = 28
    row = 4
    nb = 'Notebook CONFIG'
    row = section(book, ws, row, 'Title block (fill in)', 7)
    row = param_table(book, ws, row, [
        dict(label='Project', value='D2 I-95 Pedestrian Bridge', name='TB_Project', src='Project journal name; edit as needed'),
        dict(label='Structure', value='B4 pedestrian bridge - 2 x FIB-36, two simple spans', name='TB_Structure',
             src='FreeCAD model / Mathcad beam type; edit as needed'),
        dict(label='Job No.', value=None, name='TB_JobNo', src='Not in project sources - enter'),
        dict(label='Calculated by', value=None, name='TB_CalcBy', src='Enter'),
        dict(label='Date (calculated)', value=None, name='TB_CalcDate', src='Enter'),
        dict(label='Checked by', value=None, name='TB_CheckBy', src='Enter'),
        dict(label='Date (checked)', value=None, name='TB_CheckDate', src='Enter'),
    ])
    p, d, b = cfg['profile'], cfg['deck'], cfg['beams']
    row = section(book, ws, row, "Profile - CONFIG['profile']", 7)
    row = param_table(book, ws, row, [
        dict(label='Profile mode (brokeback or parabolic)', sym='mode', value=p['mode'], name='Prof_Mode', src=f"{nb} profile.mode"),
        dict(label='Profile start station (FFBW End Bent 1)', sym='s0', value=p['start_sta_ft'], unit='ft', name='Prof_Sta0', fmt=FMT_STA, src=f'{nb} profile.start_sta_ft'),
        dict(label='Profile end station (FFBW End Bent 2)', sym='s1', value=p['end_sta_ft'], unit='ft', name='Prof_Sta1', fmt=FMT_STA, src=f'{nb} profile.end_sta_ft'),
        dict(label='Start elevation', sym='z0', value=p['start_elev_ft'], unit='ft', name='Prof_Z0', fmt=FMT_EL, src=f'{nb} profile.start_elev_ft'),
        dict(label='End elevation', sym='z1', value=p['end_elev_ft'], unit='ft', name='Prof_Z1', fmt=FMT_EL, src=f'{nb} profile.end_elev_ft'),
        dict(label='Tangent intersection station (PI)', sym='sPI', value=p['intersection_sta_ft'], unit='ft', name='Prof_PISta', fmt=FMT_STA, src=f'{nb} profile.intersection_sta_ft'),
        dict(label='Nominal tangent grade (brokeback)', sym='g', value=p['nominal_grade'], unit='ft/ft', name='Prof_Gnom', fmt=FMT_GRADE, src=f'{nb} profile.nominal_grade'),
        dict(label='Grade change (parabolic mode only)', sym='dg', value=p['parabolic_grade_change'], unit='ft/ft', name='Prof_dG', fmt=FMT_GRADE, src=f'{nb} profile.parabolic_grade_change'),
    ])
    row = section(book, ws, row, "Deck - CONFIG['deck']", 7)
    row = param_table(book, ws, row, [
        dict(label='Deck mode (crowned or uniform)', sym='mode', value=d['mode'], name='Deck_Mode', src=f'{nb} deck.mode. Transition mode is not implemented in this workbook (QC fails).'),
        dict(label='Deck width', sym='W', value=d['width_ft'], unit='ft', name='Deck_W', fmt=FMT_FT, src=f'{nb} deck.width_ft (13 ft 7 in from DXF *D5380)'),
        dict(label='Deck center offset from PGL', sym='yc', value=d['center_offset_ft'], unit='ft', name='Deck_C', fmt=FMT_FT, src=f'{nb} deck.center_offset_ft'),
        dict(label='Deck thickness', sym='td', value=d['thickness_ft'], unit='ft', name='Deck_T', fmt=FMT_FT, src=f'{nb} deck.thickness_ft'),
        dict(label='Crown slope (each side, away from PGL)', sym='sc', value=d['crown_slope'], unit='ft/ft', name='Deck_Crown', fmt=FMT_GRADE, src=f'{nb} deck.crown_slope'),
        dict(label='Uniform cross slope (uniform mode only)', sym='su', value=d['uniform_slope'], unit='ft/ft', name='Deck_Uni', fmt=FMT_GRADE, src=f'{nb} deck.uniform_slope'),
    ])
    gaps = b['gaps_ft'] if b['gaps_ft'] is not None else None
    items = [
        dict(label='Beam type', value=ctx['mathcad'][1]['beam_type'].replace('FIB', 'FIB-'), name='Bm_Type', src='Mathcad BeamType (notebook models the FIB-36 outline)'),
        dict(label='Number of beams', sym='N', value=b['count'], name='Bm_N', src=f'{nb} beams.count (table layout is generated for this count)'),
        dict(label='Beam spacing', sym='S', value=b['spacing_ft'], unit='ft', name='Bm_S', fmt=FMT_FT, src=f'{nb} beams.spacing_ft'),
        dict(label='First beam offset (blank = centered on PGL)', sym='y1', value=b['first_offset_ft'], unit='ft', name='Bm_First', fmt=FMT_FT, src=f'{nb} beams.first_offset_ft'),
        dict(label='Beam depth', sym='db', value=b['depth_ft'], unit='ft', name='Bm_D', fmt=FMT_FT, src=f'{nb} beams.depth_ft'),
        dict(label='Top flange width', sym='bf', value=b['top_flange_width_in'], unit='in', name='Bm_Bf', src=f'{nb} beams.top_flange_width_in'),
        dict(label='Minimum haunch (at any flange point)', sym='hmin', value=b['minimum_haunch_in'], unit='in', name='Bm_Hmin', src=f'{nb} beams.minimum_haunch_in'),
    ]
    raises = b['raise_in'] if isinstance(b['raise_in'], (list, tuple)) else [b['raise_in']] * b['count']
    for k in range(b['count']):
        items.append(dict(label=f'Beam {k+1} vertical raise', sym=f'r{k+1}', value=raises[k], unit='in', name=f'Bm{k+1}_Raise', src=f'{nb} beams.raise_in'))
    for k in range(b['count'] - 1):
        if gaps is None:
            items.append(dict(label=f'Gap beam {k+1} to {k+2}', value='=Bm_S', unit='ft', name=f'Bm_Gap{k+1}', fmt=FMT_FT, kind='calc', src='Equal spacing (CONFIG gaps_ft = None)'))
        else:
            items.append(dict(label=f'Gap beam {k+1} to {k+2}', value=gaps[k], unit='ft', name=f'Bm_Gap{k+1}', fmt=FMT_FT, src=f'{nb} beams.gaps_ft'))
    row = section(book, ws, row, "Beams - CONFIG['beams']", 7)
    row = param_table(book, ws, row, items)
    row = section(book, ws, row, "Spans and cambers - CONFIG['spans'] (camber = Mathcad saved NetCamber)", 7)
    items = []
    for i, span in enumerate(ctx['spans'], start=1):
        items += [
            dict(label=f'Span {i} start station (datum)', value=span['start_ft'], unit='ft', name=f'Sp{i}_Sta0', fmt=FMT_STA, src=f'{nb} spans[{i-1}].start_ft'),
            dict(label=f'Span {i} end station (datum)', value=span['end_ft'], unit='ft', name=f'Sp{i}_Sta1', fmt=FMT_STA, src=f'{nb} spans[{i-1}].end_ft'),
            dict(label=f'Span {i} net camber at 120 days', sym='C', value=span['camber_in'], unit='in', name=f'Sp{i}_C', fmt='0.000000',
                 src=f"{nb} spans[{i-1}].camber_in; basis {cfg['camber_basis'][i-1]['variable']} = {cfg['camber_basis'][i-1]['source_equation']}"),
        ]
    row = param_table(book, ws, row, items)
    ph = cfg['physical_geometry']
    row = section(book, ws, row, "Physical beam / bearing geometry - CONFIG['physical_geometry']", 7)
    row = param_table(book, ws, row, [
        dict(label='Beam end clearance to FFBW', value=ph['beam_end_clear_in'], unit='in', name='Phys_EndClr', src=f'{nb} physical_geometry.beam_end_clear_in'),
        dict(label='Pier beam gap (total, beam end to beam end)', value=ph['pier_beam_gap_in'], unit='in', name='Phys_PierGap', src=f'{nb} physical_geometry.pier_beam_gap_in'),
        dict(label='Bearing setback (beam end to CL bearing)', value=ph['bearing_setback_in'], unit='in', name='Phys_Setback', src=f'{nb} physical_geometry.bearing_setback_in'),
        dict(label='Deck end clearance to FFBW', value=ph['deck_end_clear_in'], unit='in', name='Phys_DeckClr', src=f'{nb} physical_geometry.deck_end_clear_in'),
        dict(label='Trial pile length (stored)', value=ph['trial_pile_length_ft'], unit='ft', name='Phys_TrialPile', fmt=FMT_FT, src=f'{nb} physical_geometry.trial_pile_length_ft (trial only)'),
    ])
    pr = cfg['pier']
    row = section(book, ws, row, "Pier - CONFIG['pier']", 7)
    row = param_table(book, ws, row, [
        dict(label='Pier station (CL pier)', value=pr['station_ft'], unit='ft', name='Pier_Sta', fmt=FMT_STA, src=f'{nb} pier.station_ft'),
        dict(label='Bearing pad thickness', value=pr['pad_height_in'], unit='in', name='Pier_Pad', fmt=FMT_IN, src=f'{nb} pier.pad_height_in'),
        dict(label='Bearing pad length (longitudinal)', value=pr['pad_length_in'], unit='in', name='Pier_PadL', src=f'{nb} pier.pad_length_in'),
        dict(label='Bearing pad width (transverse)', value=pr['pad_width_in'], unit='in', name='Pier_PadW', src=f'{nb} pier.pad_width_in'),
        dict(label='Cap depth', value=pr['depth_ft'], unit='ft', name='Pier_Depth', fmt=FMT_FT, src=f'{nb} pier.depth_ft'),
        dict(label='Cap width (longitudinal)', value=pr['width_ft'], unit='ft', name='Pier_Width', fmt=FMT_FT, src=f'{nb} pier.width_ft'),
        dict(label='Fixed cap length (blank = derive from piles)', value=pr['length_ft'], unit='ft', name='Pier_LenFix', fmt=FMT_FT, src=f'{nb} pier.length_ft'),
        dict(label='Pile count', value=pr['pile_count'], name='Pier_NPile', src=f'{nb} pier.pile_count'),
        dict(label='Pile diameter', value=pr['pile_diameter_in'], unit='in', name='Pier_PileD', src=f'{nb} pier.pile_diameter_in (hollow pipe)'),
        dict(label='Pipe wall thickness', value=pr['pile_wall_in'], unit='in', name='Pier_PileWall', src=f'{nb} pier.pile_wall_in'),
        dict(label='Pile spacing in diameters', value=pr['pile_spacing_diameters'], name='Pier_PileSpD', src=f'{nb} pier.pile_spacing_diameters'),
        dict(label='Clear distance, outer pile face to cap end', value=pr['pile_end_clear_in'], unit='in', name='Pier_PileClr', src=f'{nb} pier.pile_end_clear_in'),
        dict(label='Pile embedment above cap underside', value=pr['pile_embedment_in'], unit='in', name='Pier_Emb', src=f'{nb} pier.pile_embedment_in'),
        dict(label='Minimum pedestal height at CL bearing', value=pr['minimum_pedestal_cl_in'], unit='in', name='Pier_Reveal', src=f'{nb} pier.minimum_pedestal_cl_in'),
        dict(label='Fixed cap top (blank = pedestal rule)', value=pr['top_center_ft'], unit='ft', name='Pier_TopFix', fmt=FMT_EL, src=f'{nb} pier.top_center_ft'),
        dict(label='Cap center offset', value=pr['center_offset_ft'], unit='ft', name='Pier_Ctr', fmt=FMT_FT, src=f'{nb} pier.center_offset_ft'),
        dict(label='Cap cross slope', value=pr['cross_slope'], unit='ft/ft', name='Pier_Slope', fmt=FMT_GRADE, src=f'{nb} pier.cross_slope'),
        dict(label='Pedestal width (transverse)', value=pr['pedestal_width_in'], unit='in', name='Pier_PedW', src=f'{nb} pier.pedestal_width_in'),
        dict(label='Pedestal length (longitudinal)', value=pr['pedestal_length_in'], unit='in', name='Pier_PedL', src=f'{nb} pier.pedestal_length_in'),
        dict(label='Pedestal top transverse chamfer', value=pr['pedestal_chamfer_in'], unit='in', name='Pier_Chamf', src=f'{nb} pier.pedestal_chamfer_in'),
    ])
    for key, tag, label in (('start', 'EB1', 'End Bent 1'), ('end', 'EB2', 'End Bent 2')):
        e = cfg['end_bents'][key]
        row = section(book, ws, row, f"{label} - CONFIG['end_bents']['{key}']", 7)
        src = f"{nb} end_bents.{key}"
        row = param_table(book, ws, row, [
            dict(label='Cap width (longitudinal)', value=e['width_ft'], unit='ft', name=f'{tag}_Width', fmt=FMT_FT, src=f'{src}.width_ft'),
            dict(label='Cap depth', value=e['depth_ft'], unit='ft', name=f'{tag}_Depth', fmt=FMT_FT, src=f'{src}.depth_ft'),
            dict(label='Fixed cap length (blank = derive from piles)', value=e['length_ft'], unit='ft', name=f'{tag}_LenFix', fmt=FMT_FT,
                 src=f'{src}.length_ft (the 235-in D-026 cap body is derived from the pile layout)'),
            dict(label='Pile count', value=e['pile_count'], name=f'{tag}_NPile', src=f'{src}.pile_count'),
            dict(label='Square pile width', value=e['pile_width_in'], unit='in', name=f'{tag}_PileW', src=f'{src}.pile_width_in'),
            dict(label='Pile spacing', value=e['pile_spacing_in'], unit='in', name=f'{tag}_PileSp', src=f'{src}.pile_spacing_in'),
            dict(label='Clear distance, outer pile face to cap end', value=e['pile_end_clear_in'], unit='in', name=f'{tag}_PileClr', src=f'{src}.pile_end_clear_in'),
            dict(label='Pile embedment above cap underside', value=e['pile_embedment_in'], unit='in', name=f'{tag}_Emb', src=f'{src}.pile_embedment_in'),
            dict(label='Bearing pad thickness', value=e['pad_height_in'], unit='in', name=f'{tag}_Pad', fmt=FMT_IN, src=f'{src}.pad_height_in'),
            dict(label='Bearing pad length (longitudinal)', value=e['pad_length_in'], unit='in', name=f'{tag}_PadL', src=f'{src}.pad_length_in'),
            dict(label='Bearing pad width (transverse)', value=e['pad_width_in'], unit='in', name=f'{tag}_PadW', src=f'{src}.pad_width_in'),
            dict(label='Minimum pedestal height at CL bearing', value=e['minimum_pedestal_cl_in'], unit='in', name=f'{tag}_Reveal', src=f'{src}.minimum_pedestal_cl_in'),
            dict(label='Fixed cap top (blank = pedestal rule)', value=e['top_center_ft'], unit='ft', name=f'{tag}_TopFix', fmt=FMT_EL, src=f'{src}.top_center_ft'),
            dict(label='Cap center offset', value=e['center_offset_ft'], unit='ft', name=f'{tag}_Ctr', fmt=FMT_FT, src=f'{src}.center_offset_ft'),
            dict(label='Cap cross slope', value=e['cross_slope'], unit='ft/ft', name=f'{tag}_Slope', fmt=FMT_GRADE, src=f'{src}.cross_slope'),
            dict(label='Pedestal width (transverse)', value=e['pedestal_width_in'], unit='in', name=f'{tag}_PedW', src=f'{src}.pedestal_width_in'),
            dict(label='Pedestal length (longitudinal, from FFBW)', value=e['pedestal_length_in'], unit='in', name=f'{tag}_PedL', src=f'{src}.pedestal_length_in'),
            dict(label='Pedestal top transverse chamfer', value=e['pedestal_chamfer_in'], unit='in', name=f'{tag}_Chamf', src=f'{src}.pedestal_chamfer_in'),
            dict(label='Backwall thickness', value=e['backwall_thickness_in'], unit='in', name=f'{tag}_BW', src=f'{src}.backwall_thickness_in'),
            dict(label='Backwall rear-face inset', value=e['rear_face_inset_in'], unit='in', name=f'{tag}_Inset', src=f'{src}.rear_face_inset_in'),
            dict(label='Beam end clearance (end-bent record)', value=e['beam_end_clear_in'], unit='in', name=f'{tag}_BmClr', src=f'{src}.beam_end_clear_in (QC: equals physical value)'),
            dict(label='Beam end to CL bearing (end-bent record)', value=e['beam_end_to_bearing_in'], unit='in', name=f'{tag}_BmBrg', src=f'{src}.beam_end_to_bearing_in (QC: equals physical value)'),
            dict(label='Backwall top below the deck-top line at FFBW', value=e['backwall_drop_in'], unit='in', name=f'{tag}_BwDrop', src=f'{src}.backwall_drop_in (D-017 landing seat)'),
            dict(label='Cheek web thickness (outside each cap end)', value=e['cheek_thickness_in'], unit='in', name=f'{tag}_CkT', src=f'{src}.cheek_thickness_in (D-026)'),
            dict(label='Lug projection (bottom and rear)', value=e['lug_projection_in'], unit='in', name=f'{tag}_LugP', src=f'{src}.lug_projection_in (D-026 DXF)'),
            dict(label='Lug width on the cheek outer face', value=e['lug_width_in'], unit='in', name=f'{tag}_LugW', src=f'{src}.lug_width_in (D-026 DXF)'),
            dict(label='MSE panel joint to cheek and lug faces', value=e['mse_joint_in'], unit='in', name=f'{tag}_Joint', src=f'{src}.mse_joint_in (D-026)'),
            dict(label='Cheek top, left (blank = deck edge at FFBW)', value=(e['cheek_top_ft'] or [None, None])[0], unit='ft', name=f'{tag}_CkTopL', fmt=FMT_EL,
                 src=f'{src}.cheek_top_ft[0]: D-033 barrier-return base (MSE interface input)'),
            dict(label='Cheek top, right (blank = deck edge at FFBW)', value=(e['cheek_top_ft'] or [None, None])[1], unit='ft', name=f'{tag}_CkTopR', fmt=FMT_EL,
                 src=f'{src}.cheek_top_ft[1]: D-033 barrier-return base (MSE interface input)'),
        ])
    rd = cfg['roadway']
    row = section(book, ws, row, "Roadway / clearance - CONFIG['roadway']", 7)
    row = param_table(book, ws, row, [
        dict(label='Required vertical clearance', value=rd['clearance_required_ft'], unit='ft', name='Rd_Req', fmt=FMT_FT, src=f'{nb} roadway.clearance_required_ft'),
        dict(label='Bridge station at XML station 0', value=float(rd['bridge_station_at_xml_zero_ft']), unit='ft', name='Rd_Shift', fmt=FMT_STA, src=f'{nb} roadway.bridge_station_at_xml_zero_ft (explicit project assumption)'),
        dict(label='Roadway XML file', value='Geometry Report.xml', name='Rd_File', src=f"SHA-256 {ctx['hashes']['roadway_xml']}"),
    ])
    row = section(book, ws, row, 'Beam design - Mathcad saved results (no Mathcad recalculation)', 7)
    row = mathcad_table(book, ws, row, ctx)
    row = section(book, ws, row, 'Designer assignments and BCC-template method constants (not in the notebook or Mathcad)', 7)
    row = param_table(book, ws, row, [
        dict(label='Concrete class (beam)', value='VI', name='Des_ConcClass', src="Designer input: FDOT Class VI for f'c = 8.5 ksi"),
        dict(label='Standard strand pattern type (table)', value=1, name='Des_Pattern', src='Designer input: strand-pattern detail number shown on the plans (Mathcad: 39 strands, 8 debonded)'),
        dict(label='Plan view case (zero skew)', value=1, name='Des_PlanCase', src='Template convention: Case 1 = square end (notebook supports are normal to the tangent alignment)'),
        dict(label='Bearing pad type', value='F', name='Des_PadType', src='Notebook CONFIG comment: "Type F pads: 10 x 32 x 2-9/16 in" (QC checks the dimensions)'),
        dict(label='Beveled bearing plate limit |slope| >', value=0.02, unit='ft/ft', name='Des_BevelLimit', fmt='0.000', src='BCC template rule'),
        dict(label='Dim P threshold (|grade x depth| below this -> 0)', value=1.0, unit='in', name='Des_DimP', src='BCC template rule'),
        dict(label='FIB-36 gross section area', value=806.58, unit='in2', name='Bm_Area', src='FDOT FIB-36 section property (BCC template BeamProp table)'),
        dict(label='FGE increments between CL bearings', value=FGE_INCREMENTS, name='Des_FGEInc', kind='calc', src='Generator constant (sets the table column count)'),
        dict(label='Bar 3C1/3C2 dimension C (square end)', value=17.5, unit='in', name='Bar_C', src='BCC template Reinforcing formulas (FIB standard bar bending)'),
        dict(label='Bar 3D1/3D2 dimension A (square end)', value=8.5, unit='in', name='Bar_DA', src='BCC template Reinforcing formulas'),
        dict(label='Bar 3D1/3D2 dimension B (square end)', value=15.0, unit='in', name='Bar_DB', src='BCC template Reinforcing formulas'),
        dict(label='Bar 4M1/4M2 dimension D (square end)', value=44.0, unit='in', name='Bar_MD', src='BCC template Reinforcing formulas'),
        dict(label='Bar 3D3 spacing', value=18.0, unit='in', name='Bar_D3s', src='BCC template Reinforcing formulas'),
        dict(label='FIB-36 4M3 constant (template BeamProp column H)', value=6, name='Bar_Mk', src='BCC template BeamProp (FIB 36 -> 6)'),
    ])
    row = section(book, ws, row, 'Reference: FDOT composite elastomeric bearing pad types (BCC template BeamProp)', 7)
    row = header_row(book, ws, row, 2, ['Pad type', 'L (in)', 'W (in)', 'T (in)'], 18)
    pad_rows = [('D', 8, 32, 1 + 29 / 32), ('E', 10, 32, 1 + 29 / 32), ('F', 10, 32, 2 + 9 / 16), ('G', 10, 32, 2 + 9 / 16),
                ('H', 10, 32, 3 + 3 / 16), ('J', 10, 32, 3 + 27 / 32), ('K', 12, 32, 3 + 27 / 32)]
    first = row
    for t, L, W, T in pad_rows:
        book.put(ws, f'B{row}', t, font=F_INPUT, fill=FILL_INPUT)
        book.put(ws, f'C{row}', L, font=F_INPUT, fill=FILL_INPUT)
        book.put(ws, f'D{row}', W, font=F_INPUT, fill=FILL_INPUT)
        book.put(ws, f'E{row}', T, font=F_INPUT, fill=FILL_INPUT, fmt='0.00000')
        row += 1
    book.wb.defined_names['Pad_Table'] = DefinedName('Pad_Table', attr_text=f"'Inputs'!$B${first}:$E${row-1}")
    book.names['Pad_Table'] = 'Inputs pad table'
    ws.freeze_panes = 'A4'
    return ws


MATHCAD_ROWS = [
    ('Input data file named in worksheet', 'InputDataFile', 'input_data_file', '', 'File', None),
    ('Beam type', 'BeamType', 'beam_type', '', 'Type', None),
    ('Beam length', 'L.beam', 'L_beam_ft', 'ft', 'Lbeam', FMT_FT),
    ('Bearing distance (beam end to CL bearing)', 'BearingDistance', 'bearing_distance_in', 'in', 'BrgDist', None),
    ('Pad width (longitudinal)', 'PadWidth', 'pad_width_in', 'in', 'PadW', None),
    ('Beam spacing', 'BeamSpacing', 'beam_spacing_ft', 'ft', 'Spacing', FMT_FT),
    ('Deck thickness', 't.deck', 't_deck_in', 'in', 'Tdeck', '0.000'),
    ('Design build-up', 'h.buildup', 'h_buildup_in', 'in', 'Hbu', '0.000'),
    ('Skew', 'Skew', 'skew_deg', 'deg', 'Skew', None),
    ("Beam f'c (28 day)", 'f.c.beam', 'fc_beam_ksi', 'ksi', 'fc', None),
    ("Beam f'ci (release)", 'f.ci.beam', 'fci_beam_ksi', 'ksi', 'fci', None),
    ('Eci', 'E.ci', 'Eci_ksi', 'ksi', 'Eci', '0.0'),
    ('Ec', 'E.c', 'Ec_ksi', 'ksi', 'Ec', '0.0'),
    ('Relative humidity', 'H', 'humidity_pct', '%', 'H', None),
    ('Environment', 'Environment', 'environment', '', 'Env', None),
    ('Strand type', 'StrandType', 'strand_type', '', 'StrType', None),
    ('Strand area', 'A.strand', 'strand_area_in2', 'in2', 'Astr', '0.000'),
    ('Number of strands', '#Strands', 'strands', '', 'NStr', None),
    ('Number of debonded strands', 'NumberOfDebondedTendons', 'debonded', '', 'NDeb', None),
    ('Jacking force per strand', 'JackingForce.per.strand', 'jacking_kip', 'kip', 'Pj', '0.0000'),
    ('Top (partial) strands', 'NoStrands.partial.ps', 'top_strands', '', 'NTop', None),
    ('Top strand force each', 'Force.strand.partial.ps', 'top_strand_force_kip', 'kip', 'PTop', None),
    ('Top strand diameter', 'Dia.strand.partial.ps', 'top_strand_dia_in', 'in', 'DTop', '0.000'),
    ('Creep coefficient at 120 days', 'Cr.120day', 'creep_120', '', 'Psi', '0.000000'),
    ('Shrinkage strain at 120 days', 'eps.sh.120', 'shrink_120', '', 'Esh', '0.00000000'),
    ('Release deflection, max', 'max(Defl.release)', 'defl_release_in', 'in', 'Drel', '0.0000'),
    ('120-day deflection, max', 'max(Defl.120day)', 'defl_120_in', 'in', 'D120', '0.0000'),
    ('Non-composite DL deflection, min', 'min(Defl.dl.nc)', 'defl_dl_nc_in', 'in', 'Dnc', '0.0000'),
    ('Composite DL deflection, min', 'min(Defl.dl.c)', 'defl_dl_c_in', 'in', 'Dc', '0.0000'),
    ('Net camber at 120 days', 'NetCamber', 'net_camber_in', 'in', 'Net', '0.000000'),
    ('First stirrup from beam end', 'EndSpace', 'stirrup_end_space_in', 'in', 'EndSp', '0.00'),
]


def mathcad_table(book: Book, ws, row: int, ctx):
    spans = sorted(ctx['mathcad'])
    headers = ['Parameter', 'Mathcad variable'] + [f'Span {i}' for i in spans] + ['Unit', 'Workbook names']
    row = header_row(book, ws, row, 2, headers[:2] + headers[2:2 + len(spans)] + headers[-2:], 18)
    for label, var, key, unit, short, fmt in MATHCAD_ROWS:
        book.put(ws, f'B{row}', label)
        book.put(ws, f'C{row}', var, font=F_NOTE)
        names = []
        for k, i in enumerate(spans):
            ref = f'{col(4 + k)}{row}'
            book.put(ws, ref, _value(ctx['mathcad'][i][key]), font=F_INPUT, fill=FILL_INPUT, fmt=fmt, border=BOX)
            book.define(f'MC{i}_{short}', ws, ref)
            names.append(f'MC{i}_{short}')
        book.put(ws, f'{col(4 + len(spans))}{row}', unit)
        book.put(ws, f'{col(5 + len(spans))}{row}', ', '.join(names), font=F_NOTE)
        row += 1
    for zone in ('A1', 'A2', 'A3', 'S1', 'S2', 'S3', 'S4'):
        for field, label, unit in (('spacing_in', 'spacing', 'in'), ('spaces', 'spaces', ''), ('area_in2', 'Av', 'in2')):
            book.put(ws, f'B{row}', f'Stirrup zone {zone} {label}')
            book.put(ws, f'C{row}', {'spacing_in': 'StirrupSpacing', 'spaces': 'NumberOfStirrupSpaces', 'area_in2': 'StirrupArea'}[field] + f'.{zone}', font=F_NOTE)
            names = []
            for k, i in enumerate(spans):
                z = next(zz for zz in ctx['mathcad'][i]['stirrup_zones'] if zz['zone'] == zone)
                ref = f'{col(4 + k)}{row}'
                book.put(ws, ref, z[field], font=F_INPUT, fill=FILL_INPUT, border=BOX)
                name = f'MC{i}_{zone}_{ {"spacing_in": "s", "spaces": "n", "area_in2": "A"}[field] }'.replace(' ', '')
                book.define(name, ws, ref)
                names.append(name)
            book.put(ws, f'{col(4 + len(spans))}{row}', unit)
            book.put(ws, f'{col(5 + len(spans))}{row}', ', '.join(names), font=F_NOTE)
            row += 1
    for k, i in enumerate(spans):
        book.put(ws, f'B{row + k}', f"Span {i} source: {ctx['mathcad'][i]['source_file']}", font=F_NOTE)
        book.put(ws, f'G{row + k}', f"SHA-256 {ctx['mathcad'][i]['source_sha256']}", font=F_NOTE)
    return row + len(spans) + 1


# ======================================================= Profile & Deck ==
def build_profile(book: Book, ctx):
    m = ctx['model']
    ws = book.sheet('Profile & Deck', '1F2A36')
    for c, w in zip('ABCDEFGHIJKLM', (3, 40, 11, 16, 9, 16, 14, 14, 14, 14, 14, 14, 14)):
        ws.column_dimensions[c].width = w
    row = title_block(book, ws, 'Profile & Deck', 'PGL profile, deck cross slope, beam offsets and haunch check points', 11)
    row = section(book, ws, row, 'Profile constants', 10)
    row = param_table(book, ws, row, [
        dict(label='Profile length', sym='L', value='=Prof_Sta1-Prof_Sta0', unit='ft', name='Prof_L', kind='calc', fmt=FMT_FT, src='s1 - s0'),
        dict(label='PI elevation (brokeback)', sym='zPI', value='=(Prof_Z0+Prof_Gnom*(Prof_PISta-Prof_Sta0)+Prof_Z1+Prof_Gnom*(Prof_Sta1-Prof_PISta))/2',
             unit='ft', name='Prof_PIZ', kind='calc', fmt=FMT_EL, src='Average of the two tangents projected to the PI'),
        dict(label='Grade 1 (entering)', sym='g1', value='=IF(Prof_Mode="parabolic",(Prof_Z1-Prof_Z0)/Prof_L-Prof_dG/2,(Prof_PIZ-Prof_Z0)/(Prof_PISta-Prof_Sta0))',
             unit='ft/ft', name='Prof_G1', kind='calc', fmt=FMT_GRADE, src='Notebook BridgeModel.__init__'),
        dict(label='Grade 2 (leaving)', sym='g2', value='=IF(Prof_Mode="parabolic",Prof_G1+Prof_dG,(Prof_Z1-Prof_PIZ)/(Prof_Sta1-Prof_PISta))',
             unit='ft/ft', name='Prof_G2', kind='calc', fmt=FMT_GRADE, src='Notebook BridgeModel.__init__'),
        dict(label='Parabolic coefficient (0 for brokeback)', sym='kq', value='=IF(Prof_Mode="parabolic",0.5*(Prof_G2-Prof_G1)/Prof_L,0)',
             unit='1/ft', name='Prof_Kq', kind='calc', fmt='0.000000000', src='z = z0 + g1 x + kq x^2'),
    ])
    row = section(book, ws, row, 'PGL at control stations', 10)
    row = header_row(book, ws, row, 2, ['Control point', 'Station (ft)', 'PGL (ft)', 'Station label'], 18)
    controls = [('FFBW End Bent 1 (profile start)', '=Prof_Sta0'), ('Tangent intersection (PI)', '=Prof_PISta'),
                ('FFBW End Bent 2 (profile end)', '=Prof_Sta1')]
    ctx['pgl_controls'] = []
    for label, formula in controls:
        book.put(ws, f'B{row}', label)
        book.put(ws, f'C{row}', formula, font=F_LINK, fmt=FMT_STA)
        book.put(ws, f'D{row}', '=' + PGL(f'C{row}'), fmt=FMT_EL)
        book.put(ws, f'E{row}', f'=TEXT(INT(ROUND(C{row}*100,0)/10000),"0")&"+"&TEXT(MOD(ROUND(C{row}*100,0),10000)/100,"00.00")')
        ctx['pgl_controls'].append((label, f'D{row}'))
        row += 1
    row += 1
    row = section(book, ws, row, 'Deck edges', 10)
    row = param_table(book, ws, row, [
        dict(label='Left deck edge (coping) offset', value='=Deck_C-Deck_W/2', unit='ft', name='Deck_YL', kind='calc', fmt=FMT_FT, src='Offsets measured from the PGL, + to the right looking up-station'),
        dict(label='Right deck edge (coping) offset', value='=Deck_C+Deck_W/2', unit='ft', name='Deck_YR', kind='calc', fmt=FMT_FT, src=''),
        dict(label='Deck cross-slope drop at left edge', value='=' + DELTA('Deck_YL'), unit='ft', name='Deck_dL', kind='calc', fmt=FMT_FT, src='delta(y)'),
        dict(label='Deck cross-slope drop at right edge', value='=' + DELTA('Deck_YR'), unit='ft', name='Deck_dR', kind='calc', fmt=FMT_FT, src='delta(y)'),
    ])
    n = m.count
    row = section(book, ws, row, 'Beam lines and haunch check points (flange edges, CL and any crown point on the flange)', 12)
    row = header_row(book, ws, row, 2, ['Beam', 'Offset y (ft)', 'Left flange edge', 'CL', 'Right flange edge', 'Crown point on flange',
                                         'delta left (ft)', 'delta CL (ft)', 'delta right (ft)', 'delta crown pt (ft)', 'delta min (ft)'], 30)
    for k in range(1, n + 1):
        off = '=IF(ISNUMBER(Bm_First),Bm_First,-(' + '+'.join(f'Bm_Gap{j}' for j in range(1, n)) + ')/2)' if k == 1 else f'=Bm{k-1}_Off+Bm_Gap{k-1}'
        if n == 1:
            off = '=IF(ISNUMBER(Bm_First),Bm_First,0)'
        book.put(ws, f'B{row}', f'Beam {k}')
        book.put(ws, f'C{row}', off, fmt=FMT_FT); book.define(f'Bm{k}_Off', ws, f'C{row}')
        book.put(ws, f'D{row}', f'=C{row}-Bm_Bf/24', fmt=FMT_FT); book.define(f'Bm{k}_FL', ws, f'D{row}')
        book.put(ws, f'E{row}', f'=C{row}', fmt=FMT_FT)
        book.put(ws, f'F{row}', f'=C{row}+Bm_Bf/24', fmt=FMT_FT); book.define(f'Bm{k}_FR', ws, f'F{row}')
        book.put(ws, f'G{row}', f'=MAX(D{row},MIN(0,F{row}))', fmt=FMT_FT)
        for j, src in enumerate('DEFG'):
            book.put(ws, f'{col(8 + j)}{row}', '=' + DELTA(f'{src}{row}'), fmt=FMT_FT)
        book.define(f'Bm{k}_dCL', ws, f'I{row}')
        book.put(ws, f'L{row}', f'=MIN(H{row}:K{row})', fmt=FMT_FT); book.define(f'Bm{k}_dMin', ws, f'L{row}')
        row += 1
    row += 1
    row = section(book, ws, row, 'Layout checks', 10)
    checks = [('Beam flanges inside the deck width', f'=IF(AND(Bm1_FL>=Deck_YL-1E-9,Bm{n}_FR<=Deck_YR+1E-9),"PASS","FAIL")'),
              ('Deck mode implemented (crowned or uniform)', '=IF(OR(Deck_Mode="crowned",Deck_Mode="uniform"),"PASS","FAIL")'),
              ('Profile mode implemented (brokeback or parabolic)', '=IF(OR(Prof_Mode="brokeback",Prof_Mode="parabolic"),"PASS","FAIL")'),
              ('PI inside the profile', '=IF(AND(Prof_PISta>Prof_Sta0,Prof_PISta<Prof_Sta1),"PASS","FAIL")')]
    if n > 1:
        checks.insert(1, ('Flanges do not overlap', '=IF(MIN(' + ','.join(f'Bm_Gap{j}' for j in range(1, n)) + ')>=Bm_Bf/12-1E-9,"PASS","FAIL")'))
    ctx['layout_checks'] = []
    for label, formula in checks:
        book.put(ws, f'B{row}', label)
        book.put(ws, f'C{row}', formula, align=CENTER, border=BOX)
        ctx['layout_checks'].append((label, book.xref(ws, f'C{row}')))
        row += 1
    check_formatting(ws, f'C{row-len(checks)}:C{row-1}')
    return ws


# ========================================================= Beam Geometry ==
def build_beam_geometry(book: Book, ctx):
    m = ctx['model']; spans = ctx['spans']; n = m.count
    ws = book.sheet('Beam Geometry', '1F2A36')
    widths = [3, 9, 12, 12, 13, 13, 12, 12, 12, 12, 12, 12, 11, 11, 12, 11, 11, 11, 9, 11, 11]
    for k, w in enumerate(widths, start=1):
        ws.column_dimensions[col(k)].width = w
    row = title_block(book, ws, 'Beam Geometry', 'Physical beam ends, bearings, deck limits, bearing chords and Table of Beam Variables dimensions', 15)
    row = section(book, ws, row, 'Span layout (stations in ft)', 20)
    heads = ['Span', 'Datum start', 'Datum end', 'Start support', 'End support', 'Beam start', 'Beam end', 'CL brg A (start)',
             'CL brg B (end)', 'Deck start', 'Deck end', 'Beam length', 'Brg span Lb', 'Brg midpoint', 'PGL at A', 'PGL at B',
             'Chord grade', 'Layout check', 'Mathcad L.beam', 'Beam length - Mathcad']
    row = header_row(book, ws, row, 2, heads, 42)
    last = len(spans)
    for i, span in enumerate(spans, start=1):
        r = row
        first_span, last_span = i == 1, i == last
        start_name = 'End Bent 1' if first_span else 'Pier'
        end_name = 'End Bent 2' if last_span else 'Pier'
        book.put(ws, f'B{r}', span['id'])
        book.put(ws, f'C{r}', f'=Sp{i}_Sta0', font=F_LINK, fmt=FMT_STA)
        book.put(ws, f'D{r}', f'=Sp{i}_Sta1', font=F_LINK, fmt=FMT_STA)
        book.put(ws, f'E{r}', start_name); book.put(ws, f'F{r}', end_name)
        book.put(ws, f'G{r}', f'=C{r}+' + ('Phys_EndClr/12' if first_span else 'Phys_PierGap/24'), fmt=FMT_STA)
        book.put(ws, f'H{r}', f'=D{r}-' + ('Phys_EndClr/12' if last_span else 'Phys_PierGap/24'), fmt=FMT_STA)
        book.put(ws, f'I{r}', f'=G{r}+Phys_Setback/12', fmt=FMT_STA)
        book.put(ws, f'J{r}', f'=H{r}-Phys_Setback/12', fmt=FMT_STA)
        book.put(ws, f'K{r}', f'=C{r}' + ('+Phys_DeckClr/12' if first_span else ''), fmt=FMT_STA)
        book.put(ws, f'L{r}', f'=D{r}' + ('-Phys_DeckClr/12' if last_span else ''), fmt=FMT_STA)
        book.put(ws, f'M{r}', f'=H{r}-G{r}', fmt=FMT_FT)
        book.put(ws, f'N{r}', f'=J{r}-I{r}', fmt=FMT_FT)
        book.put(ws, f'O{r}', f'=(I{r}+J{r})/2', fmt=FMT_STA)
        book.put(ws, f'P{r}', '=' + PGL(f'I{r}'), fmt=FMT_EL)
        book.put(ws, f'Q{r}', '=' + PGL(f'J{r}'), fmt=FMT_EL)
        book.put(ws, f'R{r}', f'=(Q{r}-P{r})/N{r}', fmt=FMT_GRADE)
        book.put(ws, f'S{r}', f'=IF(AND(C{r}<=G{r},G{r}<I{r},I{r}<J{r},J{r}<H{r},H{r}<=D{r},K{r}<L{r}),"PASS","FAIL")', align=CENTER, border=BOX)
        book.put(ws, f'T{r}', f'=MC{i}_Lbeam', font=F_LINK, fmt=FMT_FT)
        book.put(ws, f'U{r}', f'=M{r}-T{r}', fmt=FMT_FT)
        for name, c in (('Bm0', 'G'), ('Bm1', 'H'), ('BrgA', 'I'), ('BrgB', 'J'), ('Dk0', 'K'), ('Dk1', 'L'), ('LBeam', 'M'),
                        ('Lb', 'N'), ('Mid', 'O'), ('ZA', 'P'), ('ZB', 'Q'), ('Gc', 'R'), ('Chk', 'S')):
            book.define(f'Sp{i}_{name}', ws, f'{c}{r}')
        row += 1
    check_formatting(ws, f'S{row-last}:S{row-1}')
    row += 1
    row = section(book, ws, row, 'Beam dimensions for the Table of Beam Variables (template method, notebook geometry)', 20)
    heads = ['Beam', 'Span', 'Offset (ft)', 'Raise (in)', 'Horizontal length (ft)', 'Sloped length Ls (ft)', 'Skew (deg)',
             'Angle phi End 1', 'Angle phi End 2', 'Dim J (in)', 'Dim K1 (in)', 'Dim K2 (in)', 'Dim P (in)', 'End elev. cond.',
             'Prestress P (kip)', 'Creep factor C1', 'Shortening R (in)', 'Dim L incl. R (ft)', 'Dim L ft', 'Dim L in']
    row = header_row(book, ws, row, 2, heads, 42)
    ctx['beam_rows'] = {}
    for i, span in enumerate(spans, start=1):
        for k in range(1, n + 1):
            r = row
            book.put(ws, f'B{r}', f'{span["id"]}-{k}')
            book.put(ws, f'C{r}', span['id'])
            book.put(ws, f'D{r}', f'=Bm{k}_Off', font=F_LINK, fmt=FMT_FT)
            book.put(ws, f'E{r}', f'=Bm{k}_Raise', font=F_LINK)
            book.put(ws, f'F{r}', f'=Sp{i}_LBeam', fmt=FMT_FT)
            book.put(ws, f'G{r}', f'=F{r}*SQRT(1+Sp{i}_Gc^2)', fmt=FMT_FT)
            book.put(ws, f'H{r}', f'=MC{i}_Skew', font=F_LINK)
            book.put(ws, f'I{r}', f'=90-H{r}'); book.put(ws, f'J{r}', f'=90-H{r}')
            book.put(ws, f'K{r}', '=Phys_Setback', fmt='0.00')
            book.put(ws, f'L{r}', f'=(Sp{i}_BrgA-Sp{i}_Sta0)*12', fmt='0.00')
            book.put(ws, f'M{r}', f'=(Sp{i}_Sta1-Sp{i}_BrgB)*12', fmt='0.00')
            book.put(ws, f'N{r}', f'=IF(ABS(Sp{i}_Gc*Bm_D*12)<Des_DimP,0,-Sp{i}_Gc*Bm_D*12)', fmt='0.00')
            book.put(ws, f'O{r}', f'=IF(N{r}<0,2,IF(N{r}>0,3,1))')
            book.put(ws, f'P{r}', f'=MC{i}_NStr*MC{i}_Pj+MC{i}_NTop*MC{i}_PTop', fmt='0.00')
            book.put(ws, f'Q{r}', f'=1+MC{i}_Psi', fmt='0.0000')
            book.put(ws, f'R{r}', f'=Q{r}*P{r}*G{r}*12/(MC{i}_Eci*Bm_Area)+ABS(MC{i}_Esh)*G{r}*12', fmt='0.000')
            book.put(ws, f'S{r}', f'=G{r}+R{r}/12', fmt=FMT_FT)
            book.put(ws, f'T{r}', f'=ROUNDDOWN(S{r},0)', fmt='0')
            book.put(ws, f'U{r}', f'=(S{r}-T{r})*12', fmt='0.00')
            ctx['beam_rows'][(i, k)] = r
            row += 1
    row += 1
    book.put(ws, f'B{row}', 'Dim L follows the BCC template (sloped beam length + shortening R). Confirm against the current '
             'FDOT instructions for the Table of Beam Variables before issue; the sloped length without R is column G.', font=F_WARN)
    ctx['ws_beam'] = ws
    return ws


# ========================================================= Camber & Haunch ==
def build_camber(book: Book, ctx):
    m = ctx['model']; spans = ctx['spans']; n = m.count
    ws = book.sheet('Camber & Haunch', '1F2A36')
    widths = [3, 12, 14, 13, 13, 13, 13, 13, 13, 13, 13, 13, 13, 13]
    for k, w in enumerate(widths, start=1):
        ws.column_dimensions[col(k)].width = w
    row = title_block(book, ws, 'Camber & Haunch', 'Common support drop, build-up (Dims B, C, D) and minimum flange haunch', 12)
    row = section(book, ws, row, 'A. Haunch shape per span: h(t) = K - m t(1-t), with t = (s - brg A)/Lb', 12)
    row = header_row(book, ws, row, 2, ['Span', 'Camber C (in)', 'Lb (ft)', 't at beam start', 't at beam end', 'm (ft)',
                                         'max of m t(1-t) on beam', 'Mathcad NetCamber (in)', 'C - Mathcad (in)'], 32)
    for i, span in enumerate(spans, start=1):
        r = row
        book.put(ws, f'B{r}', span['id'])
        book.put(ws, f'C{r}', f'=Sp{i}_C', font=F_LINK, fmt='0.000000')
        book.put(ws, f'D{r}', f'=Sp{i}_Lb', font=F_LINK, fmt=FMT_FT)
        book.put(ws, f'E{r}', f'=(Sp{i}_Bm0-Sp{i}_BrgA)/Sp{i}_Lb', fmt='0.000000')
        book.put(ws, f'F{r}', f'=(Sp{i}_Bm1-Sp{i}_BrgA)/Sp{i}_Lb', fmt='0.000000')
        book.put(ws, f'G{r}', f'=Prof_Kq*Sp{i}_Lb^2+Sp{i}_C/3', fmt='0.0000000')
        book.put(ws, f'H{r}', f'=MAX(G{r}/4,G{r}*E{r}*(1-E{r}),G{r}*F{r}*(1-F{r}))', fmt='0.0000000')
        book.put(ws, f'I{r}', f'=MC{i}_Net', font=F_LINK, fmt='0.000000')
        book.put(ws, f'J{r}', f'=C{r}-I{r}', fmt='0.000000000')
        book.define(f'Sp{i}_m', ws, f'G{r}'); book.define(f'Sp{i}_W', ws, f'H{r}')
        row += 1
    row += 1
    row = section(book, ws, row, 'B. Minimum uncorrected haunch per beam and the common support drop D', 12)
    row = header_row(book, ws, row, 2, ['Beam', 'delta min (ft)', 'Raise (in)', 'Min haunch with D = 0 (ft)', 'Min flange haunch (in)'], 32)
    first = row
    for i, span in enumerate(spans, start=1):
        for k in range(1, n + 1):
            r = row
            book.put(ws, f'B{r}', f'{span["id"]}-{k}')
            book.put(ws, f'C{r}', f'=Bm{k}_dMin', font=F_LINK, fmt=FMT_FT)
            book.put(ws, f'D{r}', f'=Bm{k}_Raise', font=F_LINK)
            book.put(ws, f'E{r}', f'=C{r}-Deck_T-D{r}/12-Bm_D-Sp{i}_W', fmt=FMT_FT)
            book.put(ws, f'F{r}', f'=(E{r}+Sup_Drop)*12', fmt=FMT_IN)
            row += 1
    last_r = row - 1
    row += 1
    book.put(ws, f'B{row}', 'Common support drop D (ft)', bold=True)
    book.put(ws, f'E{row}', f'=Bm_Hmin/12-MIN(E{first}:E{last_r})', fmt='0.0000000000', border=BOX, bold=True)
    book.define('Sup_Drop', ws, f'E{row}')
    book.put(ws, f'G{row}', 'D = h_min,req - min over every beam and flange point of the haunch with D = 0', font=F_NOTE)
    row += 1
    ctx['haunch_min_range'] = (first, last_r)
    book.put(ws, f'B{row}', 'Minimum flange haunch, whole bridge (in)', bold=True)
    book.put(ws, f'E{row}', f'=MIN(F{first}:F{last_r})', fmt=FMT_IN, border=BOX, bold=True)
    book.define('Haunch_Min', ws, f'E{row}')
    book.put(ws, f'F{row}', '=IF(E{0}>=Bm_Hmin-1E-8,"PASS","FAIL")'.format(row), align=CENTER, border=BOX)
    check_formatting(ws, f'F{row}')
    row += 2
    row = section(book, ws, row, 'C. Haunch over each beam at the bearings and midspan', 13)
    heads = ['Beam', 'Location', 'Station s (ft)', 't', 'PGL (ft)', 'Chord (ft)', 'Camber (ft)', 'Beam bottom (ft)', 'Beam top (ft)',
             'Deck @ CL beam (ft)', 'Haunch @ CL (in)', 'Min flange haunch (in)']
    row = header_row(book, ws, row, 2, heads, 32)
    ctx['haunch_rows'] = {}
    for i, span in enumerate(spans, start=1):
        for k in range(1, n + 1):
            for loc, s_expr in (('Begin brg', f'Sp{i}_BrgA'), ('Midspan', f'Sp{i}_Mid'), ('End brg', f'Sp{i}_BrgB')):
                r = row
                book.put(ws, f'B{r}', f'{span["id"]}-{k}')
                book.put(ws, f'C{r}', loc)
                book.put(ws, f'D{r}', f'={s_expr}', font=F_LINK, fmt=FMT_STA)
                book.put(ws, f'E{r}', f'=(D{r}-Sp{i}_BrgA)/Sp{i}_Lb', fmt='0.000000')
                book.put(ws, f'F{r}', '=' + PGL(f'D{r}'), fmt=FMT_EL)
                book.put(ws, f'G{r}', f'=Sp{i}_ZA+(Sp{i}_ZB-Sp{i}_ZA)*E{r}', fmt=FMT_EL)
                book.put(ws, f'H{r}', f'=Sp{i}_C/3*E{r}*(1-E{r})', fmt='0.000000')
                book.put(ws, f'I{r}', f'=G{r}-Sup_Drop+H{r}+Bm{k}_Raise/12', fmt=FMT_EL)
                book.put(ws, f'J{r}', f'=I{r}+Bm_D', fmt=FMT_EL)
                book.put(ws, f'K{r}', f'=F{r}+Bm{k}_dCL', fmt=FMT_EL)
                book.put(ws, f'L{r}', f'=(K{r}-Deck_T-J{r})*12', fmt=FMT_IN)
                book.put(ws, f'M{r}', f'=(F{r}+Bm{k}_dMin-Deck_T-J{r})*12', fmt=FMT_IN)
                ctx['haunch_rows'][(i, k, loc)] = r
                row += 1
    row += 1
    row = section(book, ws, row, 'D. Build-up summary (FDOT Build-Up & Deflection table)', 12)
    heads = ['Beam', 'Dim B begin (in)', 'Dim C CL span (in)', 'Dim D end (in)', 'Mid-ordinate R (in)', 'Net camber C (in)', 'Build-up case']
    row = header_row(book, ws, row, 2, heads, 32)
    ctx['buildup_rows'] = {}
    for i, span in enumerate(spans, start=1):
        for k in range(1, n + 1):
            r = row
            ra, rm, rb = (ctx['haunch_rows'][(i, k, loc)] for loc in ('Begin brg', 'Midspan', 'End brg'))
            book.put(ws, f'B{r}', f'{span["id"]}-{k}')
            book.put(ws, f'C{r}', f'=L{ra}', fmt='0.000')
            book.put(ws, f'D{r}', f'=L{rm}', fmt='0.000')
            book.put(ws, f'E{r}', f'=L{rb}', fmt='0.000')
            book.put(ws, f'F{r}', f'=(K{rm}-(K{ra}+K{rb})/2)*12', fmt='0.000000')
            book.put(ws, f'G{r}', f'=Sp{i}_C', font=F_LINK, fmt='0.000')
            book.put(ws, f'H{r}', f'=IF(ROUND(F{r},6)<0,"Case 2",IF(ROUND(F{r},6)>G{r},"Case 4",IF(ROUND(F{r},6)>0,"Case 3","Case 1")))')
            ctx['buildup_rows'][(i, k)] = r
            row += 1
    row += 1
    row = section(book, ws, row, 'E. Mathcad deflections by span (in, + up)', 12)
    heads = ['Span', 'Release (max)', '120 day (max)', 'Non-comp. DL (min)', 'Composite DL (min)', 'Sum = NetCamber', 'Mathcad NetCamber', 'CONFIG camber', 'Check']
    row = header_row(book, ws, row, 2, heads, 32)
    for i, span in enumerate(spans, start=1):
        r = row
        book.put(ws, f'B{r}', span['id'])
        for c, nm in zip('CDEF', ('Drel', 'D120', 'Dnc', 'Dc')):
            book.put(ws, f'{c}{r}', f'=MC{i}_{nm}', font=F_LINK, fmt='0.0000')
        book.put(ws, f'G{r}', f'=D{r}+E{r}+F{r}', fmt='0.000000')
        book.put(ws, f'H{r}', f'=MC{i}_Net', font=F_LINK, fmt='0.000000')
        book.put(ws, f'I{r}', f'=Sp{i}_C', font=F_LINK, fmt='0.000000')
        book.put(ws, f'J{r}', f'=IF(AND(ABS(G{r}-H{r})<1E-9,ABS(H{r}-I{r})<1E-9),"PASS","FAIL")', align=CENTER, border=BOX)
        ctx.setdefault('deflection_checks', []).append(book.xref(ws, f'J{r}'))
        row += 1
    check_formatting(ws, f'J{row-len(spans)}:J{row-1}')
    return ws


# =================================================================== FGE ==
def fge_lines(m):
    lines = [('Left coping (deck edge)', 'Deck_YL', m.deck_left)]
    for k, off in enumerate(m.offsets, start=1):
        lines.append((f'CL Beam {{span}}-{k}', f'Bm{k}_Off', float(off)))
    lines.append(('PGL', '0', 0.0))
    lines.append(('Right coping (deck edge)', 'Deck_YR', m.deck_right))
    lines.sort(key=lambda t: (t[2], t[0] != 'PGL'))
    return lines


def build_fge(book: Book, ctx):
    m = ctx['model']; spans = ctx['spans']
    ws = book.sheet('FGE', '1F2A36')
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 12
    for k in range(5, 5 + FGE_INCREMENTS + 3):
        ws.column_dimensions[col(k)].width = 11.5
    row = title_block(book, ws, 'Finished Grade Elevations', 'FGE = PGL(s) + delta(y) at every line and station (top of deck)', 14)
    ctx['fge'] = {}
    lines = fge_lines(m)
    last = len(spans)
    for i, span in enumerate(spans, start=1):
        row = section(book, ws, row, f'Span {span["id"]}', 4 + FGE_INCREMENTS + 3)
        start_label = 'FFBW EB 1' if i == 1 else 'CL Pier'
        end_label = 'FFBW EB 2' if i == last else 'CL Pier'
        labels = [start_label, 'CL Brg'] + [str(k) for k in range(1, FGE_INCREMENTS)] + ['CL Brg', end_label]
        row = header_row(book, ws, row, 2, ['Location', 'Offset y (ft)', 'delta(y) (ft)'] + labels, 30)
        st_row = row
        book.put(ws, f'B{row}', 'Station (ft)', bold=True)
        exprs = [f'=Sp{i}_Sta0', f'=Sp{i}_BrgA'] + [f'=Sp{i}_BrgA+{k}*Sp{i}_Lb/Des_FGEInc' for k in range(1, FGE_INCREMENTS)] + [f'=Sp{i}_BrgB', f'=Sp{i}_Sta1']
        cols = [col(5 + k) for k in range(len(exprs))]
        for c, e in zip(cols, exprs):
            book.put(ws, f'{c}{row}', e, fmt=FMT_STA)
        row += 1
        book.put(ws, f'B{row}', 'Station label')
        for c in cols:
            book.put(ws, f'{c}{row}', f'=TEXT(INT(ROUND({c}{st_row}*100,0)/10000),"0")&"+"&TEXT(MOD(ROUND({c}{st_row}*100,0),10000)/100,"00.00")', align=CENTER)
        lab_row = row
        row += 1
        book.put(ws, f'B{row}', 'Deck present?')
        for c in cols:
            book.put(ws, f'{c}{row}', f'=IF(AND({c}{st_row}>=Sp{i}_Dk0-1E-9,{c}{st_row}<=Sp{i}_Dk1+1E-9),"deck","PGL ext.")', font=F_NOTE, align=CENTER)
        row += 1
        pgl_row = row
        book.put(ws, f'B{row}', 'PGL elevation (ft)', bold=True)
        for c in cols:
            book.put(ws, f'{c}{row}', '=' + PGL(f'{c}{st_row}'), fmt=FMT_EL)
        row += 1
        line_rows = {}
        for label, off, _ in lines:
            name = label.format(span=span['id'])
            book.put(ws, f'B{row}', name)
            book.put(ws, f'C{row}', f'={off}' if off != '0' else 0, font=F_LINK if off != '0' else F_BASE, fmt=FMT_FT)
            book.put(ws, f'D{row}', '=' + DELTA(f'C{row}'), fmt=FMT_FT)
            for c in cols:
                book.put(ws, f'{c}{row}', f'={c}${pgl_row}+$D{row}', fmt=FMT_EL, border=BOX)
            line_rows[name] = row
            row += 1
        book.put(ws, f'B{row}', 'Increment (ft)')
        book.put(ws, f'C{row}', f'=Sp{i}_Lb/Des_FGEInc', fmt=FMT_FT)
        inc_ref = f'C{row}'
        row += 2
        ctx['fge'][i] = dict(cols=cols, labels=labels, station_row=st_row, label_row=lab_row, lines=line_rows, inc=inc_ref, pgl_row=pgl_row)
    ws.freeze_panes = 'C7'
    return ws


# =========================================================== Bearing Seats ==
def seat_list(ctx):
    spans = ctx['spans']; n = ctx['model'].count; last = len(spans)
    seats = []
    for i in range(1, last + 1):
        start_support = ('End Bent 1', 'EB1') if i == 1 else ('Pier', 'Pier')
        end_support = ('End Bent 2', 'EB2') if i == last else ('Pier', 'Pier')
        for face, (sup, tag) in (('Start', start_support), ('End', end_support)):
            for k in range(1, n + 1):
                seats.append(dict(support=sup, tag=tag, span=i, face=face, beam=k))
    order = {'End Bent 1': 0, 'Pier': 1, 'End Bent 2': 2}
    seats.sort(key=lambda s: (order[s['support']], s['span'], s['beam']))
    return seats


def build_seats(book: Book, ctx):
    m = ctx['model']; spans = ctx['spans']; n = m.count
    ws = book.sheet('Bearing Seats', '1F2A36')
    ws.column_dimensions['A'].width = 3
    for k in range(2, 24):
        ws.column_dimensions[col(k)].width = 11.5
    ws.column_dimensions['B'].width = 12
    row = title_block(book, ws, 'Bearing Seats', 'Pad-top planes, bearing slopes, pedestal tops and pedestal corner elevations', 16)
    row = section(book, ws, row, 'A. Seat plane at each bearing (soffit chord over the pad length)', 22)
    heads = ['Support', 'Span', 'Face', 'Beam', 'Beam ID', 'Offset (ft)', 'CL brg s (ft)', 'Pad L (in)', 'Pad T (in)', 's lo (ft)', 's hi (ft)',
             't lo', 't hi', 't brg', 'Soffit @ s lo', 'Soffit @ s hi', 'Soffit @ CL brg', 'Pad-top plane (ft)', 'Bearing slope',
             'Chord gap (in)', 'Pedestal top @ CL brg (ft)', 'Pad on beam']
    row = header_row(book, ws, row, 2, heads, 42)
    ctx['seat_rows'] = {}
    for seat in seat_list(ctx):
        r = row; i = seat['span']; k = seat['beam']; tag = seat['tag']
        pad_l = 'Pier_PadL' if tag == 'Pier' else f'{tag}_PadL'
        pad_t = 'Pier_Pad' if tag == 'Pier' else f'{tag}_Pad'
        book.put(ws, f'B{r}', seat['support']); book.put(ws, f'C{r}', spans[i - 1]['id']); book.put(ws, f'D{r}', seat['face'])
        book.put(ws, f'E{r}', k); book.put(ws, f'F{r}', f'{spans[i-1]["id"]}-{k}')
        book.put(ws, f'G{r}', f'=Bm{k}_Off', font=F_LINK, fmt=FMT_FT)
        book.put(ws, f'H{r}', f'=Sp{i}_Brg' + ('A' if seat['face'] == 'Start' else 'B'), font=F_LINK, fmt=FMT_STA)
        book.put(ws, f'I{r}', f'={pad_l}', font=F_LINK)
        book.put(ws, f'J{r}', f'={pad_t}', font=F_LINK, fmt=FMT_IN)
        book.put(ws, f'K{r}', f'=H{r}-I{r}/24', fmt=FMT_STA)
        book.put(ws, f'L{r}', f'=H{r}+I{r}/24', fmt=FMT_STA)
        for c, src in (('M', 'K'), ('N', 'L'), ('O', 'H')):
            book.put(ws, f'{c}{r}', f'=({src}{r}-Sp{i}_BrgA)/Sp{i}_Lb', fmt='0.0000000')
        for c, t in (('P', 'M'), ('Q', 'N'), ('R', 'O')):
            book.put(ws, f'{c}{r}', '=' + BOTTOM(i, f'{t}{r}', f'Bm{k}_Raise'), fmt=FMT_EL)
        book.put(ws, f'S{r}', f'=(P{r}+Q{r})/2', fmt='0.000000')
        book.put(ws, f'T{r}', f'=(Q{r}-P{r})/(I{r}/12)', fmt=FMT_GRADE)
        book.put(ws, f'U{r}', f'=(R{r}-S{r})*12', fmt='0.000000')
        book.put(ws, f'V{r}', f'=S{r}-J{r}/12', fmt='0.000000')
        book.put(ws, f'W{r}', f'=IF(AND(K{r}>=Sp{i}_Bm0-1E-9,L{r}<=Sp{i}_Bm1+1E-9),"PASS","FAIL")', align=CENTER, border=BOX)
        ctx['seat_rows'][(seat['tag'], i, k)] = r
        row += 1
    check_formatting(ws, f'W{row-len(ctx["seat_rows"])}:W{row-1}')
    row += 1
    row = section(book, ws, row, 'B. End-bent pedestal tops: plane of the seat, level transversely (A/C back edge at FFBW, B/D span-side edge)', 22)
    heads = ['Support', 'Beam ID', 'Back edge s (ft)', 'Front edge s (ft)', 'CL brg s (ft)', 'Ped. top @ CL brg', 'Slope', 'EL A', 'EL B',
             'EL C', 'EL D', 'Pad inside pedestal']
    row = header_row(book, ws, row, 2, heads, 32)
    ctx['eb_ped_rows'] = {}
    last = len(spans)
    for tag, i, sign, datum in (('EB1', 1, 1, 'Sp1_Sta0'), ('EB2', last, -1, f'Sp{last}_Sta1')):
        for k in range(1, n + 1):
            r = row; s = ctx['seat_rows'][(tag, i, k)]
            book.put(ws, f'B{r}', 'End Bent 1' if tag == 'EB1' else 'End Bent 2')
            book.put(ws, f'C{r}', f'{spans[i-1]["id"]}-{k}')
            book.put(ws, f'D{r}', f'={datum}', font=F_LINK, fmt=FMT_STA)
            book.put(ws, f'E{r}', f'={datum}{"+" if sign > 0 else "-"}{tag}_PedL/12', fmt=FMT_STA)
            book.put(ws, f'F{r}', f'=H{s}', fmt=FMT_STA)
            book.put(ws, f'G{r}', f'=V{s}', fmt='0.000000')
            book.put(ws, f'H{r}', f'=T{s}', fmt=FMT_GRADE)
            book.put(ws, f'I{r}', f'=G{r}+H{r}*(D{r}-F{r})', fmt=FMT_EL)
            book.put(ws, f'J{r}', f'=G{r}+H{r}*(E{r}-F{r})', fmt=FMT_EL)
            book.put(ws, f'K{r}', f'=I{r}', fmt=FMT_EL)
            book.put(ws, f'L{r}', f'=J{r}', fmt=FMT_EL)
            book.put(ws, f'M{r}', f'=IF(AND(K{s}>=MIN(D{r},E{r})-1E-9,L{s}<=MAX(D{r},E{r})+1E-9),"PASS","FAIL")', align=CENTER, border=BOX)
            ctx['eb_ped_rows'][(tag, k)] = r
            row += 1
    check_formatting(ws, f'M{row-2*n}:M{row-1}')
    row += 1
    row = section(book, ws, row, 'C. Pier pedestal tops: Span 1 seat plane, straight transition between pad ends (E to F), Span 2 seat plane', 22)
    heads = ['Beam', 's A/C face (ft)', 's E pad end (ft)', 's F pad start (ft)', 's B/D face (ft)', 'EL A = EL C', 'EL E', 'EL F',
             'EL B = EL D', 'EL @ CL pier', 'Order check']
    row = header_row(book, ws, row, 2, heads, 32)
    ctx['pier_ped_rows'] = {}
    if last >= 2:
        for k in range(1, n + 1):
            r = row
            s1 = ctx['seat_rows'][('Pier', 1, k)]; s2 = ctx['seat_rows'][('Pier', 2, k)]
            book.put(ws, f'B{r}', f'Beam line {k}')
            book.put(ws, f'C{r}', '=Pier_Sta-Pier_PedL/24', fmt=FMT_STA)
            book.put(ws, f'D{r}', f'=L{s1}', fmt=FMT_STA)
            book.put(ws, f'E{r}', f'=K{s2}', fmt=FMT_STA)
            book.put(ws, f'F{r}', '=Pier_Sta+Pier_PedL/24', fmt=FMT_STA)
            book.put(ws, f'G{r}', f'=V{s1}+T{s1}*(C{r}-H{s1})', fmt=FMT_EL)
            book.put(ws, f'H{r}', f'=V{s1}+T{s1}*(D{r}-H{s1})', fmt=FMT_EL)
            book.put(ws, f'I{r}', f'=V{s2}+T{s2}*(E{r}-H{s2})', fmt=FMT_EL)
            book.put(ws, f'J{r}', f'=V{s2}+T{s2}*(F{r}-H{s2})', fmt=FMT_EL)
            book.put(ws, f'K{r}', f'=H{r}+(I{r}-H{r})*(Pier_Sta-D{r})/(E{r}-D{r})', fmt=FMT_EL)
            book.put(ws, f'L{r}', f'=IF(AND(C{r}<D{r},D{r}<E{r},E{r}<F{r},K{s1}>=C{r}-1E-9,L{s2}<=F{r}+1E-9),"PASS","FAIL")', align=CENTER, border=BOX)
            ctx['pier_ped_rows'][k] = r
            row += 1
        check_formatting(ws, f'L{row-n}:L{row-1}')
    row += 1
    row = section(book, ws, row, 'D. Bearing plate data (template method; slope from the notebook seat plane)', 22)
    heads = ['Brg plate mark', 'Span', 'Beam', 'Pad type', 'Beam end', 'Plan view case', 'Slope (%)', 'Beveled plate req.', 'C (in)', 'X (in)', 'Y (in)', 'Pad dims match type']
    row = header_row(book, ws, row, 2, heads, 32)
    ctx['plate_rows'] = {}
    for seat in seat_list(ctx):
        r = row; s = ctx['seat_rows'][(seat['tag'], seat['span'], seat['beam'])]
        end = 1 if seat['face'] == 'Start' else 2
        span_id = spans[seat['span'] - 1]['id']
        tag = seat['tag']; pad_l = 'Pier_PadL' if tag == 'Pier' else f'{tag}_PadL'
        pad_w = 'Pier_PadW' if tag == 'Pier' else f'{tag}_PadW'; pad_t = 'Pier_Pad' if tag == 'Pier' else f'{tag}_Pad'
        book.put(ws, f'B{r}', f'{span_id}-{seat["beam"]}-{end}')
        book.put(ws, f'C{r}', span_id); book.put(ws, f'D{r}', seat['beam'])
        book.put(ws, f'E{r}', '=Des_PadType', font=F_LINK); book.put(ws, f'F{r}', end)
        book.put(ws, f'G{r}', '=Des_PlanCase', font=F_LINK)
        book.put(ws, f'H{r}', f'=T{s}*100', fmt='0.000')
        book.put(ws, f'I{r}', f'=IF(ABS(T{s})>Des_BevelLimit,"YES","NO")', align=CENTER)
        book.put(ws, f'J{r}', f'=IF(I{r}="YES",IF(OR(E{r}="D",E{r}="E",E{r}="F",E{r}="G",E{r}="H",E{r}="J",E{r}="AA",E{r}="AB"),12,IF(E{r}="K",13.5,"N/A")),"N/A")')
        book.put(ws, f'K{r}', f'=IF(I{r}="NO","N/A",IF(T{s}>0,0.5+J{r}*T{s},0.5))')
        book.put(ws, f'L{r}', f'=IF(I{r}="NO","N/A",IF(T{s}<0,0.5-J{r}*T{s},0.5))')
        book.put(ws, f'M{r}', f'=IF(AND(ABS(INDEX(Pad_Table,MATCH(E{r},INDEX(Pad_Table,0,1),0),2)-{pad_l})<1E-9,'
                              f'ABS(INDEX(Pad_Table,MATCH(E{r},INDEX(Pad_Table,0,1),0),3)-{pad_w})<1E-9,'
                              f'ABS(INDEX(Pad_Table,MATCH(E{r},INDEX(Pad_Table,0,1),0),4)-{pad_t})<1E-9),"PASS","FAIL")', align=CENTER, border=BOX)
        ctx['plate_rows'][(seat['tag'], seat['span'], seat['beam'])] = r
        row += 1
    check_formatting(ws, f'M{row-len(ctx["plate_rows"])}:M{row-1}')
    ctx['ws_seats'] = ws
    return ws


# ============================================================ Substructure ==
def build_substructure(book: Book, ctx):
    m = ctx['model']; spans = ctx['spans']; n = m.count; last = len(spans)
    ws = book.sheet('Substructure', '1F2A36')
    ws.column_dimensions['A'].width = 3
    for k in range(2, 26):
        ws.column_dimensions[col(k)].width = 11.5
    ws.column_dimensions['B'].width = 13
    row = title_block(book, ws, 'Substructure', 'Cap tops from the lowest pedestal, pedestal heights, cap corners and pile cutoffs', 16)
    seats_ws = ctx['ws_seats']
    supports = [('Pier', 'Pier', 'pier'), ('End Bent 1', 'EB1', 'end_bent'), ('End Bent 2', 'EB2', 'end_bent')]
    row = section(book, ws, row, 'A. Cap placement', 22)
    heads = ['Support', 'Datum station', 'Cap CL station', 'Cap width (ft)', 'Cap depth (ft)', 'Piles', 'Pile size (in)', 'Pile spacing (ft)',
             'Pile end clear (in)', 'Cap length (ft)', 'Length check', 'Center offset', 'Cross slope', 'Min ped. (in)', 'Fixed top',
             'Lowest seat basis (ft)', 'Cap top (ft)', 'Cap underside (ft)', 'Cap volume (ft3)', 'Embedment (in)', 'Embedment check']
    row = header_row(book, ws, row, 2, heads, 42)
    ctx['sup_rows'] = {}
    for name, tag, kind in supports:
        r = row
        seats = [ctx['seat_rows'][key] for key in ctx['seat_rows'] if key[0] == tag]
        sx = lambda c, s: book.xref(seats_ws, f'{c}{s}')
        book.put(ws, f'B{r}', name)
        if tag == 'Pier':
            book.put(ws, f'C{r}', '=Pier_Sta', font=F_LINK, fmt=FMT_STA)
            book.put(ws, f'D{r}', '=Pier_Sta', font=F_LINK, fmt=FMT_STA)
            book.put(ws, f'I{r}', '=Pier_PileSpD*Pier_PileD/12', fmt=FMT_FT)
            book.put(ws, f'H{r}', '=Pier_PileD', font=F_LINK)
        else:
            datum, sign = ('Sp1_Sta0', '+') if tag == 'EB1' else (f'Sp{last}_Sta1', '-')
            book.put(ws, f'C{r}', f'={datum}', font=F_LINK, fmt=FMT_STA)
            book.put(ws, f'D{r}', f'=C{r}{sign}({tag}_Width/2-{tag}_BW/12-{tag}_Inset/12)', fmt=FMT_STA)
            book.put(ws, f'I{r}', f'={tag}_PileSp/12', fmt=FMT_FT)
            book.put(ws, f'H{r}', f'={tag}_PileW', font=F_LINK)
        book.put(ws, f'E{r}', f'={tag}_Width', font=F_LINK, fmt=FMT_FT)
        book.put(ws, f'F{r}', f'={tag}_Depth', font=F_LINK, fmt=FMT_FT)
        book.put(ws, f'G{r}', f'={tag}_NPile', font=F_LINK)
        book.put(ws, f'J{r}', f'={tag}_PileClr', font=F_LINK)
        book.put(ws, f'K{r}', f'=IF(ISNUMBER({tag}_LenFix),{tag}_LenFix,(G{r}-1)*I{r}+H{r}/12+2*J{r}/12)', fmt=FMT_FT)
        book.put(ws, f'L{r}', f'=IF(ISNUMBER({tag}_LenFix),IF(ABS({tag}_LenFix-((G{r}-1)*I{r}+H{r}/12+2*J{r}/12))<1E-8,"PASS","FAIL"),"PASS")', align=CENTER, border=BOX)
        book.put(ws, f'M{r}', f'={tag}_Ctr', font=F_LINK, fmt=FMT_FT)
        book.put(ws, f'N{r}', f'={tag}_Slope', font=F_LINK, fmt=FMT_GRADE)
        book.put(ws, f'O{r}', f'={tag}_Reveal', font=F_LINK)
        book.put(ws, f'P{r}', f'={tag}_TopFix', font=F_LINK, fmt=FMT_EL)
        terms = ','.join(f'{sx("V", s)}-N{r}*({sx("G", s)}-M{r})' for s in seats)
        book.put(ws, f'Q{r}', f'=MIN({terms})', fmt='0.000000')
        book.put(ws, f'R{r}', f'=IF(ISNUMBER({tag}_TopFix),{tag}_TopFix,Q{r}-O{r}/12)', fmt='0.000000', border=BOX)
        book.put(ws, f'S{r}', f'=R{r}-F{r}', fmt='0.000000')
        book.put(ws, f'T{r}', f'=E{r}*K{r}*F{r}', fmt='0.000')
        book.put(ws, f'U{r}', f'={tag}_Emb', font=F_LINK)
        book.put(ws, f'V{r}', f'=IF(AND(U{r}>=0,U{r}<=F{r}*12),"PASS","FAIL")', align=CENTER, border=BOX)
        for nm, c in (('Top', 'R'), ('Len', 'K'), ('CapSta', 'D'), ('PileSpFt', 'I'), ('Under', 'S')):
            book.define(f'{tag}_{nm}', ws, f'{c}{r}')
        ctx['sup_rows'][tag] = r
        row += 1
    check_formatting(ws, f'L{row-3}:L{row-1}'); check_formatting(ws, f'V{row-3}:V{row-1}')
    row += 1
    book.put(ws, f'B{row}', 'End caps: 235-in body derived from the pile layout (D-026); with the two 8-in cheek webs the cage '
             'envelope is 251 in (D-027). Backwalls, cheeks and lugs are in section F.', font=F_NOTE)
    row += 2
    row = section(book, ws, row, 'B. Pedestal heights at CL bearing', 22)
    heads = ['Support', 'Beam ID', 'Face', 'Offset (ft)', 'Ped. top @ CL brg', 'Local cap top', 'Pedestal height (in)', '>= minimum']
    row = header_row(book, ws, row, 2, heads, 32)
    ctx['ped_height_rows'] = {}
    for seat in seat_list(ctx):
        r = row; s = ctx['seat_rows'][(seat['tag'], seat['span'], seat['beam'])]; tag = seat['tag']; sr = ctx['sup_rows'][tag]
        book.put(ws, f'B{r}', seat['support']); book.put(ws, f'C{r}', f'{spans[seat["span"]-1]["id"]}-{seat["beam"]}')
        book.put(ws, f'D{r}', seat['face'])
        book.put(ws, f'E{r}', f"={book.xref(seats_ws, f'G{s}')}", font=F_LINK, fmt=FMT_FT)
        book.put(ws, f'F{r}', f"={book.xref(seats_ws, f'V{s}')}", font=F_LINK, fmt='0.000000')
        book.put(ws, f'G{r}', f'=R{sr}+N{sr}*(E{r}-M{sr})', fmt='0.000000')
        book.put(ws, f'H{r}', f'=(F{r}-G{r})*12', fmt='0.0000')
        book.put(ws, f'I{r}', f'=IF(H{r}>=O{sr}-1E-6,"PASS","FAIL")', align=CENTER, border=BOX)
        ctx['ped_height_rows'][(seat['tag'], seat['span'], seat['beam'])] = r
        row += 1
    check_formatting(ws, f'I{row-len(ctx["ped_height_rows"])}:I{row-1}')
    row += 1
    row = section(book, ws, row, 'C. Pedestal bevel and footprint checks', 22)
    heads = ['Support', 'Beam', 'Lowest pedestal top (ft)', 'Chamfer (in)', 'Cap top at pedestal edge', 'Bevel clears cap', 'Pedestal within cap length']
    row = header_row(book, ws, row, 2, heads, 32)
    first = row
    for name, tag, kind in supports:
        sr = ctx['sup_rows'][tag]
        for k in range(1, n + 1):
            r = row
            if tag == 'Pier':
                pr = ctx['pier_ped_rows'].get(k)
                low = f"MIN({','.join(book.xref(seats_ws, f'{c}{pr}') for c in 'GHIJ')})" if pr else 'NA()'
                chamf, pedw = 'Pier_Chamf', 'Pier_PedW'
            else:
                er = ctx['eb_ped_rows'][(tag, k)]
                low = f"MIN({book.xref(seats_ws, f'I{er}')},{book.xref(seats_ws, f'J{er}')})"
                chamf, pedw = f'{tag}_Chamf', f'{tag}_PedW'
            book.put(ws, f'B{r}', name); book.put(ws, f'C{r}', k)
            book.put(ws, f'D{r}', f'={low}', fmt='0.000000')
            book.put(ws, f'E{r}', f'={chamf}', font=F_LINK)
            book.put(ws, f'F{r}', f'=R{sr}+N{sr}*(Bm{k}_Off-M{sr})+ABS(N{sr})*{pedw}/24', fmt='0.000000')
            book.put(ws, f'G{r}', f'=IF(D{r}-E{r}/12>=F{r}-1E-9,"PASS","FAIL")', align=CENTER, border=BOX)
            book.put(ws, f'H{r}', f'=IF(AND(Bm{k}_Off-{pedw}/24>=M{sr}-K{sr}/2-1E-9,Bm{k}_Off+{pedw}/24<=M{sr}+K{sr}/2+1E-9),"PASS","FAIL")', align=CENTER, border=BOX)
            row += 1
    check_formatting(ws, f'G{first}:H{row-1}')
    ctx['bevel_check_range'] = book.xref(ws, f'G{first}:H{row-1}')
    row += 1
    row = section(book, ws, row, 'D. Cap corner elevations (transverse section)', 22)
    row = header_row(book, ws, row, 2, ['Support', 'Corner', 'Offset (ft)', 'Elevation (ft)'], 18)
    for name, tag, kind in supports:
        sr = ctx['sup_rows'][tag]
        for corner, side, bottom in (('TL', -1, False), ('TR', 1, False), ('BR', 1, True), ('BL', -1, True)):
            r = row
            book.put(ws, f'B{r}', name); book.put(ws, f'C{r}', corner)
            book.put(ws, f'D{r}', f'=M{sr}{"+" if side > 0 else "-"}K{sr}/2', fmt=FMT_FT)
            book.put(ws, f'E{r}', f'=R{sr}+N{sr}*(D{r}-M{sr})' + (f'-F{sr}' if bottom else ''), fmt='0.000000')
            row += 1
    row += 1
    row = section(book, ws, row, 'E. Pile layout and cutoff elevations (cutoff = local cap top - cap depth + embedment)', 22)
    heads = ['Support', 'Pile', 'Offset (ft)', 'Size (in)', 'Shape', 'Local cap top', 'Cap underside', 'Embedment (in)', 'Pile cutoff (ft)', 'Trial tip (ft)']
    row = header_row(book, ws, row, 2, heads, 32)
    ctx['pile_rows'] = {}
    for name, tag, kind in supports:
        sr = ctx['sup_rows'][tag]
        count = int(ctx['cfg']['pier']['pile_count'] if tag == 'Pier' else ctx['cfg']['end_bents']['start' if tag == 'EB1' else 'end']['pile_count'])
        for j in range(1, count + 1):
            r = row
            book.put(ws, f'B{r}', name); book.put(ws, f'C{r}', j)
            book.put(ws, f'D{r}', f'=M{sr}+({j}-(G{sr}+1)/2)*I{sr}', fmt=FMT_FT)
            book.put(ws, f'E{r}', f'=H{sr}')
            book.put(ws, f'F{r}', 'pipe' if tag == 'Pier' else 'square')
            book.put(ws, f'G{r}', f'=R{sr}+N{sr}*(D{r}-M{sr})', fmt='0.000000')
            book.put(ws, f'H{r}', f'=G{r}-F{sr}', fmt='0.000000')
            book.put(ws, f'I{r}', f'=U{sr}')
            book.put(ws, f'J{r}', f'=H{r}+I{r}/12', fmt='0.000000', border=BOX)
            book.put(ws, f'K{r}', '=J{0}-Phys_TrialPile'.format(r), fmt='0.000')
            ctx['pile_rows'][(tag, j)] = r
            row += 1
    row += 1
    row = end_bent_walls_section(book, ws, row, ctx)
    ctx['ws_sub'] = ws
    return ws


WALL_ROWS = [
    # key, label, formula(tag, datum, sign, R) -> str, unit, fmt, defined-name suffix
    ('rear', 'Cap rear face (toward the fill)', lambda t, d, s, R: f'={t}_CapSta{"-" if s < 0 else "+"}{t}_Width/2', 'ft', FMT_STA, 'RearSta'),
    ('lug', 'Rear-lug outer face = MSE coping stop', lambda t, d, s, R: f'={R("rear")}{"-" if s < 0 else "+"}{t}_LugP/12', 'ft', FMT_STA, 'LugFace'),
    ('bw0', 'Backwall from station', lambda t, d, s, R: f'={R("rear")}+{t}_Inset/12' if s < 0 else f'={R("rear")}-({t}_Inset+{t}_BW)/12', 'ft', FMT_STA, None),
    ('bw1', 'Backwall to station', lambda t, d, s, R: f'={R("rear")}+({t}_Inset+{t}_BW)/12' if s < 0 else f'={R("rear")}-{t}_Inset/12', 'ft', FMT_STA, None),
    ('ffbw', 'Backwall front face is the FFBW', lambda t, d, s, R: f'=IF(ABS({R("bw1") if s < 0 else R("bw0")}-{d})<1E-8,"PASS","FAIL")', '', None, None),
    ('yl', 'Cap end offset, left', lambda t, d, s, R: f'={t}_Ctr-{t}_Len/2', 'ft', FMT_FT, None),
    ('yr', 'Cap end offset, right', lambda t, d, s, R: f'={t}_Ctr+{t}_Len/2', 'ft', FMT_FT, None),
    ('out', 'Out-to-out with cheek webs', lambda t, d, s, R: f'={t}_Len*12+2*{t}_CkT', 'in', '0.000', 'OutIn'),
    ('pgl', 'PGL at the FFBW', lambda t, d, s, R: '=' + PGL(d), 'ft', FMT_EL, None),
    ('btl', 'Backwall top, left cap end', lambda t, d, s, R: f'={R("pgl")}+{DELTA(R("yl"))}-{t}_BwDrop/12', 'ft', FMT_EL, 'BwTopL'),
    ('bt0', 'Backwall top at the PGL', lambda t, d, s, R: f'={R("pgl")}+{DELTA("0")}-{t}_BwDrop/12', 'ft', FMT_EL, 'BwTop0'),
    ('btr', 'Backwall top, right cap end', lambda t, d, s, R: f'={R("pgl")}+{DELTA(R("yr"))}-{t}_BwDrop/12', 'ft', FMT_EL, 'BwTopR'),
    ('ctl', 'Cap top, left cap end', lambda t, d, s, R: f'={t}_Top+{t}_Slope*({R("yl")}-{t}_Ctr)', 'ft', FMT_EL, None),
    ('ct0', 'Cap top at the PGL', lambda t, d, s, R: f'={t}_Top+{t}_Slope*(0-{t}_Ctr)', 'ft', FMT_EL, None),
    ('ctr', 'Cap top, right cap end', lambda t, d, s, R: f'={t}_Top+{t}_Slope*({R("yr")}-{t}_Ctr)', 'ft', FMT_EL, None),
    ('hl', 'Backwall height, left end', lambda t, d, s, R: f'={R("btl")}-{R("ctl")}', 'ft', FMT_FT, None),
    ('h0', 'Backwall height at the PGL', lambda t, d, s, R: f'={R("bt0")}-{R("ct0")}', 'ft', FMT_FT, 'BwH0'),
    ('hr', 'Backwall height, right end', lambda t, d, s, R: f'={R("btr")}-{R("ctr")}', 'ft', FMT_FT, None),
    ('bwv', 'Backwall concrete', lambda t, d, s, R: (f'={t}_BW/12*IF(AND({R("yl")}<0,{R("yr")}>0),(0-{R("yl")})*({R("hl")}+{R("h0")})/2'
                                                    f'+{R("yr")}*({R("h0")}+{R("hr")})/2,({R("yr")}-{R("yl")})*({R("hl")}+{R("hr")})/2)'), 'ft3', '0.0000', 'BwVol'),
    ('zbl', 'Cheek bottom, left (cap underside at cap end)', lambda t, d, s, R: f'={R("ctl")}-{t}_Depth', 'ft', FMT_EL, 'CkBotL'),
    ('zbr', 'Cheek bottom, right', lambda t, d, s, R: f'={R("ctr")}-{t}_Depth', 'ft', FMT_EL, 'CkBotR'),
    ('lbl', 'Lug bottom, left', lambda t, d, s, R: f'={R("zbl")}-{t}_LugP/12', 'ft', FMT_EL, 'LugBotL'),
    ('lbr', 'Lug bottom, right', lambda t, d, s, R: f'={R("zbr")}-{t}_LugP/12', 'ft', FMT_EL, 'LugBotR'),
    ('ckl', 'Cheek top, left', lambda t, d, s, R: f'=IF(ISNUMBER({t}_CkTopL),{t}_CkTopL,{R("pgl")}+{DELTA("Deck_YL")})', 'ft', FMT_EL, 'CkTopLv'),
    ('ckr', 'Cheek top, right', lambda t, d, s, R: f'=IF(ISNUMBER({t}_CkTopR),{t}_CkTopR,{R("pgl")}+{DELTA("Deck_YR")})', 'ft', FMT_EL, 'CkTopRv'),
    ('cvl', 'Cheek concrete, left (web + bottom lug + rear lug)', lambda t, d, s, R: (
        f'={t}_Width*{t}_CkT/12*({R("ckl")}-{R("zbl")})+{t}_Width*{t}_LugW/12*{t}_LugP/12+{t}_LugP/12*{t}_LugW/12*({R("ckl")}-{R("zbl")}+{t}_LugP/12)'), 'ft3', '0.0000', 'CkVolL'),
    ('cvr', 'Cheek concrete, right', lambda t, d, s, R: (
        f'={t}_Width*{t}_CkT/12*({R("ckr")}-{R("zbr")})+{t}_Width*{t}_LugW/12*{t}_LugP/12+{t}_LugP/12*{t}_LugW/12*({R("ckr")}-{R("zbr")}+{t}_LugP/12)'), 'ft3', '0.0000', 'CkVolR'),
    ('chk1', 'Cheek tops above the cap', lambda t, d, s, R: f'=IF(AND({R("ckl")}>{R("ctl")},{R("ckr")}>{R("ctr")}),"PASS","FAIL")', '', None, None),
    ('chk2', 'Cheeks enclose the backwall ends', lambda t, d, s, R: f'=IF(AND({R("ckl")}>={R("btl")}-1E-9,{R("ckr")}>={R("btr")}-1E-9),"PASS","FAIL")', '', None, None),
    ('chk3', 'Lug width within the cheek', lambda t, d, s, R: f'=IF(AND({t}_LugW>0,{t}_LugW<={t}_CkT,{t}_LugP>=0),"PASS","FAIL")', '', None, None),
]


def end_bent_walls_section(book: Book, ws, row: int, ctx):
    last = len(ctx['spans'])
    row = section(book, ws, row, 'F. End-bent backwalls, cheek webs and lugs (D-026/D-027/D-033)', 22)
    book.put(ws, f'B{row}', 'Quantity', font=F_HEAD, fill=FILL_HEAD, align=CENTER, border=BOX)
    ws.merge_cells(f'B{row}:E{row}')
    row = header_row(book, ws, row, 6, ['End Bent 1', 'End Bent 2', 'Unit'], 18)
    rows = {key: row + k for k, (key, *_rest) in enumerate(WALL_ROWS)}
    ctx['wall_rows'] = rows
    for tag, c, datum, sign in (('EB1', 'F', 'Sp1_Sta0', -1), ('EB2', 'G', f'Sp{last}_Sta1', 1)):
        R = lambda key, c=c: f'{c}{rows[key]}'
        for key, label, formula, unit, fmt, name in WALL_ROWS:
            r = rows[key]
            if tag == 'EB1':
                book.put(ws, f'B{r}', label)
                ws.merge_cells(f'B{r}:E{r}')
                book.put(ws, f'H{r}', unit)
            book.put(ws, f'{c}{r}', formula(tag, datum, sign, R), fmt=fmt, border=BOX,
                     align=CENTER if fmt is None else None)
            if name:
                book.define(f'{tag}_{name}', ws, f'{c}{r}')
    first = row
    row += len(WALL_ROWS)
    check_formatting(ws, f'F{first}:G{row-1}')
    book.put(ws, f'B{row}', 'Backwall top = deck-top line at the FFBW less the landing seat; cheek tops are the D-033 barrier-return '
             'bases (MSE interface inputs). The CAD EB1 backwall seat follows the MSE landing and is to be rebuilt from this geometry.', font=F_NOTE)
    return row + 2


# ============================================================== Beam Data ==
def build_beam_data(book: Book, ctx):
    m = ctx['model']; spans = ctx['spans']; n = m.count
    ws = book.sheet('Beam Data', '1F2A36')
    ws.column_dimensions['A'].width = 3
    for k in range(2, 26):
        ws.column_dimensions[col(k)].width = 11
    ws.column_dimensions['B'].width = 14
    row = title_block(book, ws, 'Beam Data', 'Bars 5K spacing zones (Mathcad stirrup layout) and FIB standard bar dimensions (template method)', 16)
    row = section(book, ws, row, 'A. Bars 5K zones by span: Mathcad stirrup zones mapped to the FDOT S*, Sj, S1-S4 notation', 20)
    heads = ['Span', "S* spaces", "V* (in)", 'Sj spaces', 'Vj (in)', 'S1', 'V1 (in)', 'S2', 'V2 (in)', 'S3', 'V3 (in)', 'V4 (in)', 'First bar from end (in)', 'A1/A2 spacing equal']
    row = header_row(book, ws, row, 2, heads, 32)
    ctx['zone_rows'] = {}
    for i, span in enumerate(spans, start=1):
        r = row
        book.put(ws, f'B{r}', span['id'])
        book.put(ws, f'C{r}', f'=MC{i}_A1_n+MC{i}_A2_n'); book.put(ws, f'D{r}', f'=MC{i}_A1_s', font=F_LINK)
        book.put(ws, f'E{r}', f'=MC{i}_A3_n', font=F_LINK); book.put(ws, f'F{r}', f'=MC{i}_A3_s', font=F_LINK)
        for c1, c2, z in (('G', 'H', 'S1'), ('I', 'J', 'S2'), ('K', 'L', 'S3')):
            book.put(ws, f'{c1}{r}', f'=MC{i}_{z}_n', font=F_LINK); book.put(ws, f'{c2}{r}', f'=MC{i}_{z}_s', font=F_LINK)
        book.put(ws, f'M{r}', f'=MC{i}_S4_s', font=F_LINK)
        book.put(ws, f'N{r}', f'=MC{i}_EndSp', font=F_LINK)
        book.put(ws, f'O{r}', f'=IF(ABS(MC{i}_A1_s-MC{i}_A2_s)<1E-9,"PASS","FAIL")', align=CENTER, border=BOX)
        ctx['zone_rows'][i] = r
        row += 1
    check_formatting(ws, f'O{row-len(spans)}:O{row-1}')
    row += 1
    row = section(book, ws, row, 'B. Per beam: S4 count, total Bars 5K and standard bar dimensions (square ends)', 24)
    heads = ['Beam', 'Dim L (ft)', 'Half length (in)', 'Zones S*..S3 (in)', 'S4 spaces', 'Center space (in)', 'Center <= V4', 'Total 5K',
             '3C1 C (in)', '3C2 C (in)', '3D1 A (in)', '3D1 B (in)', '3D2 A (in)', '3D2 B (in)', '3D3 No.', '4M1 D (in)', '4M2 D (in)', '4M3 No.']
    row = header_row(book, ws, row, 2, heads, 32)
    ctx['bar_rows'] = {}
    bws = ctx['ws_beam']
    for i, span in enumerate(spans, start=1):
        z = ctx['zone_rows'][i]
        for k in range(1, n + 1):
            r = row; br = ctx['beam_rows'][(i, k)]
            book.put(ws, f'B{r}', f'{span["id"]}-{k}')
            book.put(ws, f'C{r}', f"={book.xref(bws, f'S{br}')}", font=F_LINK, fmt=FMT_FT)
            book.put(ws, f'D{r}', f'=C{r}*12/2', fmt='0.00')
            book.put(ws, f'E{r}', f'=N{z}+C{z}*D{z}+E{z}*F{z}+G{z}*H{z}+I{z}*J{z}+K{z}*L{z}', fmt='0.00')
            book.put(ws, f'F{r}', f'=MAX(0,ROUNDDOWN((D{r}-E{r})/M{z},0))')
            book.put(ws, f'G{r}', f'=(D{r}-E{r}-F{r}*M{z})*2', fmt='0.00')
            book.put(ws, f'H{r}', f'=IF(G{r}<=M{z}+1E-9,"PASS","CHECK")', align=CENTER, border=BOX)
            book.put(ws, f'I{r}', f'=2*(1+C{z}+E{z}+G{z}+I{z}+K{z}+F{r})')
            phi1 = f"{book.xref(bws, f'I{br}')}"; phi2 = f"{book.xref(bws, f'J{br}')}"
            book.put(ws, f'J{r}', f'=Bar_C/SIN(RADIANS({phi1}))', fmt='0.00')
            book.put(ws, f'K{r}', f'=Bar_C/SIN(RADIANS({phi2}))', fmt='0.00')
            book.put(ws, f'L{r}', f'=Bar_DA/SIN(RADIANS({phi1}))', fmt='0.00')
            book.put(ws, f'M{r}', f'=Bar_DB/SIN(RADIANS({phi1}))', fmt='0.00')
            book.put(ws, f'N{r}', f'=Bar_DA/SIN(RADIANS({phi2}))', fmt='0.00')
            book.put(ws, f'O{r}', f'=Bar_DB/SIN(RADIANS({phi2}))', fmt='0.00')
            book.put(ws, f'P{r}', f'=(1+ROUNDUP((C{r}*12-2*(N{z}+C{z}*D{z}+E{z}*F{z}+Bar_D3s))/Bar_D3s,0))*2')
            book.put(ws, f'Q{r}', f'=Bar_MD/SIN(RADIANS({phi1}))', fmt='0.00')
            book.put(ws, f'R{r}', f'=Bar_MD/SIN(RADIANS({phi2}))', fmt='0.00')
            book.put(ws, f'S{r}', f'=1+CEILING((C{r}*12-2*(3.5+2*7+Bar_Mk*6+12))/12,1)')
            ctx['bar_rows'][(i, k)] = r
            row += 1
    check_formatting(ws, f'H{row-len(ctx["bar_rows"])}:H{row-1}')
    ctx['ws_bars'] = ws
    return ws


# ============================================================== Clearance ==
def build_clearance(book: Book, ctx):
    m = ctx['model']; spans = ctx['spans']; n = m.count
    ws = book.sheet('Clearance', '1F2A36')
    ws.column_dimensions['A'].width = 3
    for k in range(2, 26):
        ws.column_dimensions[col(k)].width = 11
    row = title_block(book, ws, 'Vertical Clearance', 'Beam bottom minus roadway profile, minimized exactly on every roadway segment under a beam', 16)
    row = section(book, ws, row, 'A. Result', 12)
    res_row = row
    row += 7
    row = section(book, ws, row, 'B. Roadway segments from Geometry Report.xml (as parsed by the notebook)', 12)
    heads = ['Seg', 'XML start', 'XML end', 'Bridge start (ft)', 'Bridge end (ft)', 'z start (ft)', 'Grade', 'Quad (1/ft)', 'Overlap with bridge (ft)']
    row = header_row(book, ws, row, 2, heads, 32)
    seg_first = row
    shift = float(ctx['cfg']['roadway']['bridge_station_at_xml_zero_ft'])
    for j, seg in enumerate(ctx['road'], start=1):
        r = row
        book.put(ws, f'B{r}', j)
        book.put(ws, f'C{r}', seg['start'] - shift, font=F_INPUT, fill=FILL_INPUT, fmt='0.000000')
        book.put(ws, f'D{r}', seg['end'] - shift, font=F_INPUT, fill=FILL_INPUT, fmt='0.000000')
        book.put(ws, f'E{r}', f'=C{r}+Rd_Shift', fmt=FMT_STA)
        book.put(ws, f'F{r}', f'=D{r}+Rd_Shift', fmt=FMT_STA)
        book.put(ws, f'G{r}', seg['z'], font=F_INPUT, fill=FILL_INPUT, fmt='0.000000')
        book.put(ws, f'H{r}', seg['g'], font=F_INPUT, fill=FILL_INPUT, fmt='0.00000000')
        book.put(ws, f'I{r}', seg['q'], font=F_INPUT, fill=FILL_INPUT, fmt='0.00000000')
        book.put(ws, f'J{r}', f'=MAX(0,MIN(F{r},Prof_Sta1)-MAX(E{r},Prof_Sta0))', fmt='0.000000')
        row += 1
    seg_last = row - 1
    ctx['clear_segs'] = (seg_first, seg_last)
    row += 1
    row = section(book, ws, row, 'C. Clearance on each roadway segment piece under a beam (exact minimum of a quadratic)', 20)
    heads = ['Piece', 'Span', 'Seg', 's a', 's b', 's mid', 'road a', 'road mid', 'road b', 'bottom a', 'bottom mid', 'bottom b',
             'c a', 'c mid', 'c b', 'A', 'B', 't*', 'c*', 'c min (ft)', 's at min (ft)']
    row = header_row(book, ws, row, 2, heads, 32)
    piece_first = row
    raise_min = 'MIN(' + ','.join(f'Bm{k}_Raise' for k in range(1, n + 1)) + ')'
    piece = 0
    for i, span in enumerate(spans, start=1):
        lo, hi = m.beam_limits[i - 1]
        for j, seg in enumerate(ctx['road'], start=1):
            if not (seg['end'] > lo and seg['start'] < hi):
                continue
            piece += 1; r = row; sr = seg_first + j - 1
            book.put(ws, f'B{r}', piece); book.put(ws, f'C{r}', span['id']); book.put(ws, f'D{r}', j)
            book.put(ws, f'E{r}', f'=MAX(E{sr},Sp{i}_Bm0)', fmt=FMT_STA)
            book.put(ws, f'F{r}', f'=MIN(F{sr},Sp{i}_Bm1)', fmt=FMT_STA)
            book.put(ws, f'G{r}', f'=(E{r}+F{r})/2', fmt=FMT_STA)
            for c, s in (('H', 'E'), ('I', 'G'), ('J', 'F')):
                book.put(ws, f'{c}{r}', f'=$G{sr}+$H{sr}*({s}{r}-$E{sr})+$I{sr}*({s}{r}-$E{sr})^2', fmt='0.000000')
            for c, s in (('K', 'E'), ('L', 'G'), ('M', 'F')):
                t = f'({s}{r}-Sp{i}_BrgA)/Sp{i}_Lb'
                book.put(ws, f'{c}{r}', '=' + BOTTOM(i, t, raise_min), fmt='0.000000')
            book.put(ws, f'N{r}', f'=K{r}-H{r}', fmt='0.000000'); book.put(ws, f'O{r}', f'=L{r}-I{r}', fmt='0.000000')
            book.put(ws, f'P{r}', f'=M{r}-J{r}', fmt='0.000000')
            book.put(ws, f'Q{r}', f'=2*(P{r}-2*O{r}+N{r})', fmt='0.000000000')
            book.put(ws, f'R{r}', f'=P{r}-N{r}-Q{r}', fmt='0.000000000')
            book.put(ws, f'S{r}', f'=IF(Q{r}>1E-12,-R{r}/(2*Q{r}),-1)', fmt='0.000000')
            book.put(ws, f'T{r}', f'=IF(AND(S{r}>0,S{r}<1),N{r}+R{r}*S{r}+Q{r}*S{r}^2,N{r})', fmt='0.000000')
            book.put(ws, f'U{r}', f'=IF(F{r}>E{r},MIN(N{r},P{r},T{r}),"")', fmt='0.000000')
            book.put(ws, f'V{r}', f'=IF(U{r}="","",IF(U{r}=N{r},E{r},IF(U{r}=P{r},F{r},E{r}+(F{r}-E{r})*S{r})))', fmt=FMT_STA)
            row += 1
    piece_last = row - 1
    r = res_row
    rows = [
        ('Minimum vertical clearance (ft)', f'=MIN(U{piece_first}:U{piece_last})', 'Clr_Min', '0.0000'),
        ('Station of minimum (ft)', f'=INDEX(V{piece_first}:V{piece_last},MATCH(Clr_Min,U{piece_first}:U{piece_last},0))', 'Clr_Sta', FMT_STA),
        ('Required clearance (ft)', '=Rd_Req', None, '0.000'),
        ('Margin (ft)', '=Clr_Min-Rd_Req', 'Clr_Margin', '0.0000'),
        ('Roadway coverage of the bridge length (ft)', f'=SUM(J{seg_first}:J{seg_last})', 'Clr_Cover', '0.000'),
        ('Status', '=IF(Clr_Min<Rd_Req-1E-8,"Below requirement",IF(ABS(Clr_Cover-Prof_L)<=1E-6,"Meets requirement","Covered region only"))', 'Clr_Status', None),
    ]
    for label, formula, name, fmt in rows:
        book.put(ws, f'B{r}', label, bold=True)
        book.put(ws, f'F{r}', formula, fmt=fmt, border=BOX, bold=True)
        if name:
            book.define(name, ws, f'F{r}')
        r += 1
    book.put(ws, f'H{res_row}', 'Coverage note: the XML profile ends before the bridge does; the notebook reports '
             '"Covered region only" and never extrapolates. Clearance is evaluated only where a beam exists.', font=F_NOTE)
    return ws


# ============================================================== CAD tabs ==
def cad_header(book: Book, ws, title, subtitle, table_date=None, width=14):
    book.put(ws, 'A1', title, font=F_TITLE)
    book.put(ws, 'A2', subtitle, font=F_NOTE)
    labels = [('JOB #: ', 'TB_JobNo'), ('PROJECT NAME: ', 'TB_Project'), ('SUBJECT: ', f'"{title}"'),
              ('CALCULATED BY: ', 'TB_CalcBy&"   "&TB_CalcDate'), ('CHECKED BY: ', 'TB_CheckBy&"   "&TB_CheckDate')]
    for k, (lab, f) in enumerate(labels):
        book.put(ws, f'{col(width)}{1+k}', f'="{lab}"&{f}', font=F_HEAD)
    if table_date:
        book.put(ws, f'A{5}', f'TABLE DATE {table_date}', font=F_NOTE)
    return 7


def cad_cell(book, ws, ref, formula, fmt=None):
    return book.put(ws, ref, formula, fmt=fmt, align=CENTER, border=BOX, font=F_LINK)


def build_cad_fge(book: Book, ctx):
    spans = ctx['spans']; fge_ws = book.wb['FGE']
    ws = book.sheet('CAD FGE', 'C0392B')
    ws.column_dimensions['A'].width = 30; ws.column_dimensions['B'].width = 10
    for k in range(3, 3 + FGE_INCREMENTS + 3):
        ws.column_dimensions[col(k)].width = 10.5
    row = cad_header(book, ws, 'FINISH GRADE ELEVATIONS', 'Place these tables into CADD. Elevations in feet at top of deck.', None, 14)
    ctx['cad_fge'] = {}
    for i, span in enumerate(spans, start=1):
        f = ctx['fge'][i]
        book.put(ws, f'A{row}', f'FINISH GRADE ELEVATIONS - SPAN {span["id"]}', font=F_H2)
        row += 1
        heads = ['LOCATION', 'INCR. DIST (ft)'] + [lab.upper() for lab in f['labels']]
        row = header_row(book, ws, row, 1, heads, 30)
        book.put(ws, f'A{row}', 'STATION', font=F_HEAD, border=BOX, align=CENTER)
        cad_cell(book, ws, f'B{row}', f"={book.xref(fge_ws, f['inc'])}", '0.000')
        for k, c in enumerate(f['cols']):
            cad_cell(book, ws, f'{col(3+k)}{row}', f"={book.xref(fge_ws, f'{c}{f['label_row']}')}")
        row += 1
        for name, lr in f['lines'].items():
            book.put(ws, f'A{row}', name.upper(), font=F_HEAD, border=BOX)
            book.put(ws, f'B{row}', '', border=BOX)
            for k, c in enumerate(f['cols']):
                cad_cell(book, ws, f'{col(3+k)}{row}', f"={book.xref(fge_ws, f'{c}{lr}')}", '0.000')
            ctx['cad_fge'][(i, name)] = row
            row += 1
        book.put(ws, f'A{row}', 'FFBW columns are the PGL extended; the deck ends 3 in inside each FFBW.', font=F_NOTE)
        row += 2
    return ws


def build_cad_buildup(book: Book, ctx):
    spans = ctx['spans']; n = ctx['model'].count; cws = book.wb['Camber & Haunch']
    ws = book.sheet('CAD Build-Up', 'C0392B')
    for c, w in zip('ABCDEFGHI', (10, 10, 13, 13, 13, 17, 17, 18, 12)):
        ws.column_dimensions[c].width = w
    row = cad_header(book, ws, 'BUILD-UP & DEFLECTION DATA TABLE FOR PRESTRESSED I-BEAMS', 'Inches. Build-up over CL beam; deflections + up (Mathcad).', '07/01/17', 11)
    book.put(ws, f'A{row}', 'LOCATION', font=F_HEAD, fill=FILL_HEAD, border=BOX, align=CENTER)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    book.put(ws, f'C{row}', 'REQUIRED THEORETICAL BUILD-UP OVER CL BEAM', font=F_HEAD, fill=FILL_HEAD, border=BOX, align=CENTER)
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
    for c, t in zip('FGHI', ('NET BEAM CAMBER (PRESTRESS - DEAD LOAD OF BEAM) @ RELEASE', 'NET BEAM CAMBER (PRESTRESS - DEAD LOAD OF BEAM) @ 120 DAYS',
                             'DEAD LOAD DEFLECTION DURING DECK POUR @ 120 DAYS DIM A', 'BUILD-UP CASE NO.')):
        book.put(ws, f'{c}{row}', t, font=F_HEAD, fill=FILL_HEAD, border=BOX, align=CENTER)
        ws.merge_cells(f'{c}{row}:{c}{row+1}')
    ws.row_dimensions[row].height = 45
    row += 1
    for c, t in zip('ABCDE', ('SPAN NO.', 'BEAM NO.', 'AT BEGIN SPAN DIM B', 'AT CL SPAN DIM C', 'AT END SPAN DIM D')):
        book.put(ws, f'{c}{row}', t, font=F_HEAD, fill=FILL_HEAD, border=BOX, align=CENTER)
    ws.row_dimensions[row].height = 30
    row += 1
    ctx['cad_buildup'] = {}
    for i, span in enumerate(spans, start=1):
        for k in range(1, n + 1):
            br = ctx['buildup_rows'][(i, k)]
            book.put(ws, f'A{row}', str(span['id']), border=BOX, align=CENTER)
            book.put(ws, f'B{row}', f'{span["id"]}-{k}', border=BOX, align=CENTER)
            for c, src in zip('CDE', 'CDE'):
                cad_cell(book, ws, f'{c}{row}', f"={book.xref(cws, f'{src}{br}')}", '0.00')
            cad_cell(book, ws, f'F{row}', f'=MC{i}_Drel', '0.00')
            cad_cell(book, ws, f'G{row}', f'=MC{i}_D120', '0.00')
            cad_cell(book, ws, f'H{row}', f'=ABS(MC{i}_Dnc)', '0.00')
            cad_cell(book, ws, f'I{row}', f"={book.xref(cws, f'H{br}')}")
            ctx['cad_buildup'][(i, k)] = row
            row += 1
    row += 1
    book.put(ws, f'A{row}', 'Dim A is the non-composite (deck pour) deflection. The build-up uses the Mathcad NetCamber = 120-day camber '
             '+ non-composite DL + composite DL (barrier) deflections, so G - H differs from the camber used by the composite DL deflection.', font=F_NOTE)
    return ws


def build_cad_beam_variables(book: Book, ctx):
    spans = ctx['spans']; n = ctx['model'].count; bws = ctx['ws_beam']; bars = ctx['ws_bars']
    ws = book.sheet('CAD Beam Variables', 'C0392B')
    widths = [7, 8, 9, 7, 8, 8, 8, 7, 6, 6, 9, 9, 7, 7, 6, 6, 6, 6, 6, 6, 6, 7, 7, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6, 6]
    for k, w in enumerate(widths, start=1):
        ws.column_dimensions[col(k)].width = w
    row = cad_header(book, ws, 'FLORIDA-I BEAM - TABLE OF BEAM VARIABLES', 'Dimensions: ft-in split as feet text + inches, as in the BCC template.', '01-01-16', 30)
    top = row
    groups = [('LOCATION', 1, 2), ('BEAM TYPE', 3, 3), ('CONCRETE PROPERTIES', 4, 6), ('STND. PTRN. TYPE', 7, 7), ('END ELEV. COND.', 8, 8),
              ('PLAN VIEW CASE', 9, 10), ('BRG. PLATE MARK', 11, 12), ('END OF BEAM & BEARING DIMENSIONS', 13, 20), ('BEAM DIMENSIONS', 21, 23),
              ('REINFORCING STEEL', 24, 47)]
    for label, a, b in groups:
        book.put(ws, f'{col(a)}{top}', label, font=F_HEAD, fill=FILL_HEAD, border=BOX, align=CENTER)
        if b > a:
            ws.merge_cells(start_row=top, start_column=a, end_row=top, end_column=b)
    second = [(1, 'SPAN NO.'), (2, 'BEAM NO.'), (4, 'CLASS'), (5, 'STRENGTHS (psi)'), (9, 'END 1'), (10, 'END 2'), (11, 'END 1'), (12, 'END 2'),
              (13, 'ANGLE PHI'), (15, 'DIM P'), (16, 'DIM J'), (17, 'DIM K1'), (19, 'DIM K2'), (21, 'DIM L'), (23, 'DIM R'),
              (24, '3C1'), (26, '3C2'), (28, '3D1'), (32, '3D2'), (36, '3D3'), (37, '4M1'), (39, '4M2'), (41, '4M3'), (42, '5K'),
              (43, 'NO. OF SPACES BARS 5K'), (47, 'SP. BARS 5K')]
    for c, label in second:
        book.put(ws, f'{col(c)}{top+1}', label, font=F_HEAD, fill=FILL_HEAD, border=BOX, align=CENTER)
    for a, b in ((5, 6), (13, 14), (17, 18), (19, 20), (21, 22), (24, 25), (26, 27), (28, 31), (32, 35), (37, 38), (39, 40), (43, 46)):
        ws.merge_cells(start_row=top + 1, start_column=a, end_row=top + 1, end_column=b)
    third = [(5, '28 DAY'), (6, 'RELEASE'), (13, 'END 1'), (14, 'END 2'), (21, 'ft'), (22, 'in'), (24, 'C'), (26, 'C'), (28, 'A'), (30, 'B'),
             (32, 'A'), (34, 'B'), (36, 'NO.'), (37, 'D'), (39, 'D'), (41, 'NO.'), (42, 'NO.'), (43, 'S1'), (44, 'S2'), (45, 'S3'), (46, 'S4'), (47, 'V1')]
    for c, label in third:
        book.put(ws, f'{col(c)}{top+2}', label, font=F_HEAD, fill=FILL_HEAD, border=BOX, align=CENTER)
    for a, b in ((24, 25), (26, 27), (28, 29), (30, 31), (32, 33), (34, 35), (37, 38), (39, 40)):
        ws.merge_cells(start_row=top + 2, start_column=a, end_row=top + 2, end_column=b)
    for r in (top, top + 1, top + 2):
        ws.row_dimensions[r].height = 30
    row = top + 3
    ctx['cad_bv'] = {}

    def ftin(c_ft, c_in, expr):
        cad_cell(book, ws, f'{col(c_ft)}{row}', f'=TRUNC({expr})&"\'-"')
        cad_cell(book, ws, f'{col(c_in)}{row}', f'=({expr}-TRUNC({expr}))*12', 'General')

    for i, span in enumerate(spans, start=1):
        for k in range(1, n + 1):
            br = ctx['beam_rows'][(i, k)]; bar = ctx['bar_rows'][(i, k)]; z = ctx['zone_rows'][i]
            B = lambda c: book.xref(bws, f'{c}{br}')
            R = lambda c: book.xref(bars, f'{c}{bar}')
            Z = lambda c: book.xref(bars, f'{c}{z}')
            book.put(ws, f'A{row}', str(span['id']), border=BOX, align=CENTER)
            book.put(ws, f'B{row}', f'{span["id"]}-{k}', border=BOX, align=CENTER)
            cad_cell(book, ws, f'C{row}', '=Bm_Type')
            cad_cell(book, ws, f'D{row}', '=Des_ConcClass')
            cad_cell(book, ws, f'E{row}', f'=MC{i}_fc*1000', '0')
            cad_cell(book, ws, f'F{row}', f'=MC{i}_fci*1000', '0')
            cad_cell(book, ws, f'G{row}', '=Des_Pattern')
            cad_cell(book, ws, f'H{row}', f'={B("O")}')
            cad_cell(book, ws, f'I{row}', '=Des_PlanCase'); cad_cell(book, ws, f'J{row}', '=Des_PlanCase')
            cad_cell(book, ws, f'K{row}', f'="{span["id"]}-{k}-1"'); cad_cell(book, ws, f'L{row}', f'="{span["id"]}-{k}-2"')
            cad_cell(book, ws, f'M{row}', f'={B("I")}', 'General'); cad_cell(book, ws, f'N{row}', f'={B("J")}', 'General')
            cad_cell(book, ws, f'O{row}', f'={B("N")}', 'General')
            cad_cell(book, ws, f'P{row}', f'={B("K")}', 'General')
            ftin(17, 18, f'{B("L")}/12')
            ftin(19, 20, f'{B("M")}/12')
            cad_cell(book, ws, f'U{row}', f'={B("T")}', '0')
            cad_cell(book, ws, f'V{row}', f'={B("U")}', '0.00')
            cad_cell(book, ws, f'W{row}', f'={B("R")}', '0.00')
            ftin(24, 25, f'{R("J")}/12')
            ftin(26, 27, f'{R("K")}/12')
            ftin(28, 29, f'{R("L")}/12'); ftin(30, 31, f'{R("M")}/12')
            ftin(32, 33, f'{R("N")}/12'); ftin(34, 35, f'{R("O")}/12')
            cad_cell(book, ws, f'{col(36)}{row}', f'={R("P")}')
            ftin(37, 38, f'{R("Q")}/12'); ftin(39, 40, f'{R("R")}/12')
            cad_cell(book, ws, f'{col(41)}{row}', f'={R("S")}')
            cad_cell(book, ws, f'{col(42)}{row}', f'={R("I")}')
            cad_cell(book, ws, f'{col(43)}{row}', f'={Z("G")}'); cad_cell(book, ws, f'{col(44)}{row}', f'={Z("I")}')
            cad_cell(book, ws, f'{col(45)}{row}', f'={Z("K")}'); cad_cell(book, ws, f'{col(46)}{row}', f'={R("F")}')
            cad_cell(book, ws, f'{col(47)}{row}', f'={Z("H")}', 'General')
            ctx['cad_bv'][(i, k)] = row
            row += 1
    row += 1
    for note in ('Dim L = sloped beam length + shortening R (BCC template rule) - confirm against the current FDOT instructions before issue.',
                 'Bar dimensions use the BCC template constants for square ends; 3D3 and S4 use the Mathcad first-stirrup distance (EndSpace).',
                 'Strand pattern type and concrete class are designer assignments on the Inputs tab.'):
        book.put(ws, f'A{row}', note, font=F_NOTE); row += 1
    return ws


def build_cad_plates(book: Book, ctx):
    sws = ctx['ws_seats']
    ws = book.sheet('CAD Bearing Plates', 'C0392B')
    for c, w in zip('ABCDEFGHIJK', (13, 9, 9, 8, 8, 10, 10, 13, 8, 8, 8)):
        ws.column_dimensions[c].width = w
    row = cad_header(book, ws, 'BEARING PLATE DATA TABLE - TYPE 2', 'Slope (%) is the notebook seat-plane slope, + up-station.', '07-01-13', 13)
    book.put(ws, f'A{row}', 'GENERAL BEARING PLATE DATA', font=F_HEAD, fill=FILL_HEAD, border=BOX, align=CENTER)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    book.put(ws, f'H{row}', 'BEVELED PLATE REQUIRED (Yes/No)', font=F_HEAD, fill=FILL_HEAD, border=BOX, align=CENTER)
    ws.merge_cells(f'H{row}:H{row+1}')
    book.put(ws, f'I{row}', 'BEVELED PLATE DIMENSIONS (PLATE B) (inches)', font=F_HEAD, fill=FILL_HEAD, border=BOX, align=CENTER)
    ws.merge_cells(start_row=row, start_column=9, end_row=row, end_column=11)
    ws.row_dimensions[row].height = 30
    row += 1
    for c, t in zip('ABCDEFGIJK', ('BRG. PLATE MARK', 'SPAN NO(s).', 'BEAM NO(s).', 'PAD TYPE', 'BEAM END', 'PLAN VIEW CASE', 'SLOPE (%)', 'C', 'X', 'Y')):
        book.put(ws, f'{c}{row}', t, font=F_HEAD, fill=FILL_HEAD, border=BOX, align=CENTER)
    ws.row_dimensions[row].height = 30
    row += 1
    ctx['cad_plates'] = {}
    for key, r in ctx['plate_rows'].items():
        for c, src in zip('ABCDEFGHIJK', 'BCDEFGHIJKL'):
            cad_cell(book, ws, f'{c}{row}', f"={book.xref(sws, f'{src}{r}')}", '0.000' if c == 'G' else None)
        ctx['cad_plates'][key] = row
        row += 1
    return ws


def build_cad_pedestals(book: Book, ctx):
    sws = ctx['ws_seats']; sub = ctx['ws_sub']; spans = ctx['spans']; n = ctx['model'].count
    ws = book.sheet('CAD Pedestal Elevs', 'C0392B')
    for c, w in zip('ABCDEFGHI', (14, 10, 11, 11, 11, 11, 11, 11, 11)):
        ws.column_dimensions[c].width = w
    row = cad_header(book, ws, 'TOP OF PEDESTAL ELEVATIONS', 'Feet. Corner letters per the Method tab plan view; tops are level transversely.', None, 11)
    book.put(ws, f'A{row}', 'END BENTS', font=F_H2); row += 1
    row = header_row(book, ws, row, 1, ['END BENT NO.', 'BEAM NO.', 'EL. A (FT)', 'EL. B (FT)', 'EL. C (FT)', 'EL. D (FT)'], 30)
    ctx['cad_ped'] = {}
    for tag, label in (('EB1', '1'), ('EB2', '2')):
        for k in range(1, n + 1):
            r = ctx['eb_ped_rows'][(tag, k)]
            book.put(ws, f'A{row}', label, border=BOX, align=CENTER)
            book.put(ws, f'B{row}', f"={book.xref(sws, f'C{r}')}", border=BOX, align=CENTER, font=F_LINK)
            for c, src in zip('CDEF', 'IJKL'):
                cad_cell(book, ws, f'{c}{row}', f"={book.xref(sws, f'{src}{r}')}", '0.000')
            ctx['cad_ped'][(tag, k)] = row
            row += 1
    row += 1
    if ctx['pier_ped_rows']:
        book.put(ws, f'A{row}', 'PIER (one pedestal per beam line; E and F are the pad-end break lines)', font=F_H2); row += 1
        row = header_row(book, ws, row, 1, ['PIER NO.', 'BEAM LINE', 'EL. A (FT)', 'EL. B (FT)', 'EL. C (FT)', 'EL. D (FT)', 'EL. E (FT)', 'EL. F (FT)', 'EL. @ CL PIER'], 30)
        for k in range(1, n + 1):
            r = ctx['pier_ped_rows'][k]
            book.put(ws, f'A{row}', '1', border=BOX, align=CENTER)
            book.put(ws, f'B{row}', k, border=BOX, align=CENTER)
            for c, src in zip('CDEFGHI', 'GJGJHIK'):
                cad_cell(book, ws, f'{c}{row}', f"={book.xref(sws, f'{src}{r}')}", '0.000')
            ctx['cad_ped'][('Pier', k)] = row
            row += 1
        row += 1
    book.put(ws, f'A{row}', 'CAP AND PILE CUTOFF ELEVATIONS', font=F_H2); row += 1
    row = header_row(book, ws, row, 1, ['SUPPORT', 'PILE NO.', 'OFFSET (FT)', 'TOP OF CAP', 'BOTTOM OF CAP', 'PILE CUTOFF'], 30)
    for (tag, j), r in ctx['pile_rows'].items():
        book.put(ws, f'A{row}', {'Pier': 'Pier', 'EB1': 'End Bent 1', 'EB2': 'End Bent 2'}[tag], border=BOX, align=CENTER)
        book.put(ws, f'B{row}', j, border=BOX, align=CENTER)
        for c, src, fmt in (('C', 'D', '0.000'), ('D', 'G', '0.000'), ('E', 'H', '0.000'), ('F', 'J', '0.000')):
            cad_cell(book, ws, f'{c}{row}', f"={book.xref(sub, f'{src}{r}')}", fmt)
        row += 1
    row += 1
    book.put(ws, f'A{row}', 'END BENT WALL ELEVATIONS (D-026/D-033)', font=F_H2); row += 1
    heads = ['END BENT NO.', 'TOP OF BACKWALL @ LT CAP END', 'TOP OF BACKWALL @ PGL', 'TOP OF BACKWALL @ RT CAP END',
             'TOP OF CHEEK LT', 'TOP OF CHEEK RT', 'BOTTOM OF CHEEK (CAP UNDERSIDE)', 'BOTTOM OF LUGS', 'REAR LUG FACE / COPING STOP STA.']
    row = header_row(book, ws, row, 1, heads, 54)
    for tag, label in (('EB1', '1'), ('EB2', '2')):
        book.put(ws, f'A{row}', label, border=BOX, align=CENTER)
        for c, name, fmt in (('B', 'BwTopL', '0.000'), ('C', 'BwTop0', '0.000'), ('D', 'BwTopR', '0.000'), ('E', 'CkTopLv', '0.000'),
                             ('F', 'CkTopRv', '0.000'), ('G', 'CkBotL', '0.000'), ('H', 'LugBotL', '0.000'), ('I', 'LugFace', '0.00')):
            cad_cell(book, ws, f'{c}{row}', f'={tag}_{name}', fmt)
        ctx.setdefault('cad_walls', {})[tag] = row
        row += 1
    book.put(ws, f'A{row}', 'Cheek webs 8 in outside each cap end (251 in out-to-out); 6-in bottom and rear lugs, 5.5 in wide, on the outer face. '
             'Cheek tops are the D-033 barrier-return bases.', font=F_NOTE)
    return ws


# ==================================================================== QC ==
def qc_items(ctx):
    """(label, workbook formula, notebook value, tolerance)."""
    m = ctx['model']; items = []
    items.append(('Common support drop D (ft)', '=Sup_Drop', float(m.support_drop), 1e-9))
    items.append(('Grade 1', '=Prof_G1', float(m.g1), 1e-12))
    items.append(('Grade 2', '=Prof_G2', float(m.g2), 1e-12))
    if m.profile_mode == 'brokeback':
        items.append(('PI elevation (ft)', '=Prof_PIZ', float(m.peak_elev), 1e-9))
    for i in range(1, len(ctx['spans']) + 1):
        a, b = m.bearing_limits[i - 1]; lo, hi = m.beam_limits[i - 1]
        items += [(f'Span {i} bearing A station', f'=Sp{i}_BrgA', a, 1e-9), (f'Span {i} bearing B station', f'=Sp{i}_BrgB', b, 1e-9),
                  (f'Span {i} beam start', f'=Sp{i}_Bm0', lo, 1e-9), (f'Span {i} beam end', f'=Sp{i}_Bm1', hi, 1e-9),
                  (f'Span {i} PGL at bearing A', f'=Sp{i}_ZA', float(m.pgl(a)), 1e-9),
                  (f'Span {i} PGL at bearing B', f'=Sp{i}_ZB', float(m.pgl(b)), 1e-9)]
    items.append(('Minimum flange haunch, bridge (in)', '=Haunch_Min', float(m.haunch_min * 12), 1e-6))
    cws = "'Camber & Haunch'"
    for (i, k, loc), r in ctx['haunch_rows'].items():
        span = i - 1; a, b = m.bearing_limits[span]
        s = {'Begin brg': a, 'Midspan': (a + b) / 2, 'End brg': b}[loc]
        e = m.evaluate([s], span)
        items.append((f'Haunch @ CL beam {i}-{k} {loc} (in)', f"={cws}!L{r}", float(e['haunch'][0, k - 1, 1] * 12), 1e-6))
        items.append((f'Min flange haunch {i}-{k} {loc} (in)', f"={cws}!M{r}", float(e['haunch'][0, k - 1].min() * 12), 1e-6))
    for p in m.supports:
        tag = {'Pier': 'Pier', 'End Bent 1': 'EB1', 'End Bent 2': 'EB2'}[p['name']]
        items.append((f'{p["name"]} cap top (ft)', f'={tag}_Top', float(p['top']), 1e-9))
        items.append((f'{p["name"]} cap CL station (ft)', f'={tag}_CapSta', float(p['station']), 1e-9))
        items.append((f'{p["name"]} cap length (ft)', f'={tag}_Len', float(p['length']), 1e-9))
        for seat in p['seats'].to_dict('records'):
            span = int([s['id'] for s in ctx['spans']].index(seat['Span'])) + 1
            r = ctx['seat_rows'][(tag, span, int(seat['Beam']))]
            items.append((f'{p["name"]} seat {span}-{int(seat["Beam"])} pedestal top (ft)', f"='Bearing Seats'!V{r}", float(seat['Pedestal top (ft)']), 1e-9))
            items.append((f'{p["name"]} seat {span}-{int(seat["Beam"])} bearing slope', f"='Bearing Seats'!T{r}", float(seat['Bearing slope']), 1e-9))
            items.append((f'{p["name"]} seat {span}-{int(seat["Beam"])} pad-top plane (ft)', f"='Bearing Seats'!S{r}", float(seat['Pad top (ft)']), 1e-9))
            hr = ctx['ped_height_rows'][(tag, span, int(seat['Beam']))]
            items.append((f'{p["name"]} seat {span}-{int(seat["Beam"])} pedestal height (in)', f"='Substructure'!H{hr}", float(seat['Pedestal height (in)']), 1e-6))
        for row in p['pile_elevations'].to_dict('records'):
            r = ctx['pile_rows'][(tag, int(row['Pile']))]
            items.append((f'{p["name"]} pile {int(row["Pile"])} offset (ft)', f"='Substructure'!D{r}", float(row['Offset (ft)']), 1e-9))
            items.append((f'{p["name"]} pile {int(row["Pile"])} cutoff (ft)', f"='Substructure'!J{r}", float(row['Pile cutoff (ft)']), 1e-9))
        for ped in p['solid_geometry']['pedestals']:
            k = ped['beam']
            if tag == 'Pier':
                r = ctx['pier_ped_rows'][k]
                for c, z, lab in zip('GHIJ', ped['tops'], ('A/C', 'E', 'F', 'B/D')):
                    items.append((f'Pier pedestal line {k} EL {lab} (ft)', f"='Bearing Seats'!{c}{r}", float(z), 1e-9))
            else:
                r = ctx['eb_ped_rows'][(tag, k)]
                back, front = (ped['tops'][0], ped['tops'][1]) if tag == 'EB1' else (ped['tops'][1], ped['tops'][0])
                items.append((f'{p["name"]} pedestal {k} EL A (ft)', f"='Bearing Seats'!I{r}", float(back), 1e-9))
                items.append((f'{p["name"]} pedestal {k} EL B (ft)', f"='Bearing Seats'!J{r}", float(front), 1e-9))
    for eb, tag in zip(m.end_bents, ('EB1', 'EB2')):
        w = eb.get('wall_summary')
        if not w:
            continue
        name = eb['name']; prof = eb['backwall_profile'].set_index('Offset (ft)')['Backwall top (ft)']
        lugs = eb['walls'].set_index(['Component', 'Side'])['Bottom (ft)']
        items += [(f'{name} out-to-out with cheek webs (in)', f'={tag}_OutIn', float(w['out_to_out_in']), 1e-9),
                  (f'{name} rear-lug face / coping stop (ft)', f'={tag}_LugFace', float(w['rear_lug_face_station']), 1e-9),
                  (f'{name} backwall top, left cap end (ft)', f'={tag}_BwTopL', float(prof.iloc[0]), 1e-9),
                  (f'{name} backwall top at PGL (ft)', f'={tag}_BwTop0', float(prof.loc[0.0]), 1e-9),
                  (f'{name} backwall top, right cap end (ft)', f'={tag}_BwTopR', float(prof.iloc[-1]), 1e-9),
                  (f'{name} backwall height at PGL (ft)', f'={tag}_BwH0', float(w['backwall_crown_height_ft']), 1e-9),
                  (f'{name} backwall concrete (ft3)', f'={tag}_BwVol', float(w['backwall_volume_ft3']), 1e-9)]
        for side in ('L', 'R'):
            items += [(f'{name} cheek top {side} (ft)', f'={tag}_CkTop{side}v', float(w[f'cheek_top_{side}_ft']), 1e-12),
                      (f'{name} cheek {side} concrete (ft3)', f'={tag}_CkVol{side}', float(w[f'cheek_{side}_volume_ft3']), 1e-9),
                      (f'{name} cheek {side} lug bottom (ft)', f'={tag}_LugBot{side}', float(lugs.loc[('Bottom lug', side)]), 1e-9),
                      (f'{name} cheek {side} bottom (ft)', f'={tag}_CkBot{side}', float(lugs.loc[('Cheek web', side)]), 1e-9)]
    if np.isfinite(m.clearance_min):
        items.append(('Minimum vertical clearance (ft)', '=Clr_Min', float(m.clearance_min), 1e-7))
        items.append(('Station of minimum clearance (ft)', '=Clr_Sta', float(m.clearance_station), 1e-6))
        items.append(('Roadway coverage (ft)', '=Clr_Cover', float(m.covered_length), 1e-6))
    for i in range(1, len(ctx['spans']) + 1):
        items.append((f'Span {i} beam length vs Mathcad L.beam (ft)', f'=Sp{i}_LBeam', float(ctx['mathcad'][i]['L_beam_ft']), 1e-9))
        items.append((f'Span {i} setback vs Mathcad BearingDistance (in)', '=Phys_Setback', float(ctx['mathcad'][i]['bearing_distance_in']), 1e-9))
        items.append((f'Span {i} deck thickness vs Mathcad t.deck (in)', '=Deck_T*12', float(ctx['mathcad'][i]['t_deck_in']), 1e-9))
        items.append((f'Span {i} beam spacing vs Mathcad (ft, Mathcad rounds)', '=Bm_S', float(ctx['mathcad'][i]['beam_spacing_ft']), 1e-3))
        items.append((f'Span {i} camber vs Mathcad NetCamber (in)', f'=Sp{i}_C', float(ctx['mathcad'][i]['net_camber_in']), 1e-12))
    return items


def input_snapshot(book: Book, ws, top=8):
    """Hidden QC columns O:R: every generated input, the value written at build, and a changed flag.

    The Method-tab banners and the QC validity row read the count (Inputs_Edited), so an edit
    made in Excel is reported everywhere even though the tables and charts follow it.
    """
    for k, head in enumerate(('Input cell', 'Value at build', 'Current value', 'Changed')):
        book.put(ws, f'{col(15 + k)}{top}', head, font=F_HEAD)
    r = top + 1
    for sheet, ref, value in book.inputs:
        cell = f"'{sheet}'!{ref}"
        book.put(ws, f'O{r}', f'{sheet}!{ref}', font=F_NOTE)
        if value is not None:
            ws[f'P{r}'] = value
        ws[f'Q{r}'] = f'=IF(ISBLANK({cell}),"",{cell})'
        if value is None:
            flag = f'=IF(ISBLANK({cell}),0,1)'
        elif isinstance(value, (int, float)) and not isinstance(value, bool):
            flag = f'=IF(ISNUMBER({cell}),IF(ABS({cell}-P{r})<=1E-12*MAX(1,ABS(P{r})),0,1),1)'
        else:
            flag = f'=IF(ISBLANK({cell}),1,IF(EXACT({cell}&"",P{r}&""),0,1))'
        ws[f'R{r}'] = flag
        r += 1
    book.put(ws, f'O{r}', 'Inputs edited since the build', font=F_BOLD)
    book.put(ws, f'R{r}', f'=SUM(R{top + 1}:R{r - 1})', font=F_BOLD)
    book.define('Inputs_Edited', ws, f'R{r}')
    for c in 'OPQR':
        ws.column_dimensions[c].hidden = True


def build_qc(book: Book, ctx):
    m = ctx['model']
    ws = book.sheet('QC', '2F855A')
    for c, w in zip('ABCDEFG', (3, 58, 17, 17, 14, 11, 10)):
        ws.column_dimensions[c].width = w
    row = title_block(book, ws, 'QC - workbook vs. notebook', 'Workbook formulas against the geometry notebook values captured at build time', 9)
    row = section(book, ws, row, 'A. Numerical agreement with the notebook model and Mathcad', 7)
    row = header_row(book, ws, row, 2, ['Check', 'Workbook value', 'Notebook / source value', 'Difference', 'Tolerance', 'Status'], 30)
    first = row
    for label, formula, value, tol in qc_items(ctx):
        book.put(ws, f'B{row}', label)
        book.put(ws, f'C{row}', formula, font=F_LINK, fmt='0.000000000')
        book.put(ws, f'D{row}', value, font=F_INPUT, fill=FILL_INPUT, fmt='0.000000000')
        book.put(ws, f'E{row}', f'=C{row}-D{row}', fmt='0.0E+00')
        book.put(ws, f'F{row}', tol, fmt='0.0E+00')
        book.put(ws, f'G{row}', f'=IF(ABS(E{row})<=F{row},"PASS","FAIL")', align=CENTER, border=BOX)
        row += 1
    numeric_last = row - 1
    row += 1
    row = section(book, ws, row, 'B. FGE grid agreement (every FGE cell)', 7)
    fge = book.wb['FGE']
    snap_rows = []
    book.put(ws, f'B{row}', 'Maximum |FGE workbook - notebook deck_top| (ft)')
    # Hidden snapshot block to the right: notebook deck_top at each FGE cell.
    snap_col = 12
    diffs = []
    sr = row
    for i, f in ctx['fge'].items():
        for name, lr in f['lines'].items():
            offset = next(o for lab, _, o in fge_lines(m) if lab.format(span=ctx['spans'][i - 1]['id']) == name)
            for c in f['cols']:
                station = ctx['fge_stations'][i][f['cols'].index(c)]
                value = float(m.deck_top(station, offset))
                ref = f'{col(snap_col)}{sr}'
                book.put(ws, ref, value, font=F_INPUT, fmt='0.000000000')
                book.put(ws, f'{col(snap_col+1)}{sr}', f"=ABS('FGE'!{c}{lr}-{ref})", fmt='0.0E+00')
                diffs.append(f'{col(snap_col+1)}{sr}')
                sr += 1
    book.put(ws, f'L{row-1}', 'Notebook deck_top snapshot', font=F_HEAD)
    book.put(ws, f'C{row}', f'=MAX({col(snap_col+1)}{row}:{col(snap_col+1)}{sr-1})', fmt='0.0E+00')
    book.put(ws, f'F{row}', 1e-9, fmt='0.0E+00')
    book.put(ws, f'G{row}', f'=IF(C{row}<=F{row},"PASS","FAIL")', align=CENTER, border=BOX)
    fge_check = row
    row += 2
    row = section(book, ws, row, 'C. Validity checks', 7)
    validity = [(label, '=' + ref) for label, ref in ctx['layout_checks']]
    validity.append(('Brokeback PI not inside any beam (closed-form drop requires a straight PGL per beam)',
                     '=IF(Prof_Mode="brokeback",IF(AND(' + ','.join(f'OR(Prof_PISta<=Sp{i}_Bm0,Prof_PISta>=Sp{i}_Bm1)' for i in range(1, len(ctx['spans']) + 1)) + '),"PASS","FAIL"),"PASS")'))
    validity.append(('Spans contiguous and covering the profile',
                     '=IF(AND(ABS(Sp1_Sta0-Prof_Sta0)<1E-8,ABS(Sp{0}_Sta1-Prof_Sta1)<1E-8'.format(len(ctx['spans'])) +
                     ''.join(f',ABS(Sp{i}_Sta1-Sp{i+1}_Sta0)<1E-8' for i in range(1, len(ctx['spans']))) + '),"PASS","FAIL")'))
    validity.append(('End-bent records agree with physical beam-end values',
                     '=IF(AND(EB1_BmClr=Phys_EndClr,EB2_BmClr=Phys_EndClr,EB1_BmBrg=Phys_Setback,EB2_BmBrg=Phys_Setback),"PASS","FAIL")'))
    validity.append(('Pier station at the span joint', '=IF(ABS(Pier_Sta-Sp1_Sta1)<1E-8,"PASS","FAIL")'))
    for i in range(1, len(ctx['spans']) + 1):
        validity.append((f'Span {i} physical layout order', f'=Sp{i}_Chk'))
    validity += [(lab, f'={ref}') for lab, ref in [(f'Deflection sum = NetCamber = CONFIG camber (span {k+1})', r) for k, r in enumerate(ctx['deflection_checks'])]]
    validity.append(('All seat, pedestal, plate and cap checks on the calculation tabs',
                     "=IF(COUNTIF('Bearing Seats'!B1:Z400,\"FAIL\")+COUNTIF('Substructure'!B1:Z400,\"FAIL\")+COUNTIF('Beam Geometry'!B1:Z400,\"FAIL\")"
                     "+COUNTIF('Beam Data'!B1:Z400,\"FAIL\")+COUNTIF('Camber & Haunch'!B1:Z400,\"FAIL\")+COUNTIF('Profile & Deck'!B1:Z400,\"FAIL\")=0,\"PASS\",\"FAIL\")"))
    validity.append(('Clearance status (informational)', '=IF(Clr_Status="Below requirement","FAIL",IF(Clr_Status="Meets requirement","PASS","CHECK"))'))
    input_snapshot(book, ws)
    validity.append(('Inputs unchanged since the notebook build (every generated input vs. its build value, hidden columns O:R)',
                     '=IF(Inputs_Edited=0,"PASS","CHECK")'))
    vfirst = row
    for label, formula in validity:
        book.put(ws, f'B{row}', label, align=WRAP)
        book.put(ws, f'G{row}', formula, align=CENTER, border=BOX)
        row += 1
    check_formatting(ws, f'G{first}:G{row}')
    row += 1
    book.put(ws, f'B{row}', 'Overall', bold=True)
    book.put(ws, f'G{row}', f'=IF(COUNTIF(G{first}:G{row-2},"FAIL")=0,"PASS","FAIL")', align=CENTER, border=BOX, bold=True)
    book.define('QC_Overall', ws, f'G{row}')
    book.put(ws, f'C{row}', f'=COUNTIF(G{first}:G{row-2},"PASS")&" pass / "&COUNTIF(G{first}:G{row-2},"FAIL")&" fail / "&COUNTIF(G{first}:G{row-2},"CHECK")&" check"')
    book.define('QC_Counts', ws, f'C{row}')
    check_formatting(ws, f'G{row}')
    ws.column_dimensions['L'].hidden = True
    ws.column_dimensions['M'].hidden = True
    ws.freeze_panes = 'A9'
    return ws


# ================================================================= Cover ==
def build_cover(book: Book, ctx, sheet_index):
    ws = book.wb.create_sheet('Cover', 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = '2B6CB0'
    for c, w in zip('ABCD', (3, 34, 90, 20)):
        ws.column_dimensions[c].width = w
    book.put(ws, 'B2', 'B4 Pedestrian Bridge - Geometry and CAD Tables', font=Font(name=FONT, size=18, bold=True, color='1F2A36'))
    book.put(ws, 'B3', '=TB_Project&" | "&TB_Structure', font=F_LINK)
    rows = [
        ('Geometry authority', 'Colab geometry notebook Bridge_Geometry_2 Beam_V2.ipynb (CONFIG + BridgeModel). This workbook reproduces its '
                               'equations in Excel; the CAD model and these CAD tables are built from it, never the reverse.'),
        ('Beam design data', 'Saved results of the two registered Mathcad runs (Span 1, Span 2). Mathcad was not recalculated.'),
        ('Template methods', 'Where the notebook is silent (Table of Beam Variables bar dimensions, Dim L/R, build-up case, beveled plates) '
                             "the BCC example workbook's formulas are used and flagged on the Method tabs."),
        ('QC status', '=QC_Overall&"  ("&QC_Counts&")"'),
        ('Input sync', '=IF(Inputs_Edited=0,"Inputs match the notebook build: tables and Method-tab charts show the notebook geometry.",'
                       'Inputs_Edited&" input(s) edited since the build. Tables and charts follow the edits, which no longer match the '
                       'notebook. Update the notebook CONFIG and rebuild.")'),
        ('Rebuild', r'C:\github\D2MPDB\.venv\Scripts\python.exe scripts\build_geometry_cad_tables.py   (rebuilds, recalculates in Excel, verifies)'),
        ('Build fingerprint', ctx['fingerprint']),
        ('Generator version', GENERATOR_VERSION),
        ('Built', ctx['built']),
        ('Notebook model ID', ctx['model'].model_id),
    ]
    r = 5
    for label, value in rows:
        book.put(ws, f'B{r}', label, font=F_BOLD, align=WRAP)
        c = book.put(ws, f'C{r}', value, font=F_LINK if str(value).startswith('=') else F_BASE, align=WRAP)
        if label == 'QC status':
            ws.conditional_formatting.add(f'C{r}', FormulaRule(formula=[f'LEFT(C{r},4)="PASS"'], fill=CF_PASS))
            ws.conditional_formatting.add(f'C{r}', FormulaRule(formula=[f'LEFT(C{r},4)="FAIL"'], fill=CF_FAIL))
        if label == 'Input sync':
            ws.conditional_formatting.add(f'C{r}', FormulaRule(formula=['Inputs_Edited=0'], fill=CF_PASS))
            ws.conditional_formatting.add(f'C{r}', FormulaRule(formula=['Inputs_Edited>0'], fill=CF_WARN))
        for name, key in (('Build_Fingerprint', 'Build fingerprint'), ('Build_Model', 'Notebook model ID'), ('Build_Time', 'Built')):
            if label == key:
                book.define(name, ws, f'C{r}')
        ws.row_dimensions[r].height = 30 if len(str(value)) > 90 else 16
        r += 1
    r += 1
    book.put(ws, f'B{r}', 'Sources', font=F_H2); r += 1
    for label, path, digest in ctx['source_rows']:
        book.put(ws, f'B{r}', label, font=F_BOLD)
        book.put(ws, f'C{r}', f'{path}   SHA-256 {digest}', font=F_NOTE)
        r += 1
    r += 1
    book.put(ws, f'B{r}', 'Tabs', font=F_H2); r += 1
    for name, desc in sheet_index:
        c = book.put(ws, f'B{r}', name, font=Font(name=FONT, size=10, color='2B6CB0', underline='single'))
        c.hyperlink = f"#'{name}'!A1"
        book.put(ws, f'C{r}', desc, align=WRAP)
        r += 1
    r += 1
    book.put(ws, f'B{r}', 'Cell legend', font=F_H2); r += 1
    for text, font, fill in (('Input (generated from notebook / Mathcad)', F_INPUT, FILL_INPUT), ('Calculated in this tab', F_BASE, None),
                             ('Linked from another tab', F_LINK, None), ('PASS', F_BASE, FILL_PASS), ('FAIL', F_BASE, FILL_FAIL), ('CHECK / review', F_BASE, FILL_WARN)):
        book.put(ws, f'B{r}', text, font=font, fill=fill); r += 1
    r += 1
    book.put(ws, f'B{r}', 'Open items', font=F_H2); r += 1
    for text in ctx['open_items']:
        book.put(ws, f'C{r}', '- ' + text, align=WRAP)
        ws.row_dimensions[r].height = 30 if len(text) > 95 else 16
        r += 1
    return ws


# ============================================================ entry point ==
def source_hashes():
    hashes = {
        'notebook_lf_text': sources.notebook_text_sha256(),
        'roadway_xml': sources.sha256(sources.ROADWAY_XML),
        'mathcad_span1': sources.sha256(sources.MATHCAD[1]),
        'mathcad_span2': sources.sha256(sources.MATHCAD[2]),
        'generator': hashlib.sha256(b''.join((PACKAGE / f).read_bytes().replace(b'\r\n', b'\n')
                                             for f in ('sources.py', 'figures.py', 'charts.py', 'workbook.py', 'methods.py'))).hexdigest(),
    }
    return hashes


def fingerprint(hashes) -> str:
    return hashlib.sha256(json.dumps(hashes, sort_keys=True).encode()).hexdigest()[:20]


def build(output: Path = OUTPUT) -> dict:
    from . import methods
    cfg, model, env = sources.load_notebook_model()
    spans = sorted(cfg['spans'], key=lambda s: s['start_ft'])
    if len(spans) != 2 or [s['id'] for s in spans] != [1, 2]:
        raise ValueError('The workbook maps notebook spans 1-2 to the two Mathcad runs; span layout changed.')
    if cfg['deck']['mode'] not in ('crowned', 'uniform'):
        raise NotImplementedError('Deck transition mode is not implemented in the workbook.')
    if cfg['profile']['mode'] == 'brokeback' and any(lo < model.peak < hi for lo, hi in model.beam_limits):
        raise NotImplementedError('A brokeback PI inside a beam needs piecewise haunch minimization.')
    hashes = source_hashes()
    ctx = dict(cfg=cfg, model=model, env=env, spans=spans, mathcad={i: sources.mathcad_beam_data(i) for i in (1, 2)},
               road=sources.roadway_segments(model), hashes=hashes, fingerprint=fingerprint(hashes),
               built=_dt.datetime.now().strftime('%Y-%m-%d %H:%M'))
    ctx['fge_stations'] = {}
    for i, span in enumerate(spans, start=1):
        a, b = model.bearing_limits[i - 1]
        ctx['fge_stations'][i] = [span['start_ft'], a] + [a + k * (b - a) / FGE_INCREMENTS for k in range(1, FGE_INCREMENTS)] + [b, span['end_ft']]
    ctx['source_rows'] = [
        ('Geometry notebook', 'Bridge_Geometry_2 Beam_V2.ipynb (lf_text)', hashes['notebook_lf_text']),
        ('Roadway XML', 'Geometry Report.xml', hashes['roadway_xml']),
        ('Mathcad Span 1', ctx['mathcad'][1]['source_file'], hashes['mathcad_span1']),
        ('Mathcad Span 2', ctx['mathcad'][2]['source_file'], hashes['mathcad_span2']),
    ]
    ctx['open_items'] = [
        'End caps, backwalls, cheek webs and lugs (D-026/D-027/D-033) are notebook geometry (D-036). The CAD EB1 backwall seat follows the '
        'MSE landing and holds 0.117 ft3 more concrete than the notebook rule; rebuild it from the notebook.',
        'Cheek tops are the four D-033 MSE barrier-return base elevations, entered as notebook inputs.',
        'Deck barrier (D-031, 12 ft clear) is not in the notebook, so no barrier-face FGE line is tabulated.',
        'Job number, calculated/checked by and dates are blank on the Inputs tab.',
        'Strand pattern type, concrete class and pad type are designer assignments (Inputs tab).',
        'Dim L includes R per the BCC template - confirm against the FDOT Table of Beam Variables instructions.',
        'Mathcad .dat inputs are missing from the project record; values are the worksheets\' saved results.',
        'The roadway XML ends about 5 ft before End Bent 2, so clearance is reported for the covered region only.',
    ]
    book = Book()
    tabs = [
        ('Inputs', build_inputs, 'All inputs with provenance (notebook CONFIG, Mathcad saved results, designer/template constants).'),
        ('Profile & Deck', build_profile, 'PGL grades, deck edges, beam offsets and haunch check points.'),
        ('Beam Geometry', build_beam_geometry, 'Physical beam ends, bearings, chords; Dims J, K, P, L and shortening R.'),
        ('Camber & Haunch', build_camber, 'Common support drop D, build-up Dims B/C/D, minimum haunch, Mathcad deflections.'),
        ('FGE', build_fge, 'Finished grade elevations at every line and station.'),
        ('Bearing Seats', build_seats, 'Seat planes, slopes, pedestal tops and corners, bearing plate data.'),
        ('Substructure', build_substructure, 'Cap tops, pedestal heights, cap corners and pile cutoffs.'),
        ('Beam Data', build_beam_data, 'Bars 5K zones from Mathcad and FIB standard bar dimensions.'),
        ('Clearance', build_clearance, 'Exact minimum vertical clearance under the beams.'),
        ('CAD FGE', build_cad_fge, 'CAD table: Finish Grade Elevations by span.'),
        ('CAD Build-Up', build_cad_buildup, 'CAD table: Build-Up & Deflection Data.'),
        ('CAD Beam Variables', build_cad_beam_variables, 'CAD table: Florida-I Beam Table of Beam Variables.'),
        ('CAD Bearing Plates', build_cad_plates, 'CAD table: Bearing Plate Data, Type 2.'),
        ('CAD Pedestal Elevs', build_cad_pedestals, 'CAD tables: top of pedestal, cap and pile cutoff elevations.'),
        ('QC', build_qc, 'Every key result compared with the notebook; validity checks.'),
    ]
    index = []
    for name, builder, desc in tabs:
        builder(book, ctx)
    # Method tabs after every calculation tab, so their live charts can reference any table row.
    for name, builder, desc in tabs:
        method = methods.build_method(book, ctx, name)
        index.append((name, desc))
        index.append((method.title, f'Method for {name}: equations, explanation, schematics and live charts.'))
    build_cover(book, ctx, index)
    # Interleave each Method tab immediately after its calculation tab.
    order = ['Cover']
    for name, _, _ in tabs:
        order += [name, f'{name} Method']
    book.wb._sheets = [book.wb[n] for n in order]
    book.wb.calculation.fullCalcOnLoad = True
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    book.wb.save(output)
    return dict(output=output, fingerprint=ctx['fingerprint'], hashes=hashes, model_id=model.model_id,
                qc=qc_items(ctx), sheets=order)
