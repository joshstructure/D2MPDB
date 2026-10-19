"""Case, checks, report, Pile Data Table, Mathcad, Blockpad and bundle exports for the pile calculation."""
import csv
from datetime import datetime, timezone
import importlib
import io
import json
from pathlib import Path
import zipfile
from .model import DATA_TABLE_COLUMNS, INPUTS, CONFIRMATIONS, upgrade_case

CASE = 'pile_case.json'
CHECKS = 'pile_checks.csv'
REPORT = 'pile_calculation_report.html'
TABLE = 'pile_data_table.xlsx'
MATHCAD = 'pile_design.xmcd'
BPAD = 'pile_design.bpad'                    # internal key; the file on disk carries the support name ("End Bent 1 Pile Design.bpad")
BUNDLE = 'pile_bundle.zip'
MATHCAD_NOTE = 'pile_design_xmcd_unavailable.txt'
BPAD_NOTE = 'pile_design_bpad_unavailable.txt'
VERIFICATION = 'verification.json'


class MathcadUnavailable(RuntimeError):
    """Raised when pile_design.mathcad (the distilled Mathcad generator) cannot be imported."""


class BpadUnavailable(RuntimeError):
    """Raised when pile_design.bpad (the distilled Blockpad generator) or what it needs cannot be imported."""


def case_json(case):
    return json.dumps(case, indent=1, sort_keys=False)


def load_case(content):
    data = json.loads(bytes(content).decode('utf-8-sig'))
    return upgrade_case(data)


# ----------------------------------------------------------------------------- checks
def checks_rows(e):
    for c in e.checks:
        yield dict(group=c.group, check=c.label, status=c.status, ratio='' if c.ratio is None else f'{c.ratio:.4f}',
                   demand=c.demand, capacity=c.capacity, basis=c.basis, note=c.note,
                   review='' if not c.review else f'{c.review["action"]}: {c.review["reason"]} ({c.review["utc"]})')


def checks_csv(e):
    buffer = io.StringIO()
    rows = list(checks_rows(e))
    writer = csv.DictWriter(buffer, fieldnames=list(rows[0]))
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue()


# ----------------------------------------------------------------------------- Pile Data Table
def _basis_rows(e):
    p, conf, v = e.case['inputs'], e.case['confirmations'], e.value
    analysis = e.case.get('analysis') or {}
    fd = e.details['fdl']
    state = lambda key: 'Confirmed' if conf.get(key) else 'PENDING'
    rows = [
        ('Support', e.case['support']),
        ('Pile family', 'Square prestressed concrete' if e.family == 'psc' else 'Steel pipe (noncomposite)'),
        ('FBMP export', analysis.get('filename', 'No analysis loaded')),
        ('FBMP SHA-256', analysis.get('sha256', 'N/A')),
        ('Source register', (f'Matches {analysis["registered_id"]}' if analysis.get('registered_id') else 'Not a registered export (new upload)') if analysis else 'N/A'),
        ('Factored Design Load, column I (tons)', f'{v("FDL_tons"):.2f} (maximum strength pile-head compression from FBMP, downdrag excluded)'),
        ('Downdrag, column J (tons)', f'{p["DD_tons"]:g} per pile, drag load DR. Source: {e.case["sources"]["DD_tons"]["source"]} ({state("downdrag_source")} documentation)'),
        ('Drag load factors', f'γDR = {p["gamma_DR_str"]:g} at the strength limit states and {p["gamma_DR_ser"]:g} at Service I (LRFD Table 3.4.1-1); '
                              'the geotechnical limit state uses DR = 0. The factored drag load enters the structural demands only; it is not part of column I.'),
        ('Net scour resistance, column L (tons)', f'{p["R_scour_net_tons"]:g}'),
        ('Total scour resistance, column K (tons)', f'{p["R_scour_total_tons"]:g}'),
        ('Resistance factor φ, column P', f'{p["phi_geo"]:g} (SDG Table 3.5.7-1; footnote 7: Pile Data Tables use the default 0.75)'),
        ('Nominal Bearing Resistance equation, column C', 'Rn = (Factored Design Load + Net Scour Resistance + Downdrag) / φ  (SDG Eq. 3-1, Section 3.5.13)'),
        ('Nominal Bearing Resistance (tons)', f'{v("NBR_tons"):.2f} computed; {e.data_table.get("Nominal Bearing Resistance (tons)", "N/A")} in the table (rounded up)'),
    ]
    if e.family == 'psc':
        rows.append(('Table 3.5.13-1 maximum driving resistance (tons)', f'{v("R_drive_tons"):.0f} for the {p["D_pile"]:g} in prestressed concrete pile'))
    else:
        rows.append(('Table 3.5.13-1', 'Covers prestressed concrete piles only; confirm the driving resistance limit for the steel pipe with the geotechnical engineer.'))
    rows += [
        ('Tension Resistance, column D (tons)', 'Maximum strength uplift from FBMP' + (f' ({e.details["uplift_kip"] / 2:.2f} tons)' if e.details['uplift_kip'] > 0 else ' (none in the loaded combinations)')),
        ('Minimum Tip Elevation, column E', 'Geotechnical handoff (SDG 3.5.9), rounded down to a whole foot' if p.get('EL_min_tip') is not None else 'N/A until the handoff is received'),
        ('Pile cut-off elevation, column Q', f'{p["EL_cutoff"]:.4f} ft from the FBMP pile head, rounded to 0.1 ft'),
        ('Rounding (SDG 3.5.15.D)', 'Loads up to the nearest ton; minimum tip elevations down and pile lengths up to the nearest foot; cut-off elevations to 0.1 ft.'),
        ('Jet and preform elevations, columns G and H', 'Not applicable: no jetting or preforming is specified.'),
        ('Design scour elevation, column N', 'Temporary bridges only: N/A.'),
    ]
    for key, text in CONFIRMATIONS.items():
        if key in ('steel_grade', 'corrosion') and e.family != 'pipe':
            continue
        rows.append((f'Confirmation: {key}', f'{state(key)} · {text}'))
    return rows


def data_table_xlsx(e, extra_rows=()):
    """SDM 11.4 Pile Data Table workbook: columns A-Q in order, then a Basis sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    wb = Workbook()
    ws = wb.active
    ws.title = 'Pile Data Table'
    bold = Font(bold=True)
    head = PatternFill('solid', fgColor='DCECE0')
    thin = Side(style='thin', color='9AA8B4')
    ws.append([header for _, header in DATA_TABLE_COLUMNS])
    for cell in ws[1]:
        cell.font = bold
        cell.fill = head
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for row in [e.data_table, *extra_rows]:
        ws.append([row.get(header, 'N/A') for _, header in DATA_TABLE_COLUMNS])
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    widths = (14, 24, 16, 14, 14, 20, 14, 14, 14, 12, 14, 14, 14, 18, 14, 12, 14)
    for (letter, _), width in zip(DATA_TABLE_COLUMNS, widths):
        ws.column_dimensions[letter].width = width
    ws.row_dimensions[1].height = 62
    ws.freeze_panes = 'A2'
    basis = wb.create_sheet('Basis')
    basis.append(['Item', 'Basis'])
    for cell in basis[1]:
        cell.font = bold
        cell.fill = head
    for item, text in _basis_rows(e):
        basis.append([item, text])
    basis.column_dimensions['A'].width = 48
    basis.column_dimensions['B'].width = 110
    for r in basis.iter_rows(min_row=2):
        r[1].alignment = Alignment(wrap_text=True, vertical='top')
        r[0].alignment = Alignment(vertical='top', wrap_text=True)
    inputs = wb.create_sheet('Inputs')
    inputs.append(['Input', 'Value', 'Unit', 'Meaning', 'Source', 'Status'])
    for cell in inputs[1]:
        cell.font = bold
        cell.fill = head
    for name, spec in INPUTS.items():
        src = e.case['sources'].get(name, {})
        value = e.case['inputs'].get(name)
        inputs.append([name, 'N/A' if value is None else value, spec['unit'] or '', spec['caption'], src.get('source', ''), src.get('status', '')])
    for col, width in zip('ABCDEF', (18, 12, 10, 56, 70, 22)):
        inputs.column_dimensions[col].width = width
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ----------------------------------------------------------------------------- report and Mathcad
def _report(e):
    from .report import build_report
    return build_report(e).encode('utf-8')


def mathcad_xmcd(e):
    """Distilled Mathcad 15 worksheet for the current case (pile_design.mathcad.build)."""
    try:
        mathcad = importlib.import_module(f'{__package__}.mathcad')
    except ImportError as exc:
        raise MathcadUnavailable(f'Mathcad generator not available: pile_design.mathcad could not be imported ({exc}).') from exc
    if not callable(getattr(mathcad, 'build', None)):
        raise MathcadUnavailable('Mathcad generator not available: pile_design.mathcad has no build(e) function.')
    out = mathcad.build(e)
    return out.encode('utf-8') if isinstance(out, str) else bytes(out)


def bpad_outputs(e):
    """Distilled Blockpad document of the current case with its verification record (pile_design.bpad.build_outputs)."""
    try:
        bpad = importlib.import_module(f'{__package__}.bpad')
        bpad.require()
    except ImportError as exc:
        raise BpadUnavailable(f'Blockpad generator not available: pile_design.bpad could not be imported ({exc}).') from exc
    if not callable(getattr(bpad, 'build_outputs', None)):
        raise BpadUnavailable('Blockpad generator not available: pile_design.bpad has no build_outputs(e) function.')
    return bpad.build_outputs(e)


def bpad_files(e):
    """{file name: bytes} of the Blockpad export: the .bpad named after the support, and its verification.json."""
    out = bpad_outputs(e)
    record = json.dumps(out['verification'], indent=1, ensure_ascii=False, sort_keys=True) + '\n'
    return {out['filename']: bytes(out['document']), VERIFICATION: record.encode('utf-8')}, out['filename']


def bpad_document(e):
    files, name = bpad_files(e)
    return files[name]


WRITERS = {
    CASE: lambda e: case_json(e.case).encode('utf-8'),
    CHECKS: lambda e: checks_csv(e).encode('utf-8'),
    REPORT: _report,
    TABLE: data_table_xlsx,
    MATHCAD: mathcad_xmcd,
    BPAD: bpad_document,
}


def _folder(root):
    stamp = f'pile-{datetime.now(timezone.utc):%Y%m%d-%H%M%S}'
    base = Path(root)
    folder, n = base / stamp, 1
    while folder.exists():
        n += 1
        folder = base / f'{stamp}-{n}'
    folder.mkdir(parents=True, exist_ok=False)
    return folder


def export_file(e, root, name):
    """Write one export (or the full bundle) to a new timestamped folder; originals are never overwritten."""
    if name == BUNDLE:
        return export_bundle(e, root)[BUNDLE]
    if name not in WRITERS:
        raise ValueError(f'Unknown export {name!r}; choose {", ".join([*WRITERS, BUNDLE])}.')
    if name == BPAD:
        files, main = bpad_files(e)       # build first so a failed generator leaves no empty folder
        folder = _folder(root)
        for file_name, data in files.items():
            (folder / file_name).write_bytes(data)
        return folder / main
    data = WRITERS[name](e)               # build first so a failed generator leaves no empty folder
    path = _folder(root) / name
    path.write_bytes(data)
    return path


def export_bundle(e, root):
    """All exports plus the ZIP. If the Mathcad generator is missing, the bundle says so instead of failing."""
    files = {}
    for name, writer in WRITERS.items():
        try:
            if name == BPAD:
                files.update(bpad_files(e)[0])
            else:
                files[name] = writer(e)
        except MathcadUnavailable as exc:
            files[MATHCAD_NOTE] = str(exc).encode('utf-8')
        except BpadUnavailable as exc:
            files[BPAD_NOTE] = str(exc).encode('utf-8')
    folder = _folder(root)
    for name, data in files.items():
        (folder / name).write_bytes(data)
    bundle = folder / BUNDLE
    with zipfile.ZipFile(bundle, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    return {name: folder / name for name in files} | {BUNDLE: bundle}
