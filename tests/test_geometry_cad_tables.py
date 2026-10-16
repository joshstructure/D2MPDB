"""Regression tests for the generated geometry CAD-table workbook (geometry_tables/).

Rebuild after any notebook, Mathcad, roadway XML or generator change:
    .venv\\Scripts\\python.exe scripts\\build_geometry_cad_tables.py
"""
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import numpy as np  # noqa: E402
import openpyxl  # noqa: E402

from geometry_tables import sources, verify, workbook  # noqa: E402

REBUILD = r'Rebuild with: .venv\Scripts\python.exe scripts\build_geometry_cad_tables.py'


class GeometryCadTablesTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.values = openpyxl.load_workbook(workbook.OUTPUT, data_only=True)
        cls.formulas = openpyxl.load_workbook(workbook.OUTPUT)
        cls.cfg, cls.model, cls.env = sources.load_notebook_model()

    def value(self, name):
        return verify.named_value(self.values, name)

    def test_workbook_is_current_with_sources_and_generator(self):
        self.assertEqual(self.value('Build_Fingerprint'), workbook.fingerprint(workbook.source_hashes()), REBUILD)

    def test_values_are_cached_and_error_free(self):
        self.assertEqual(verify.formula_errors(self.values), [])
        self.assertIsNotNone(self.value('Sup_Drop'), 'Workbook was saved without Excel recalculation. ' + REBUILD)

    def test_qc_passes(self):
        self.assertEqual(self.value('QC_Overall'), 'PASS')
        failures = [row for row in verify.qc_rows(self.values) if row[3] == 'FAIL']
        self.assertEqual(failures, [])

    def test_results_match_live_notebook_model(self):
        m = self.model
        expected = {
            'Sup_Drop': m.support_drop, 'Prof_G1': m.g1, 'Prof_G2': m.g2, 'Haunch_Min': m.haunch_min * 12,
            'Pier_Top': m.pier['top'], 'EB1_Top': m.end_bents[0]['top'], 'EB2_Top': m.end_bents[1]['top'],
            'Pier_Len': m.pier['length'], 'EB1_Len': m.end_bents[0]['length'], 'EB1_CapSta': m.end_bents[0]['station'],
            'EB2_CapSta': m.end_bents[1]['station'], 'Sp1_BrgA': m.bearing_limits[0][0], 'Sp2_BrgB': m.bearing_limits[1][1],
        }
        if np.isfinite(m.clearance_min):
            expected['Clr_Min'] = m.clearance_min
        for name, value in expected.items():
            with self.subTest(name=name):
                self.assertAlmostEqual(float(self.value(name)), float(value), delta=1e-8)

    def test_fge_table_matches_notebook_deck_top(self):
        ws = self.values['FGE']
        m = self.model
        checked = 0
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            label = row[1].value
            if not isinstance(label, str) or not (label.startswith('CL Beam') or label.endswith('(deck edge)') or label == 'PGL'):
                continue
            offset = float(row[2].value)
            # The station row is the first "Station (ft)" row above this line.
            r = row[0].row
            while ws.cell(row=r, column=2).value != 'Station (ft)':
                r -= 1
            for c in range(5, 5 + workbook.FGE_INCREMENTS + 3):
                station = float(ws.cell(row=r, column=c).value)
                self.assertAlmostEqual(float(ws.cell(row=row[0].row, column=c).value), float(m.deck_top(station, offset)), delta=1e-9)
                checked += 1
        self.assertEqual(checked, 2 * 5 * (workbook.FGE_INCREMENTS + 3))

    def test_every_tab_is_followed_by_its_method_tab(self):
        names = self.values.sheetnames
        self.assertEqual(names[0], 'Cover')
        for name in names[1:]:
            if name.endswith(' Method'):
                self.assertTrue(self.formulas[name]._images, f'{name} has no equations or schematics')
                continue
            self.assertIn(f'{name} Method', names)
            self.assertEqual(names.index(f'{name} Method'), names.index(name) + 1)

    def test_cad_tables_are_live_formulas(self):
        for sheet in ('CAD FGE', 'CAD Build-Up', 'CAD Beam Variables', 'CAD Bearing Plates', 'CAD Pedestal Elevs'):
            ws = self.formulas[sheet]
            formulas = sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
            with self.subTest(sheet=sheet):
                self.assertGreater(formulas, 20)

    def test_mathcad_reader_matches_registered_camber_basis(self):
        for basis in self.cfg['camber_basis']:
            data = sources.mathcad_beam_data(basis['span'])
            with self.subTest(span=basis['span']):
                self.assertEqual(data['source_sha256'], basis['source_sha256'])
                self.assertAlmostEqual(data['net_camber_in'], basis['net_camber_in'], places=12)
                total = data['defl_120_in'] + data['defl_dl_nc_in'] + data['defl_dl_c_in']
                self.assertAlmostEqual(total, data['net_camber_in'], places=9)
                self.assertEqual(data['strands'], 39)
                self.assertEqual([z['zone'] for z in data['stirrup_zones']], ['A1', 'A2', 'A3', 'S1', 'S2', 'S3', 'S4'])


if __name__ == '__main__':
    unittest.main()
