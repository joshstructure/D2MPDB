"""Pile design notebook: portable package, exports, Pile Data Table workbook, report, widgets and figures."""
import base64
import hashlib
import io
import json
import math
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest import mock
import zipfile
from openpyxl import load_workbook
from pile_design import default_case, evaluate, exports, visuals
from pile_design.model import DATA_TABLE_COLUMNS, PRESETS
from pile_design.report import build_report

ROOT = Path(__file__).resolve().parent.parent
CURRENT = ROOT / 'project_current'

SDM_HEADERS = ['Pier / Bent', 'Pile Size & Type', 'Nominal Bearing Resistance (tons)', 'Tension Resistance (tons)', 'Minimum Tip Elevation (ft)',
               'Test Pile Length / Pile Order Length (ft)', 'Required Jet Elevation (ft)', 'Required Preform Elevation (ft)', 'Factored Design Load (tons)',
               'Downdrag (tons)', 'Total Scour Resistance (tons)', 'Net Scour Resistance (tons)', '100-Year Scour Elevation (ft)',
               'Design Scour Elevation (temporary bridges)', 'Long-Term Scour Elevation (ft)', 'Resistance Factor', 'Pile Cut-Off Elevation (ft)']

_EVALUATIONS = {}


def evaluation(support, load_xml=True):
    key = (support, load_xml)
    if key not in _EVALUATIONS:
        _EVALUATIONS[key] = evaluate(default_case(support, load_xml=load_xml))
    return _EVALUATIONS[key]


def registered_xml_available(support):
    from pile_design.sources import registered_path
    return registered_path(PRESETS[support]['xml']).exists()


class PortableNotebook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from scripts.build_pile_notebook import NOTEBOOK, SETUP_ID, package_files
        cls.notebook = json.loads((ROOT / NOTEBOOK).read_text(encoding='utf-8'))
        cls.setup = ''.join(next(c for c in cls.notebook['cells'] if c.get('id') == SETUP_ID)['source'])
        cls.revision = cls.setup.split('revision = "', 1)[1].split('"', 1)[0]
        cls.package = base64.b64decode(cls.setup.split('PACKAGE = "', 1)[1].split('"', 1)[0])
        cls.files = package_files(ROOT)

    def test_embedded_package_is_current_and_intact(self):
        self.assertEqual(hashlib.sha256(self.package).hexdigest(), self.revision)
        embedded = zipfile.ZipFile(io.BytesIO(self.package))
        for file in self.files:
            name = file.relative_to(ROOT).as_posix()
            # Checkouts may convert line endings; compare content, not the platform newline.
            self.assertEqual(embedded.read(name).replace(b'\r\n', b'\n'), file.read_bytes().replace(b'\r\n', b'\n'),
                             f'{name} changed; run scripts/build_pile_notebook.py')
        self.assertEqual(sorted(embedded.namelist()), sorted(f.relative_to(ROOT).as_posix() for f in self.files))
        self.assertTrue(all(c['outputs'] == [] and c['execution_count'] is None for c in self.notebook['cells'] if c['cell_type'] == 'code'))

    def test_package_contents(self):
        names = set(zipfile.ZipFile(io.BytesIO(self.package)).namelist())
        for required in ('pile_design/widgets.py', 'pile_design/model.py', 'pile_design/exports.py', 'pile_design/report.py', 'pile_design/visuals.py',
                         'pile_design/data/inputs.json', 'pile_design/data/standard_plans.json', 'pier_cap/pile_review.py', 'pier_cap/engine.py',
                         'scripts/mathcad_sheet.py', 'scripts/mathcad_plot.py', 'cap_mathcad/templates/psbeam_check_summary_plot.bin',
                         'cap_mathcad/templates/psbeam_check_summary_plot.json', 'requirements-pier-cap-colab.txt'):
            self.assertIn(required, names)
        self.assertFalse([n for n in names if n.startswith(('tests', 'end_bent_walls'))])
        self.assertEqual(sorted(n for n in names if n.startswith('project_current')),
                         ['project_current/fbmp/end_bents/End Bent 1_Fixed.XML', 'project_current/fbmp/end_bents/End Bent 2_Fixed.XML',
                          'project_current/fbmp/pier/Pier_Fixed.XML'])

    def test_embedded_exports_are_the_registered_bytes(self):
        from pile_design.model import SOURCES
        embedded = zipfile.ZipFile(io.BytesIO(self.package))
        for sid in ('EB-XML', 'EB2-XML', 'PIER-XML'):
            entry = SOURCES['sources'][sid]
            self.assertEqual(hashlib.sha256(embedded.read(entry['path'])).hexdigest(), entry['sha256'], sid)

    def test_builder_refuses_an_export_that_is_not_the_registered_file(self):
        from scripts import build_pile_notebook as builder
        real = builder.registered_exports(ROOT)
        wrong = {path: '0' * 64 for path in real}
        with mock.patch.object(builder, 'registered_exports', return_value=wrong):
            with self.assertRaisesRegex(ValueError, 'does not match the registered SHA-256'):
                builder.package_files(ROOT)

    def test_setup_cell_wiring(self):
        self.assertIn('pile-tool-', self.setup)
        self.assertIn('from pile_design.widgets import PileNotebook', self.setup)
        for module in ('pier_cap', 'pile_design', 'scripts'):
            self.assertIn(f'"{module}"', self.setup)
        sources = ''.join(''.join(c['source']) for c in self.notebook['cells'] if c['cell_type'] == 'code')
        self.assertIn('PileNotebook(_case, export_root=ROOT / "exports")', sources)
        self.assertIn('app.display("sources")', sources)
        self.assertIn('app.display("design")', sources)

    def test_extracted_package_imports_and_builds_without_the_repository(self):
        with tempfile.TemporaryDirectory() as folder:
            zipfile.ZipFile(io.BytesIO(self.package)).extractall(folder)
            code = ('import sys; sys.path.insert(0, %r)\n'
                    'from pathlib import Path\n'
                    'from pile_design.widgets import PileNotebook\n'
                    'from pile_design import evaluate, default_case, exports\n'
                    'import pile_design, pier_cap\n'
                    'assert Path(pile_design.__file__).resolve().is_relative_to(Path(%r).resolve()), pile_design.__file__\n'
                    'app = PileNotebook(export_root=%r)\n'
                    'assert app.e is not None and app.recalc.state == "done"\n'
                    'assert app.e.records and app.case["analysis"]["matches_register"], "embedded registered XML not loaded"\n'
                    'assert app.case["analysis"]["registered_id"] == "EB-XML"\n'
                    'assert "MATCHES SOURCE REGISTER" in app.xml_status.value\n'
                    'for support, family in (("End Bent 2", "psc"), ("Pier", "pipe")):\n'
                    '    e = evaluate(default_case(support))\n'
                    '    assert e.family == family and e.records and e.case["analysis"]["matches_register"], support\n'
                    'print("OK", app.e.status)\n') % (folder, folder, str(Path(folder) / 'exports'))
            run = subprocess.run([sys.executable, '-I', '-c', code], cwd=folder, capture_output=True, text=True, timeout=300)
            self.assertEqual(run.returncode, 0, run.stderr[-2000:])
            self.assertIn('OK', run.stdout)


class ExportsRoundTrip(unittest.TestCase):
    def test_case_json_reload_gives_the_same_check_ratios(self):
        for support in PRESETS:
            e = evaluation(support)
            reloaded = evaluate(exports.load_case(exports.case_json(e.case).encode('utf-8')))
            before = {c.key: (c.ratio, c.status) for c in e.checks}
            after = {c.key: (c.ratio, c.status) for c in reloaded.checks}
            self.assertEqual(before.keys(), after.keys(), support)
            for key, (ratio, status) in before.items():
                if ratio is None:
                    self.assertIsNone(after[key][0], key)
                else:
                    self.assertAlmostEqual(ratio, after[key][0], places=9, msg=f'{support} {key}')
                self.assertEqual(status, after[key][1], f'{support} {key}')
            self.assertEqual(e.data_table, reloaded.data_table)

    def test_load_case_rejects_other_json(self):
        with self.assertRaises(ValueError):
            exports.load_case(b'{"format": "something-else"}')

    def test_checks_csv_has_one_row_per_check(self):
        e = evaluation('Pier')
        rows = exports.checks_csv(e).splitlines()
        self.assertTrue(rows[0].startswith('group,check,status,ratio'))
        import csv
        self.assertEqual(len(list(csv.DictReader(io.StringIO(exports.checks_csv(e))))), len(e.checks))

    def test_export_files_and_bundle(self):
        e = evaluation('End Bent 1')
        with tempfile.TemporaryDirectory() as root:
            first = exports.export_file(e, root, exports.CASE)
            second = exports.export_file(e, root, exports.CASE)
            self.assertNotEqual(first.parent, second.parent, 'a second export must not overwrite the first')
            self.assertTrue(first.parent.name.startswith('pile-'))
            self.assertEqual(json.loads(first.read_text(encoding='utf-8'))['format'], 'd2mpdb-pile-design')
            with self.assertRaises(ValueError):
                exports.export_file(e, root, 'unknown.txt')
            with mock.patch.dict(sys.modules, {'pile_design.mathcad': None}):
                with self.assertRaises(exports.MathcadUnavailable) as raised:
                    exports.export_file(e, root, exports.MATHCAD)
                self.assertIn('Mathcad generator not available', str(raised.exception))
                bundle = exports.export_bundle(e, root)
            names = zipfile.ZipFile(bundle[exports.BUNDLE]).namelist()
            for name in (exports.CASE, exports.CHECKS, exports.REPORT, exports.TABLE):
                self.assertIn(name, names)
            self.assertIn(exports.MATHCAD_NOTE, names)
            self.assertNotIn(exports.MATHCAD, names)

    def test_mathcad_writer_uses_the_shared_generator_when_present(self):
        e = evaluation('End Bent 1')
        fake = mock.Mock()
        fake.build.return_value = '<?xml version="1.0"?><worksheet/>'
        with mock.patch.dict(sys.modules, {'pile_design.mathcad': fake}):
            data = exports.mathcad_xmcd(e)
        fake.build.assert_called_once_with(e)
        self.assertEqual(data, b'<?xml version="1.0"?><worksheet/>')


class PileDataTableWorkbook(unittest.TestCase):
    def sheet(self, e):
        wb = load_workbook(io.BytesIO(exports.data_table_xlsx(e)))
        return wb

    def test_columns_a_to_q_in_sdm_order(self):
        self.assertEqual([header for _, header in DATA_TABLE_COLUMNS], SDM_HEADERS)
        wb = self.sheet(evaluation('End Bent 1'))
        ws = wb['Pile Data Table']
        self.assertEqual([ws.cell(row=1, column=i).value for i in range(1, 18)], SDM_HEADERS)
        self.assertIsNone(ws.cell(row=1, column=18).value)
        self.assertEqual([letter for letter, _ in DATA_TABLE_COLUMNS], list('ABCDEFGHIJKLMNOPQ'))
        self.assertEqual(wb.sheetnames[:2], ['Pile Data Table', 'Basis'])

    def test_not_applicable_entries_are_strings_and_values_follow_the_model(self):
        e = evaluation('End Bent 1')
        ws = self.sheet(e)['Pile Data Table']
        row = {SDM_HEADERS[i]: ws.cell(row=2, column=i + 1).value for i in range(17)}
        self.assertEqual(row, {k: v for k, v in e.data_table.items()})
        self.assertEqual(row['Pier / Bent'], 'End Bent 1')
        self.assertEqual(row['Pile Size & Type'], '18" Sq. Prestr. Conc.')
        for header in ('Minimum Tip Elevation (ft)', 'Required Jet Elevation (ft)', 'Required Preform Elevation (ft)', 'Net Scour Resistance (tons)',
                       '100-Year Scour Elevation (ft)', 'Design Scour Elevation (temporary bridges)', 'Long-Term Scour Elevation (ft)'):
            self.assertEqual(row[header], 'N/A', header)
        self.assertEqual(row['Downdrag (tons)'], 18)
        self.assertEqual(row['Resistance Factor'], 0.75)
        self.assertAlmostEqual(row['Pile Cut-Off Elevation (ft)'], 41.2, places=6)
        pier = {SDM_HEADERS[i]: self.sheet(evaluation('Pier'))['Pile Data Table'].cell(row=2, column=i + 1).value for i in range(17)}
        self.assertEqual(pier['Downdrag (tons)'], 'N/A')
        self.assertEqual(pier['Pile Size & Type'], '20" Ø Pipe')

    def test_provided_geotechnical_values_are_rounded_per_sdg_3_5_15(self):
        case = default_case('End Bent 1')
        case['inputs'].update(EL_min_tip=-31.4, L_test_pile=71.2, EL_scour_100=10.26, EL_scour_LT=12.0)
        e = evaluate(case)
        ws = self.sheet(e)['Pile Data Table']
        row = {SDM_HEADERS[i]: ws.cell(row=2, column=i + 1).value for i in range(17)}
        self.assertEqual(row['Minimum Tip Elevation (ft)'], -32)                    # down
        self.assertEqual(row['Test Pile Length / Pile Order Length (ft)'], '72 / 74')   # up
        self.assertEqual(row['100-Year Scour Elevation (ft)'], 10.3)
        self.assertEqual(row['Long-Term Scour Elevation (ft)'], 12.0)

    def test_basis_sheet_documents_the_equation_and_sources(self):
        e = evaluation('End Bent 1')
        wb = self.sheet(e)
        text = ' '.join(str(c.value) for row in wb['Basis'].iter_rows() for c in row if c.value)
        for needle in ('Rn = (Factored Design Load + Net Scour Resistance + Downdrag) / φ', 'γDR', 'excluded', 'SHA-256', 'Geotechnical engineer', 'Table 3.5.13-1'):
            self.assertIn(needle, text)
        self.assertIn('Inputs', wb.sheetnames)

    def test_export_without_fbmp_results_still_writes_every_column(self):
        ws = self.sheet(evaluation('Pier', load_xml=False))['Pile Data Table']
        self.assertEqual([ws.cell(row=1, column=i).value for i in range(1, 18)], SDM_HEADERS)
        self.assertEqual(ws.cell(row=2, column=9).value, 'N/A')


class CalculationReport(unittest.TestCase):
    def test_report_sections_equations_and_figures(self):
        expected = {'End Bent 1': 'Interaction diagram and curve table', 'Pier': 'Interaction envelope'}
        for support, interaction in expected.items():
            html = build_report(evaluation(support))
            for heading in ('Overview', 'Result summary', 'FBMP demands and governing records', interaction, 'Pile Data Table', 'Check register',
                            'Pending items', 'Sources, decisions and confirmations', '1. Pile layout and downdrag', '8. Geotechnical resistance and Pile Data Table'):
                self.assertIn(heading, html, f'{support}: {heading}')
            self.assertGreater(html.count('<math'), 60, support)
            self.assertGreaterEqual(html.count('Plotly.newPlot'), 5, support)
            self.assertNotIn('<script src=', html)                       # self-contained: Plotly is inlined
            self.assertIn('Equation &amp; values', html)
            self.assertIn('Factored Design Load', html)

    def test_psc_report_has_the_interaction_curve_table(self):
        html = build_report(evaluation('End Bent 1'))
        self.assertIn('c (in)', html)
        self.assertIn('M<sub>n</sub> (kip-ft)', html)
        self.assertIn('Index 455-018', html)
        self.assertIn('Prestressed pile', html)

    def test_pipe_report_has_envelope_vertices(self):
        html = build_report(evaluation('Pier'))
        self.assertIn('Pure compression', html)
        self.assertIn('6.9.2.2.1', html)

    def test_report_without_fbmp_results_builds(self):
        html = build_report(evaluation('End Bent 1', load_xml=False))
        self.assertIn('No FBMP results are loaded', html)


class FiguresBuild(unittest.TestCase):
    def test_every_figure_builds_for_both_presets(self):
        for support in ('End Bent 1', 'Pier'):
            e = evaluation(support)
            for name, build in visuals.FIGURES:
                fig = build(e)
                self.assertGreater(len(fig.data), 0, f'{support} {name}')
            self.assertIn('Pile Cut-Off Elevation', visuals.data_table_html(e))

    def test_figures_build_without_fbmp_results(self):
        for support in ('End Bent 1', 'Pier'):
            e = evaluation(support, load_xml=False)
            for name, build in visuals.FIGURES:
                build(e)

    def test_interaction_figure_marks_the_governing_demand(self):
        fig = visuals.interaction_figure(evaluation('End Bent 1'))
        names = [t.name for t in fig.data]
        self.assertIn('Governing resultant-axis demand', names)
        self.assertTrue(any(n and n.startswith('φ-curve (Mr, Pr)') for n in names))
        pipe = [t.name for t in visuals.interaction_figure(evaluation('Pier')).data]
        self.assertIn('Governing demand', pipe)
        self.assertTrue(any(n and 'Eq. 6.9.2.2.1' in n for n in pipe))

    def test_profile_marks_governing_records(self):
        names = {t.name for t in visuals.demand_profile_figure(evaluation('End Bent 1')).data}
        self.assertTrue({'Governing axial', 'Governing combined', 'Governing shear'} <= names)


class WidgetsHeadless(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import ipywidgets as W
        from pile_design.widgets import PileNotebook
        cls.W = W
        cls.root = tempfile.TemporaryDirectory()
        cls.app = PileNotebook(export_root=cls.root.name)

    @classmethod
    def tearDownClass(cls):
        cls.app.close()
        cls.root.cleanup()

    def setUp(self):
        self.app.load(default_case('End Bent 1'), 'test reset')

    def titles(self):
        tabs = self.app.input_tabs
        return [tabs.get_title(i) for i in range(len(tabs.children))]

    def test_builds_and_evaluates(self):
        app = self.app
        self.assertIsNotNone(app.e)
        self.assertEqual(app.recalc.state, 'done')
        self.assertEqual([app.result_tabs.get_title(i) for i in range(len(app.result_tabs.children))],
                         ['Section', 'Demands', 'Interaction', 'Shear & service', 'Pile Data Table', 'D/C checks', 'Equations', 'Exports'])
        self.assertIn('Governing structural D/C', app.banner.value)
        self.assertIn('Table 3.5.13-1', app.banner.value)
        self.assertIn('Downdrag used', app.banner.value)
        self.assertTrue(all(not b.disabled for b in app.export_buttons))
        self.assertEqual(len(app.export_buttons), 6)

    def test_categorical_inputs_are_dropdowns_and_optional_inputs_are_text(self):
        W = self.W
        c = self.app.controls
        for name in ('soil_type', 'environment', 'phi_geo'):
            self.assertIsInstance(c[name], W.Dropdown, name)
        self.assertEqual([label for label, _ in c['soil_type'].options], ['Sand (nh)', 'Clay (Su)'])
        self.assertEqual([value for _, value in c['environment'].options], [1, 2, 3])
        for name in ('EL_scour_100', 'EL_scour_LT', 'EL_min_tip', 'L_test_pile'):
            self.assertIsInstance(c[name], W.Text, name)
            self.assertEqual(c[name].value, '')

    def test_blank_optional_input_is_na_and_a_number_is_accepted(self):
        app = self.app
        app.controls['EL_min_tip'].value = '-31.4'
        self.assertEqual(app.case['inputs']['EL_min_tip'], -31.4)
        self.assertEqual(app.recalc.state, 'dirty')
        app.recalculate()
        self.assertEqual(app.e.data_table['Minimum Tip Elevation (ft)'], -32)
        app.controls['EL_min_tip'].value = ' '
        self.assertIsNone(app.case['inputs']['EL_min_tip'])
        app.controls['EL_min_tip'].value = 'n/a'
        app.recalculate()
        self.assertEqual(app.e.data_table['Minimum Tip Elevation (ft)'], 'N/A')
        app.controls['EL_min_tip'].value = 'not a number'
        self.assertEqual(app.recalc.state, 'invalid')
        self.assertIsNone(app.case['inputs']['EL_min_tip'])

    def test_family_dropdown_shows_and_hides_the_family_input_tabs(self):
        app = self.app
        self.assertIn('Prestress and strand', self.titles())
        self.assertNotIn('Steel pipe pile', self.titles())
        self.assertEqual(app.psc_row.layout.display, '')
        app.family_dd.value = 'pipe'
        self.assertEqual(app.case['family'], 'pipe')
        self.assertIn('Steel pipe pile', self.titles())
        self.assertNotIn('Prestress and strand', self.titles())
        self.assertEqual(app.psc_row.layout.display, 'none')
        self.assertEqual(app.case['inputs']['D_pile'], app.case['inputs']['D_out'])
        self.assertIn('What-if', app.family_info.value)
        app.recalculate()
        self.assertEqual(app.recalc.state, 'done')
        self.assertEqual(app.e.family, 'pipe')
        app.family_dd.value = 'psc'
        app.recalculate()
        self.assertEqual(app.e.family, 'psc')
        self.assertEqual(app.case['inputs']['D_pile'], 18.0)

    def test_standard_plans_index_and_pattern_dropdowns(self):
        app = self.app
        self.assertEqual([v for _, v in app.index_dd.options], ['455-014', '455-018', '455-024', '455-030'])
        app.index_dd.value = '455-024'
        self.assertEqual(app.case['psc']['index'], '455-024')
        self.assertEqual(app.case['inputs']['D_pile'], 24.0)
        self.assertEqual([v for _, v in app.pattern_dd.options][0], '16~0.6')
        app.pattern_dd.value = '24~1/2'
        self.assertEqual(app.case['psc']['pattern'], '24~1/2')
        app.recalculate()
        self.assertEqual(app.recalc.state, 'done')
        self.assertEqual(app.e.data_table['Pile Size & Type'], '24" Sq. Prestr. Conc.')

    def test_preset_requires_confirmation(self):
        app = self.app
        app.preset_dd.value = 'Pier'
        app._ask_preset()
        self.assertEqual(app.ask_box.layout.display, '')
        self.assertEqual(app.case['support'], 'End Bent 1')          # nothing replaced yet
        app._ask_no()
        self.assertEqual(app.ask_box.layout.display, 'none')
        self.assertEqual(app.case['support'], 'End Bent 1')
        app._ask_preset()
        app._ask_yes()
        self.assertEqual(app.case['support'], 'Pier')
        self.assertEqual(app.case['family'], 'pipe')
        self.assertEqual(app.e.family, 'pipe')

    def test_confirmation_with_evidence_note_and_pipe_only_visibility(self):
        app = self.app
        self.assertEqual(app.confirm_rows['steel_grade'].layout.display, 'none')
        app.confirm['exposure'].value = True
        self.assertTrue(app.case['confirmations']['exposure'])
        self.assertIn('Manually confirmed', app.case['confirmation_evidence']['exposure']['note'])
        app.family_dd.value = 'pipe'
        self.assertEqual(app.confirm_rows['steel_grade'].layout.display, '')

    def test_every_input_table_lists_every_input(self):
        import html
        from pile_design.model import INPUT_SPEC
        table = self.app.sources_table.value
        for spec in INPUT_SPEC:
            self.assertIn(html.escape(spec['caption']), table, spec['name'])

    @unittest.skipUnless(registered_xml_available('Pier'), 'registered FBMP exports are not in this checkout')
    def test_xml_upload_applies_the_registered_receipt(self):
        from pile_design.sources import registered_path
        app = self.app
        path = registered_path('PIER-XML')
        app._load_source('FBMP-XML', path.read_bytes(), path.name)
        self.assertEqual(app.case['family'], 'pipe')
        self.assertTrue(app.case['analysis']['matches_register'])
        self.assertIn('MATCHES SOURCE REGISTER', app.xml_status.value)
        self.assertIn('PIER-XML', app.xml_status.value)
        self.assertEqual(app.recalc.state, 'dirty')
        app.recalculate()
        self.assertEqual(app.e.family, 'pipe')
        app._load_source('FBMP-XML', b'<not-xml', 'broken.xml')
        self.assertIn('SOURCE NOT APPLIED', app.xml_status.value)
        self.assertEqual(app.case['family'], 'pipe')                 # the failed upload changed nothing

    def test_case_json_upload_restores_the_case(self):
        app = self.app
        saved = exports.case_json(app.case).encode('utf-8')
        app.controls['N_pile'].value = 6
        self.assertEqual(app.case['inputs']['N_pile'], 6)
        app._load_case(saved, 'saved.json')
        self.assertEqual(app.case['inputs']['N_pile'], 4)
        self.assertIn('CASE LOADED', app.case_status.value)
        app._load_case(b'{"format": "x"}', 'bad.json')
        self.assertIn('CASE NOT LOADED', app.case_status.value)

    def test_exports_write_files(self):
        app = self.app
        fake = mock.Mock()                              # the Mathcad generator has its own tests (tests.test_pile_mathcad)
        fake.build.return_value = '<?xml version="1.0"?><worksheet/>'
        with mock.patch.dict(sys.modules, {'pile_design.mathcad': fake}):
            for name in (exports.CASE, exports.CHECKS, exports.REPORT, exports.TABLE, exports.MATHCAD, exports.BUNDLE):
                app._export(name)
                self.assertTrue(app.last_files[name].is_file(), name)
        self.assertIn('EXPORT WRITTEN', app.export_status.value)
        self.assertIn(exports.MATHCAD, zipfile.ZipFile(app.last_files[exports.BUNDLE]).namelist())

    def test_widgets_without_fbmp_results(self):
        app = self.app
        app.load(default_case('Pier', load_xml=False), 'no xml')
        self.assertEqual(app.recalc.state, 'done')
        self.assertIn('No FBMP results are loaded', app.banner.value)
        self.assertIn('NO FBMP RESULTS LOADED', app.xml_status.value)


if __name__ == '__main__':
    unittest.main()
