"""Geotechnical report contents and complete soil/trial exports."""
from copy import deepcopy
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from lxml import etree as E
from openpyxl import load_workbook

from pier_cap.geotech_handoff import write_handoff_xlsx
from pier_cap.pile_review import import_pile_xml, parse_trials, evaluate_trials, validate_saved_review
from pier_cap.soil_properties import soil_properties_from_xml, soil_report_rows, validate_soil_properties


SOIL = '''<SOIL_INPUT_DATA>
<AXIAL><LAYER number="1" soilset="1"><MODEL>1</MODEL>
<TOP><ELEV units="ft">21.5</ELEV><G units="ksi">0.6</G><POISSON>0.3</POISSON></TOP>
<BOTTOM><ELEV units="ft">-1.5</ELEV><G units="ksi">1.2</G><POISSON>0.35</POISSON></BOTTOM>
</LAYER></AXIAL>
<LATERAL><LAYER number="1" soilset="1"><MODEL>2</MODEL>
<TOP><PHI units="deg">30</PHI><RK units="lbs/in^3">31</RK></TOP>
<BOTTOM><PHI units="deg">35</PHI><RK units="lbs/in^3">32</RK></BOTTOM></LAYER></LATERAL>
<TORSION><LAYER number="1" soilset="1"><MODEL>1</MODEL>
<TOP><TAU_MAX units="psf">418</TAU_MAX></TOP><BOTTOM><TAU_MAX units="psf">600</TAU_MAX></BOTTOM>
</LAYER></TORSION>
<TIP><SOIL_SET number="1" model="1"><Q_ULT units="kips">9999</Q_ULT></SOIL_SET></TIP>
<SOIL_MOVEMENT_OPTION><SOIL_SET number="1"><LAYER number="1"><OPTION_FLAG>0</OPTION_FLAG></LAYER></SOIL_SET></SOIL_MOVEMENT_OPTION>
<INITIAL_OFFSETS_XP_AND_YP><SOIL_SET number="1"><LAYER number="1">
<OFFSET_TOP_XP units="in">0</OFFSET_TOP_XP><OFFSET_BOT_XP units="in">0.2</OFFSET_BOT_XP>
</LAYER></SOIL_SET></INITIAL_OFFSETS_XP_AND_YP>
<P_SCALE_FACTORS><SOIL_SET number="1"><LAYER number="1">
<PSCALE_FACTOR_TOP_XP>1</PSCALE_FACTOR_TOP_XP><PSCALE_FACTOR_BOT_XP>0.8</PSCALE_FACTOR_BOT_XP>
</LAYER></SOIL_SET></P_SCALE_FACTORS>
<CUSTOM><SOIL_SET number="2"><CURVE name="curve A"><POINT number="1"><P units="kip">4</P></POINT>
<POINT number="2"><P units="kip">8</P></POINT></CURVE><DESCRIPTION>=1+1</DESCRIPTION></SOIL_SET></CUSTOM>
</SOIL_INPUT_DATA>'''


def soil_group():
    group = E.Element('PILE_GROUP')
    group.append(E.fromstring(SOIL))
    return group


class SoilPropertyTests(unittest.TestCase):
    def test_all_leaf_values_model_attributes_and_custom_data_are_retained(self):
        group = soil_group()
        rows = soil_properties_from_xml(group)
        leaves = group.xpath('SOIL_INPUT_DATA//*[not(*)]')
        # Non-context attributes: tip model, curve name, and both point numbers.
        self.assertEqual(len(rows), len(leaves)+4)
        for node in leaves:
            raw = node.text
            value = raw if raw == '=1+1' else float(raw)
            matches = [r for r in rows if r['property'].split('/')[-1] == node.tag and r['value'] == value]
            self.assertTrue(matches, E.tostring(node))
            self.assertTrue(any(r['units'] == node.get('units', '') for r in matches))
        self.assertTrue(any(r['category']=='TIP' and r['property']=='MODEL' and r['value']==1 for r in rows))
        points = [r['property'] for r in rows if r['property'].endswith('/P')]
        self.assertEqual(points, ['CURVE/POINT[1]/P', 'CURVE/POINT[2]/P'])
        pivot = soil_report_rows(rows)
        modulus = next(r for r in pivot if r['category']=='Axial' and r['property']=='Shear modulus (G)')
        self.assertEqual((modulus['top'], modulus['bottom'], modulus['units']), (0.6,1.2,'ksi'))
        offset = next(r for r in pivot if r['property']=='Initial XP offset')
        self.assertEqual((offset['top'],offset['bottom']), (0,0.2))
        self.assertEqual(sum(k in row for row in pivot for k in ('top','bottom','value')),len(rows))

    def test_import_and_json_round_trip_preserve_soil_records(self):
        root=E.parse(str(Path(__file__).parent/'fixtures/fbmp_610_piles.xml')).getroot()
        root.find('.//PILE_GROUP').append(E.fromstring(SOIL))
        review=import_pile_xml(E.tostring(root))
        self.assertEqual(review['soil_properties'], soil_properties_from_xml(soil_group()))
        saved=json.loads(json.dumps(review))
        validate_saved_review(saved)
        self.assertEqual(saved['soil_properties'], review['soil_properties'])
        saved.pop('soil_properties')
        validate_saved_review(saved)  # Old saved studies remain usable.

    def test_invalid_properties_reject_import_or_restore(self):
        group=soil_group()
        group.find('.//TOP/G').text='NaN'
        with self.assertRaisesRegex(ValueError,'nonfinite'):
            soil_properties_from_xml(group)
        for value in (None, [], True, float('inf')):
            rows=soil_properties_from_xml(soil_group())
            rows[0]['value']=value
            with self.assertRaisesRegex(ValueError,'soil property value'):
                validate_soil_properties(rows)


class GeotechReportTests(unittest.TestCase):
    def data(self):
        trial_text=('trial,combination,pile,embedment_ft,displacement_in,series,dc,dx_in,dy_in,converged,iterations,tolerance_kip\n'
                    '1,2,4,20,1,Study A,0.8,0.6,0.8,yes,5,0.001\n'
                    '2,3,1,15,1.05,Study A,1.1,,,no,,')
        result=evaluate_trials(parse_trials(trial_text))
        return dict(section=[dict(item='Pile count',value=4)], loads=[dict(load='Compression',short_tons=75)],
                    selected=True, elevations=dict(soil_top_ft=21.5,pile_head_ft=41.4175,cutoff_ft=40.25),
                    soil_properties=soil_properties_from_xml(soil_group()),trials=result['rows'])

    def test_three_sheets_include_geometry_soils_and_every_trial(self):
        data=self.data()
        result=dict(tip_elevation_ft=-30,total_length_ft=71,comparison_complete=True,
                    controlling_criterion='Second zero crossing')
        with TemporaryDirectory() as folder:
            path=write_handoff_xlsx(Path(folder)/'Pier 2.xlsx',data,result,20,trial_source='Pier 2')
            wb=load_workbook(path)
            self.addCleanup(wb.close)
            self.assertEqual(wb.sheetnames,['Pile summary','Soil properties','Minimum tip trials'])
            for sheet in wb:
                self.assertEqual(sheet['A1'].value,'Pier 2')
                self.assertNotIn('handoff',str(list(sheet.values)).lower())
                self.assertFalse(any(c.data_type=='f' for row in sheet for c in row))
            summary={r[0]:r[1] for r in wb.active.values}
            for label,value in [('Soil-top elevation (ft)',21.5),('Pile-head elevation (ft)',41.4175),
                                ('Total pile length at minimum tip (ft)',71),('Adopted pile cutoff elevation (ft)',40.25),
                                ('Design ground / scour elevation (ft)',20)]:
                self.assertEqual(summary[label],value)
            self.assertIn('Total pile length is rounded up to whole feet.',summary)
            soil_rows=list(wb['Soil properties'].values)[4:]
            self.assertEqual(len(soil_rows),len(soil_report_rows(data['soil_properties'])))
            self.assertTrue(any(r[3]=='Shear modulus (G)' and r[4:7]==('ksi',0.6,1.2) for r in soil_rows))
            trial=wb['Minimum tip trials']
            records=[dict(zip([c.value for c in trial[4]],r)) for r in list(trial.values)[4:]]
            self.assertEqual(len(records),2)
            self.assertEqual(records[0]['Load combination'],'2')
            self.assertEqual(records[0]['Embedment (ft)'],20)
            self.assertEqual(records[0]['DX (in)'],0.6)
            self.assertEqual(records[0]['Converged'],'Yes')
            self.assertEqual(records[1]['Converged'],'No')
            self.assertEqual(records[1]['DX (in)'],'—')
            self.assertAlmostEqual(records[0]['Change to next shallower trial (in)'],0.05)
            self.assertEqual(records[1]['Change to next shallower trial (in)'],'—')
            self.assertEqual(records[0]['Trial Lcrit'],'Yes')

    def test_missing_soil_or_trial_data_never_appear_as_zero(self):
        data=self.data()
        data.pop('soil_properties');data.pop('trials')
        with TemporaryDirectory() as folder:
            path=write_handoff_xlsx(Path(folder)/'report.xlsx',data,None,20)
            wb=load_workbook(path)
            self.addCleanup(wb.close)
            self.assertEqual(wb.active['A1'].value,'Geotechnical report')
            self.assertIn('unavailable',wb['Soil properties']['A5'].value)
            self.assertIn('No minimum-tip trial data',wb['Minimum tip trials']['A5'].value)


if __name__=='__main__':unittest.main()
