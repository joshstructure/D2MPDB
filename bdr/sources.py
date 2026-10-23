"""Read-only access to the project record for the Bridge Development Report.

Everything the report states about the bridge is read here from registered
project sources (the journal, the source register, decisions, the assumptions
register, the coherence review, the saved cases, the pile data tables, the
check registers, the Mathcad beam runs, the geotechnical sections snapshot and
the substructure drawing manifests). Nothing is recalculated; saved results are
read as saved. The facts the record cannot supply live in ``data/facts.json``.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import sys
import xml.etree.ElementTree as ET
from fractions import Fraction
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PKG = Path(__file__).resolve().parent
PC = ROOT / "project_current"

# Registered and generated files the report reads. The manifest records each
# file's SHA-256 so the staleness test can tell when the record moved on.
SOURCES = {
    "JOURNAL": ROOT / "journal" / "D2_I95_Project_Journal.bpad",
    "REGISTER": PC / "source_register.csv",
    "DECISIONS": PC / "decisions.md",
    "ASSUMPTIONS": PC / "coherence" / "assumptions.json",
    "LATEST-REVIEW": PC / "coherence" / "latest_review.md",
    "PIER-CASE": PC / "cap_design" / "pier" / "selected_case.json",
    "EB-CASE": PC / "cap_design" / "end_bents" / "selected_case.json",
    "EB2-CASE": PC / "cap_design" / "end_bent_2" / "selected_case.json",
    "BACKWALL-CASE": PC / "cap_design" / "backwall" / "selected_case.json",
    "PIER-WALL-CASE": PC / "pier_wall" / "selected_case.json",
    "CHEEK-CASE": PC / "pier_cheek_wall" / "selected_case.json",
    "PILE-EB1-CASE": PC / "pile_design" / "end_bent_1" / "selected_case.json",
    "PILE-EB2-CASE": PC / "pile_design" / "end_bent_2" / "selected_case.json",
    "PILE-PIER-CASE": PC / "pile_design" / "pier" / "selected_case.json",
    "PILE-EB1-TABLE": PC / "pile_design" / "end_bent_1" / "End Bent 1 Pile Data Table 10-22.xlsx",
    "PILE-EB2-TABLE": PC / "pile_design" / "end_bent_2" / "End Bent 2 Pile Data Table 10-22.xlsx",
    "PILE-PIER-TABLE": PC / "pile_design" / "pier" / "Pier Pile Data Table 10-22.xlsx",
    "PILE-EB1-CHECKS": PC / "pile_design" / "end_bent_1" / "End Bent 1 Pile Design 10-22 checks.csv",
    "PILE-EB2-CHECKS": PC / "pile_design" / "end_bent_2" / "End Bent 2 Pile Design 10-22 checks.csv",
    "PILE-PIER-CHECKS": PC / "pile_design" / "pier" / "Pier Pile Design 10-22 checks.csv",
    "PIER-WALL-CHECKS": PC / "pier_wall" / "Pier Wall Design 10-22c checks.csv",
    "CHEEK-CHECKS": PC / "pier_cheek_wall" / "Pier Cheek Wall Design 10-22 checks.csv",
    "PIER-README": PC / "cap_design" / "pier" / "README.md",
    "EB-README": PC / "cap_design" / "end_bents" / "README.md",
    "EB2-README": PC / "cap_design" / "end_bent_2" / "README.md",
    "BACKWALL-README": PC / "cap_design" / "backwall" / "README.md",
    "STRAP-LOADS": ROOT / "substructure_drawings" / "data" / "strap_loads.json",
    "MCAD-S1": PC / "mathcad" / "PrestressedBeamV6.2_Pedestrian_2Beam_Span1.xmcd",
    "MCAD-S2": PC / "mathcad" / "PrestressedBeamV6.2_Pedestrian_2Beam_Span2.xmcd",
    "GEOTECH-SNAPSHOT": ROOT / "geotech_sections" / "data" / "sections_snapshot.json",
    "GEOTECH-MANIFEST": ROOT / "geotech_sections" / "output" / "manifest.json",
    "DXF-MANIFEST": ROOT / "substructure_drawings" / "output" / "manifest.json",
    "PLAN-MANIFEST": ROOT / "substructure_drawings" / "output" / "model_plan_manifest.json",
    "PLAN-PNG": ROOT / "substructure_drawings" / "output" / "B4_Model_Plan.png",
    "BAR-REGISTER": ROOT / "substructure_drawings" / "output" / "bar_mark_register.csv",
    "TYPSEC-MANIFEST": ROOT / "substructure_drawings" / "output" / "typical_section_manifest.json",
    "TYPSEC-PNG": ROOT / "substructure_drawings" / "output" / "B4_Typical_Section.png",
    "FACTS": PKG / "data" / "facts.json",
}
PREVIEWS = ROOT / "substructure_drawings" / "output" / "previews"
FBMP_README = PC / "fbmp" / "README.md"


def sha256(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_json(path: Path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


# ------------------------------------------------------------------ journal --
FRACTION = re.compile(r"(\d+)\s+(\d+)/(\d+)|(\d+)/(\d+)")


def _number(text: str) -> float:
    text = text.strip()
    m = re.fullmatch(r"(\d+)\s+(\d+)/(\d+)", text)
    if m:
        return int(m[1]) + Fraction(int(m[2]), int(m[3]))
    m = re.fullmatch(r"(\d+)/(\d+)", text)
    if m:
        return float(Fraction(int(m[1]), int(m[2])))
    return float(text)


def length_ft(text: str):
    """'19 ft 2 in', '107.35 ft', '42 in', '2 ft 2.562 in', '6 ft 9.5 in' -> feet; else None."""
    if text is None:
        return None
    t = text.strip().lstrip("=").strip()
    m = re.fullmatch(r"(-?[\d.]+(?:\s+\d+/\d+)?)\s*ft(?:\s+([\d.]+(?:\s+\d+/\d+)?)\s*in)?", t)
    if m:
        return float(_number(m[1])) + (float(_number(m[2])) / 12.0 if m[2] else 0.0)
    m = re.fullmatch(r"(-?[\d.]+(?:\s+\d+/\d+)?)\s*in", t)
    if m:
        return float(_number(m[1])) / 12.0
    m = re.fullmatch(r"(-?[\d.]+)'\s*(?:([\d.]+(?:\s+\d+/\d+)?)\")?", t)
    if m:
        return float(m[1]) + (float(_number(m[2])) / 12.0 if m[2] else 0.0)
    return None


def plain_number(text: str):
    if text is None:
        return None
    t = text.strip().lstrip("=").strip()
    m = re.fullmatch(r"(-?[\d.]+(?:[eE][+-]?\d+)?)(?:\s+\S+)?", t)
    return float(m[1]) if m else None


class JV:
    """A journal definition: its formula, the saved result text and parsed numbers."""

    def __init__(self, module, symbol, formula, result):
        self.module, self.symbol, self.formula = module, symbol, formula
        rhs = formula.split("=", 1)[1].strip() if "=" in formula else ""
        self.rhs = rhs
        self.result = (result or "").strip()
        # Display text: the saved result when there is one, else the literal.
        shown = self.result.lstrip("=").strip() if self.result else rhs
        self.text = shown[1:-1] if len(shown) >= 2 and shown[0] == '"' and shown[-1] == '"' else shown

    @property
    def ft(self):
        return length_ft(self.text)

    @property
    def num(self):
        return plain_number(self.text)

    @property
    def inches(self):
        f = self.ft
        return None if f is None else f * 12.0

    def __str__(self):
        return self.text


def _cell_text(cell):
    cached = cell.find("num")
    if cached is None:
        cached = cell.find("textvalue")
    if cached is None:
        return ""
    return " ".join(" ".join(cached.itertext()).split())


class Journal:
    def __init__(self, path: Path):
        self.path = path
        doc = ET.parse(path).getroot()
        self.modules = []          # (tag, name, title)
        self.defs = {}             # module -> {symbol: JV}  (first definition wins)
        self.sheets = {}           # spreadsheet name -> [[cell text...], ...]
        for m in doc:
            if m.tag not in ("report", "spreadsheet"):
                continue
            name, title = m.get("name", ""), m.get("title", m.get("name", ""))
            self.modules.append((m.tag, name, title))
            if m.tag == "report":
                defs = self.defs.setdefault(name, {})
                for item in m.iter("dynexp"):
                    formula = item.get("formula", "")
                    found = re.match(r"^\s*([\w]+)\s*=(?!=)", formula)
                    if not found:
                        continue
                    res = item.find("expresult")
                    rt = " ".join(" ".join(res.itertext()).split()) if res is not None else ""
                    defs.setdefault(found[1], JV(name, found[1], formula, rt))
            else:
                rows = []
                for row in m.findall("row"):
                    rows.append([_cell_text(c) for c in row if c.tag in ("c", "cell")])
                self.sheets[name] = rows

    def value(self, module: str, symbol: str) -> JV:
        return self.defs[module][symbol]

    def text(self, module, symbol, default="—"):
        try:
            return self.value(module, symbol).text or default
        except KeyError:
            return default

    def calc_modules(self):
        out = []
        for tag, name, title in self.modules:
            m = re.match(r"^(C\d{3}) - (.+)$", title)
            if m:
                out.append({"number": m[1], "title": m[2], "module": name})
        return out

    def sheet_volumes(self, name):
        """(label, ft^3 value) pairs from a quantities sheet: the last ft^3 cell of a row,
        labelled by the row's first text cell or, when that is blank, the previous label."""
        out, label = [], ""
        for row in self.sheets.get(name, []):
            first = next((c for c in row if c), "")
            vols = [c for c in row if re.fullmatch(r"[\d.]+ ft\^3", c)]
            if first and not re.fullmatch(r"[\d.]+ ft\^[23]", first):
                label = first
            if vols:
                out.append((label, float(vols[-1].split()[0])))
        return out

    def mse_walls(self):
        walls = []
        for row in self.sheets.get("QuantitiesMSEWall", []):
            if row and row[0].startswith("MSE Wall "):
                area = next((c for c in row[1:] if c.endswith("ft^2")), "")
                base = next((c for c in row[1:] if "ft" in c and not c.endswith("ft^2")), "")
                if area:
                    walls.append({"wall": row[0].replace("MSE Wall ", ""), "base_length": base,
                                  "area_ft2": float(area.split()[0])})
        totals = {}
        for row in self.sheets.get("QuantitiesMSEWall", []):
            if "Sum" in row:
                totals["area_ft2"] = float(next(c for c in row if c.endswith("ft^2")).split()[0])
        return walls, totals


# ------------------------------------------------------- project_current --
def read_register(path: Path):
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def read_decisions(path: Path):
    rows = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.startswith("| D-"):
            cells = [c.strip() for c in line.strip().strip("|").split(" | ")]
            if len(cells) >= 4:
                rows.append({"id": cells[0], "date": cells[1], "decision": cells[2], "basis": " | ".join(cells[3:])})
    return rows


def read_latest_review(path: Path):
    text = path.read_text(encoding="utf-8")
    checked = re.search(r"Checked at (\S+)\.", text)
    outstanding = re.search(r"\*\*(\d+) outstanding items remain\.\*\*", text)
    items = []
    section = text.split("## Outstanding items", 1)[1].split("## ", 1)[0]
    for line in section.splitlines():
        if line.startswith("| COH-"):
            cells = [c.strip() for c in line.strip().strip("|").split(" | ")]
            items.append({"id": cells[0], "status": cells[1], "subject": cells[2], "next": " | ".join(cells[3:])})
    registered = re.search(r"Registered sources: (\d+)\. Tracker items: (\d+)\.", text)
    return {"checked_at": checked[1] if checked else "", "outstanding": int(outstanding[1]) if outstanding else len(items),
            "items": items, "registered_sources": int(registered[1]) if registered else None,
            "tracker_items": int(registered[2]) if registered else None}


def read_pile_table(path: Path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {"columns": [], "row": [], "basis": {}, "inputs": {}}
    ws = wb["Pile Data Table"]
    rows = [r for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]
    out["columns"], out["row"] = list(rows[0]), list(rows[1])
    for r in wb["Basis"].iter_rows(values_only=True):
        if r[0] and r[0] != "Item":
            out["basis"][r[0]] = r[1]
    for r in wb["Inputs"].iter_rows(values_only=True):
        if r[0] and r[0] != "Input":
            out["inputs"][r[0]] = {"value": r[1], "unit": r[2], "meaning": r[3], "source": r[4], "status": r[5]}
    return out


def read_checks(path: Path):
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def pile_check_summary(rows):
    # Structural strength and service checks only: not the policy checks (pile size, spacing),
    # the geotechnical rows, the prerequisites or the FBMP reference row.
    strength = re.compile(r"axial|flexure|shear|tension|compression|buckling|slenderness", re.I)
    structural = [r for r in rows if r["group"] not in ("Geotechnical", "Prerequisites", "Reference")
                  and r["status"].split(" ")[0] in ("PASS", "FAIL")
                  and r["ratio"] not in ("", None) and strength.search(r["check"])]
    best = max(structural, key=lambda r: float(r["ratio"]), default=None)
    counts = {}
    for r in rows:
        key = r["status"].split(" ")[0]
        counts[key] = counts.get(key, 0) + 1
    pending = [r["check"] for r in rows if r["status"].startswith("PENDING")
               or "PENDING" in r["status"]]
    return {"max_structural_ratio": float(best["ratio"]) if best else None,
            "governing_check": best["check"] if best else None, "governing_note": (best or {}).get("note", ""),
            "counts": counts, "pending": pending}


def generic_check_summary(rows, id_key="id", title_key="title", status_key="status", ratio_key="ratio"):
    counts, failed, pending, maxrow = {}, [], [], None
    for r in rows:
        st = (r.get(status_key) or "").split(" ")[0]
        counts[st] = counts.get(st, 0) + 1
        if st == "FAIL":
            failed.append(r.get(id_key) or r.get(title_key))
        if "PENDING" in (r.get(status_key) or ""):
            pending.append(r.get(id_key) or r.get(title_key))
        try:
            ratio = float(r.get(ratio_key) or "")
        except ValueError:
            continue
        if st == "PASS" and (maxrow is None or ratio > maxrow[0]):
            maxrow = (ratio, r.get(title_key) or r.get(id_key) or r.get("check"))
    return {"counts": counts, "failed": failed, "pending": pending, "max_pass": maxrow}


def cap_readme_counts(path: Path):
    text = path.read_text(encoding="utf-8")
    fail = re.search(r"(\d+)\s+FAIL entries", text)
    pend = re.search(r"(\d+)\s+pending/conditional", text)
    decision = re.search(r"selected under (D-\d{3})", text)
    report = re.search(r"\[(\d+ October) report\]\(([^)]+)\)", text)
    return {"fail": int(fail[1]) if fail else None, "pending": int(pend[1]) if pend else None,
            "decision": decision[1] if decision else "", "report_date": report[1] if report else "",
            "report_file": report[2].replace("%20", " ") if report else ""}


# --------------------------------------------------------------- mathcad --
def mathcad_beams():
    sys.path.insert(0, str(ROOT))
    from geometry_tables.sources import mathcad_beam_data
    return {span: mathcad_beam_data(span) for span in (1, 2)}


# -------------------------------------------------------------- drawings --
def typical_section_sheet():
    """Sheet S-1, the superstructure typical section cut from the model."""
    man = read_json(SOURCES["TYPSEC-MANIFEST"])
    return {"number": man["sheet"]["number"], "title": man["sheet"]["title"], "scales": [man["sheet"]["scale"]],
            "png": SOURCES["TYPSEC-PNG"], "generator": f"{man['generator']} {man['version']}",
            "model_sha256": man["model_sha256"], "dxf": man["dxf"], "section": man["section"], "cut": man["cut"]}


def drawing_sheets():
    man = read_json(SOURCES["DXF-MANIFEST"])
    sheets = []
    for i, s in enumerate(man["sheets"], 1):
        png = PREVIEWS / f"{i:02d}_{s['number']}.png"
        sheets.append({"number": s["number"], "title": s["title"], "scales": s.get("scales", []),
                       "png": png if png.exists() else None})
    return {"generator": f"{man['generator']} {man['version']}", "model_sha256": man["model_sha256"],
            "dxf": man["dxf"], "sheets": sheets}


# ------------------------------------------------------------- the bundle --
def collect():
    """Everything the report renders, as plain data."""
    J = Journal(SOURCES["JOURNAL"])
    facts = read_json(SOURCES["FACTS"])
    register = read_register(SOURCES["REGISTER"])
    decisions = read_decisions(SOURCES["DECISIONS"])
    assumptions = read_json(SOURCES["ASSUMPTIONS"])["entries"]
    review = read_latest_review(SOURCES["LATEST-REVIEW"])
    as_of = max(r["registered_on"] for r in register if r.get("registered_on"))

    cases = {k: read_json(SOURCES[k]) for k in ("PIER-CASE", "EB-CASE", "EB2-CASE", "BACKWALL-CASE",
                                                 "PIER-WALL-CASE", "CHEEK-CASE",
                                                 "PILE-EB1-CASE", "PILE-EB2-CASE", "PILE-PIER-CASE")}
    pile_tables = {k: read_pile_table(SOURCES[f"PILE-{k}-TABLE"]) for k in ("EB1", "EB2", "PIER")}
    pile_checks = {k: pile_check_summary(read_checks(SOURCES[f"PILE-{k}-CHECKS"])) for k in ("EB1", "EB2", "PIER")}
    wall_checks = generic_check_summary(read_checks(SOURCES["PIER-WALL-CHECKS"]))
    cheek_checks = generic_check_summary(read_checks(SOURCES["CHEEK-CHECKS"]), id_key="check", title_key="check")
    caps = {k: cap_readme_counts(SOURCES[f"{k}-README"]) for k in ("PIER", "EB", "EB2")}
    backwall_readme = SOURCES["BACKWALL-README"].read_text(encoding="utf-8")
    strap = read_json(SOURCES["STRAP-LOADS"])
    beams = mathcad_beams()
    geotech = read_json(SOURCES["GEOTECH-SNAPSHOT"])
    geotech_manifest = read_json(SOURCES["GEOTECH-MANIFEST"])
    drawings = drawing_sheets()
    typsec = typical_section_sheet()
    plan = read_json(SOURCES["PLAN-MANIFEST"])

    fbmp_readme = FBMP_README.read_text(encoding="utf-8") if FBMP_README.exists() else ""

    return {
        "as_of": as_of, "facts": facts, "journal": J, "register": register, "decisions": decisions,
        "assumptions": assumptions, "review": review, "cases": cases, "pile_tables": pile_tables,
        "pile_checks": pile_checks, "wall_checks": wall_checks, "cheek_checks": cheek_checks, "caps": caps,
        "backwall_readme": backwall_readme, "strap": strap, "beams": beams, "geotech": geotech,
        "geotech_manifest": geotech_manifest, "drawings": drawings, "typsec": typsec, "plan": plan,
        "fbmp_readme": fbmp_readme,
    }


def source_hashes():
    return {k: {"path": str(p.relative_to(ROOT)).replace("\\", "/"), "sha256": sha256(p)}
            for k, p in SOURCES.items()}
